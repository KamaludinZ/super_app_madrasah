"""Impor paket data wilayah (CSV/XLSX: kode, nama, kode_pos) ke master wilayah.

Aturan baris HARUS sama dengan `pratinjauPaketWilayah` di frontend/src/lib/wilayah.js:
baris kosong dilewati, baris pertama yang kodenya bukan kode wilayah dianggap judul, kode harus
berformat Kemendagri, nama wajib, kode pos (opsional) 5 digit, kode ganda dalam satu berkas ditolak.
Unggah ulang menambah/memperbarui data (upsert per kode) tanpa menghapus yang sudah ada.
"""
import csv
import io
import re
from datetime import datetime, timezone
from typing import Dict, List

from pymongo import UpdateOne

from wilayah_master import KOLEKSI_WILAYAH, POLA_KODE_WILAYAH, TINGKAT_WILAYAH, dokumen_wilayah, ensure_wilayah_indexes

BATAS_GALAT = 200
UKURAN_BATCH = 1000


def _baris_csv(content: bytes) -> List[List[str]]:
    teks = content.decode('utf-8-sig', errors='replace')
    try:
        dialek = csv.Sniffer().sniff(teks[:4096], delimiters=',;\t')
    except csv.Error:
        dialek = csv.excel
    return [list(r) for r in csv.reader(io.StringIO(teks), dialek)]


def _baris_xlsx(content: bytes) -> List[List[str]]:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    out = []
    for row in ws.iter_rows(max_col=3, values_only=True):
        out.append(['' if v is None else (str(int(v)) if isinstance(v, float) and v.is_integer() else str(v)) for v in row])
    wb.close()
    return out


def baca_paket_wilayah(content: bytes, nama_berkas: str) -> Dict:
    """-> {valid: [(kode, nama, kode_pos)], salah: [{baris, pesan}], per_tingkat}"""
    if (nama_berkas or '').lower().endswith('.xlsx') or content[:2] == b'PK':
        rows = _baris_xlsx(content)
    else:
        rows = _baris_csv(content)
    valid, salah, ada = [], [], set()
    per_tingkat = {t: 0 for t in TINGKAT_WILAYAH}
    for i, row in enumerate(rows):
        sel = [str(x or '').strip() for x in (list(row) + ['', '', ''])[:3]]
        kode, nama, kode_pos = sel
        if not kode and not nama:
            continue
        if i == 0 and not POLA_KODE_WILAYAH.match(kode):
            continue  # baris judul
        nomor = i + 1
        if not POLA_KODE_WILAYAH.match(kode):
            salah.append({'baris': nomor, 'pesan': f'Kode "{kode}" tidak sesuai format (mis. 35, 35.73, 35.73.05, 35.73.05.1001)'})
            continue
        if not nama:
            salah.append({'baris': nomor, 'pesan': f'Nama untuk kode {kode} kosong'})
            continue
        if kode_pos and not re.fullmatch(r'\d{5}', kode_pos):
            salah.append({'baris': nomor, 'pesan': f'Kode pos "{kode_pos}" harus 5 digit'})
            continue
        if kode in ada:
            salah.append({'baris': nomor, 'pesan': f'Kode {kode} ganda'})
            continue
        ada.add(kode)
        valid.append((kode, nama, kode_pos))
        per_tingkat[TINGKAT_WILAYAH[len(kode.split('.')) - 1]] += 1
    return {'valid': valid, 'salah': salah, 'per_tingkat': per_tingkat}


async def impor_paket_wilayah(database, content: bytes, nama_berkas: str) -> Dict:
    """Upsert paket ke koleksi wilayah. -> {total, baru, diperbarui, gagal, yatim, per_tingkat, errors}
    `yatim` = wilayah dari paket yang induknya belum ada di master (mis. desa tanpa kecamatan)."""
    data = baca_paket_wilayah(content, nama_berkas)
    # Upsert per kode butuh index kode (dibuat saat startup); pastikan ada agar tidak memindai seluruh koleksi.
    await ensure_wilayah_indexes(database)
    sekarang = datetime.now(timezone.utc).isoformat()
    koll = database[KOLEKSI_WILAYAH]
    baru = diperbarui = 0
    for awal in range(0, len(data['valid']), UKURAN_BATCH):
        ops = []
        for kode, nama, kode_pos in data['valid'][awal:awal + UKURAN_BATCH]:
            doc = dokumen_wilayah(kode, nama, kode_pos or None, sekarang)
            doc['sumber'] = 'paket'
            if doc['tingkat'] == 'desa' and not kode_pos:
                doc.pop('kode_pos')  # kode pos kosong di paket tidak menghapus kode pos yang sudah ada
            ops.append(UpdateOne({'kode': kode}, {'$set': doc}, upsert=True))
        if ops:
            r = await koll.bulk_write(ops, ordered=False)
            baru += r.upserted_count
            diperbarui += r.matched_count

    induk_dibutuhkan = {k.rsplit('.', 1)[0] for k, _, _ in data['valid'] if '.' in k}
    induk_ada = set()
    daftar = list(induk_dibutuhkan)
    for awal in range(0, len(daftar), 5000):
        induk_ada |= {d['kode'] async for d in koll.find({'kode': {'$in': daftar[awal:awal + 5000]}}, {'_id': 0, 'kode': 1})}
    yatim = sum(1 for k, _, _ in data['valid'] if '.' in k and k.rsplit('.', 1)[0] not in induk_ada)

    return {'total': len(data['valid']) + len(data['salah']), 'baru': baru, 'diperbarui': diperbarui,
            'gagal': len(data['salah']), 'yatim': yatim, 'per_tingkat': data['per_tingkat'],
            'errors': data['salah'][:BATAS_GALAT]}
