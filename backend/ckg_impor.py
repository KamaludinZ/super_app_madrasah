"""Pembaca & pemroses impor Data CKG dari berkas template.

Mendukung dua format:
- 'baku' : template baru — 16 kolom baku (NO ... GDS) + Tanggal Periksa + ID Pasien (kolom_template_ckg).
- 'lama' : template lama (kolom pertama "ID (jangan diubah)") — tetap diterima agar berkas yang
           sudah terlanjur diisi tidak sia-sia.
"""
import io
import re
from datetime import date, datetime
from typing import Any, Dict, List, Optional

from openpyxl import load_workbook

from ckg_kolom import kolom_template_ckg

SHEET_TEMPLATE_CKG = 'Template CKG'
JUDUL_TEMPLATE_LAMA = 'ID (jangan diubah)'
KOLOM_HASIL_CKG = ('bb', 'tb', 'td', 'jumlah_karies', 'visus_mata', 'kesehatan_kulit', 'fungsi_pendengaran', 'hemoglobin', 'gds')
# Kolom template lama -> key baku/field lama (urutan kolom template lama).
KOLOM_TEMPLATE_LAMA = ['pasien_id', 'identitas', 'nama_lengkap', 'kelas', 'tanggal', 'tb', 'bb', 'td', 'nadi', 'suhu', 'spo2',
                       'pemeriksaan_mata', 'pemeriksaan_gigi', 'kesimpulan', 'rekomendasi', 'keterangan']


def teks_sel(v: Any) -> str:
    if v is None:
        return ''
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def cek_header_ckg(header: List[Any]) -> List[str]:
    got = [teks_sel(h) for h in header]
    while got and got[-1] == '':
        got.pop()
    want = [k['label'] for k in kolom_template_ckg()]
    if got == want:
        return []
    hilang = [x for x in want if x not in got]
    asing = [x for x in got if x and x not in want]
    masalah = []
    if hilang:
        masalah.append('Kolom tidak ditemukan: ' + ', '.join(hilang[:10]))
    if asing:
        masalah.append('Kolom tidak dikenal: ' + ', '.join(asing[:10]))
    if not masalah:
        masalah.append('Urutan kolom berbeda dari template CKG; unduh template terbaru')
    return masalah


def baca_berkas_ckg(content: bytes) -> Dict[str, Any]:
    """{format: 'baku'|'lama', masalah_header, baris: [{baris, data}], dilewati: {kosong, belum_diperiksa}}.
    Melempar ValueError bila bukan berkas Excel yang valid."""
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        raise ValueError('Berkas bukan Excel .xlsx yang valid') from e
    ws = wb[SHEET_TEMPLATE_CKG] if SHEET_TEMPLATE_CKG in wb.sheetnames else wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = list(next(rows, []) or [])
    lama = teks_sel(header[0] if header else '') == JUDUL_TEMPLATE_LAMA
    keys = KOLOM_TEMPLATE_LAMA if lama else [k['key'] for k in kolom_template_ckg()]
    hasil: Dict[str, Any] = {'format': 'lama' if lama else 'baku', 'masalah_header': [] if lama else cek_header_ckg(header),
                             'baris': [], 'dilewati': {'kosong': 0, 'belum_diperiksa': 0}}
    if hasil['masalah_header']:
        wb.close()
        return hasil
    hasil_kunci = ('tb', 'bb', 'td', 'nadi', 'suhu', 'spo2', 'pemeriksaan_mata', 'pemeriksaan_gigi', 'kesimpulan') if lama else KOLOM_HASIL_CKG
    for nomor, row in enumerate(rows, 2):
        row = list(row or [])
        data = {k: teks_sel(row[i]) if i < len(row) else '' for i, k in enumerate(keys)}
        if not any(data.values()):
            hasil['dilewati']['kosong'] += 1
            continue
        if not any(data.get(k) for k in hasil_kunci):
            hasil['dilewati']['belum_diperiksa'] += 1
            continue
        hasil['baris'].append({'baris': nomor, 'data': data})
    wb.close()
    return hasil


def angka(teks: str, cast=float) -> Optional[float]:
    """Teks angka (koma/titik desimal) -> float/int; '' -> None; ValueError bila bukan angka."""
    t = (teks or '').strip().replace(',', '.')
    if t == '':
        return None
    v = float(t)
    return int(v) if cast is int else v


def dokumen_ckg(data: Dict[str, str]) -> Dict[str, Any]:
    """Field pemeriksaan dokumen uks_ckg dari satu baris berkas (format baku atau lama)."""
    doc = {
        'tanggal': data.get('tanggal') or None,
        'berat_badan': angka(data.get('bb', '')),
        'tinggi_badan': angka(data.get('tb', '')),
        'tekanan_darah': re.sub(r'\s+', '', data.get('td', '')) or None,
        'jumlah_karies': angka(data.get('jumlah_karies', ''), int),
        'visus_mata': data.get('visus_mata') or None,
        'kesehatan_kulit': data.get('kesehatan_kulit') or None,
        'fungsi_pendengaran': data.get('fungsi_pendengaran') or None,
        'hemoglobin': angka(data.get('hemoglobin', '')),
        'gds': angka(data.get('gds', '')),
    }
    for k in ('nadi', 'spo2'):
        if k in data:
            doc[k] = angka(data.get(k, ''), int)
    if 'suhu' in data:
        doc['suhu'] = angka(data.get('suhu', ''))
    for k in ('pemeriksaan_mata', 'pemeriksaan_gigi', 'kesimpulan', 'rekomendasi', 'keterangan'):
        if k in data:
            doc[k] = data.get(k) or None
    return doc


# ------------------------------------------------------------
# VALIDASI BARIS (sama dengan frontend validasiBarisCkg)
# ------------------------------------------------------------
BATAS_ANGKA_CKG = {
    'bb': ('BB', 1, 300), 'tb': ('TB', 30, 250), 'jumlah_karies': ('Jumlah Karies', 0, 32),
    'hemoglobin': ('Hemoglobin', 0, 30), 'gds': ('GDS', 0, 1000),
}
_POLA_TANGGAL = re.compile(r'^\d{4}-\d{2}-\d{2}$')
_POLA_TD = re.compile(r'^\d{2,3}\s*/\s*\d{2,3}$')


def validasi_baris_ckg(data: Dict[str, str], hari_ini: Optional[str] = None) -> List[Dict[str, str]]:
    """Daftar kesalahan [{kolom, pesan}] satu baris impor CKG; kosong = valid.
    Aturan & pesan identik dengan frontend (lib/ckgImpor.js) agar pratinjau = hasil server."""
    hari_ini = hari_ini or datetime.utcnow().date().isoformat()
    salah: List[Dict[str, str]] = []
    nik = data.get('nik') or ''
    if not data.get('pasien_id') and not nik:
        salah.append({'kolom': 'ID Pasien / NIK', 'pesan': 'Identitas pasien kosong; isi ID Pasien (dari template) atau NIK'})
    if nik and not re.fullmatch(r'\d{16}', re.sub(r'\D', '', nik)):
        salah.append({'kolom': 'NIK', 'pesan': 'NIK harus 16 digit angka'})
    tgl = data.get('tanggal') or ''
    if not tgl:
        salah.append({'kolom': 'Tanggal Periksa', 'pesan': 'Tanggal pemeriksaan wajib diisi'})
    else:
        valid = bool(_POLA_TANGGAL.match(tgl))
        if valid:
            try:
                datetime.strptime(tgl, '%Y-%m-%d')
            except ValueError:
                valid = False
        if not valid:
            salah.append({'kolom': 'Tanggal Periksa', 'pesan': 'Format tanggal harus TAHUN-BULAN-TANGGAL (mis. 2026-09-01)'})
        elif tgl > hari_ini:
            salah.append({'kolom': 'Tanggal Periksa', 'pesan': 'Tanggal pemeriksaan tidak boleh di masa depan'})
    for key, (label, lo, hi) in BATAS_ANGKA_CKG.items():
        v = data.get(key) or ''
        if v == '':
            continue
        try:
            n = float(v.replace(',', '.'))
        except ValueError:
            salah.append({'kolom': label, 'pesan': f'{label} harus berupa angka'})
            continue
        if n < lo or n > hi:
            salah.append({'kolom': label, 'pesan': f'{label} di luar batas wajar ({lo}-{hi})'})
    td = data.get('td') or ''
    if td and not _POLA_TD.match(td):
        salah.append({'kolom': 'TD', 'pesan': 'Tekanan darah harus berformat sistolik/diastolik, mis. 110/70'})
    return salah
