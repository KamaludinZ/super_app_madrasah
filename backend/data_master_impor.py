"""Skema riwayat impor pelengkapan data master (Siswa / GTK) dan jejak audit per field.

Koleksi:
- data_master_impor           : satu dokumen per sesi impor (satu berkas yang diunggah).
- data_master_impor_perubahan : satu dokumen per field yang diubah impor (nilai lama -> baru),
                                agar setiap perubahan hasil impor dapat ditelusuri & dipulihkan.
- data_master_impor_baris     : satu dokumen per baris berkas yang diproses (berhasil, tanpa
                                perubahan, atau gagal beserta alasannya) — jejak lengkap tanpa batas.
"""
import re
import uuid
from datetime import datetime
from typing import Any, Dict, List, Literal, Optional

from pydantic import BaseModel, Field

KOLEKSI_IMPOR = 'data_master_impor'
KOLEKSI_PERUBAHAN = 'data_master_impor_perubahan'
KOLEKSI_BARIS = 'data_master_impor_baris'

JenisImpor = Literal['siswa', 'gtk']
ModeImpor = Literal['isi_kosong', 'timpa']
StatusBaris = Literal['berhasil', 'tanpa_perubahan', 'gagal']


class RingkasanImpor(BaseModel):
    total: int = 0            # baris yang dikirim ke server
    berhasil: int = 0         # baris dengan minimal satu field berubah
    tanpa_perubahan: int = 0  # baris cocok tetapi tidak ada field yang perlu diubah
    gagal: int = 0            # baris ditolak (identitas tidak cocok / nilai tidak valid)
    field_diisi: int = 0      # field kosong yang diisi
    field_ditimpa: int = 0    # field berisi yang diganti (mode timpa)
    # Pencocokan alamat ke master wilayah (baris yang mengubah nama wilayah)
    wilayah_cocok: int = 0
    wilayah_sebagian: int = 0
    wilayah_tidak_cocok: int = 0


class ImporModel(BaseModel):
    """Sesi impor: dibuat pada batch pertama, diperbarui tiap batch berikutnya."""
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    jenis: JenisImpor
    mode: ModeImpor = 'isi_kosong'
    nama_berkas: Optional[str] = None
    status: Literal['berjalan', 'selesai'] = 'berjalan'
    ringkasan: RingkasanImpor = Field(default_factory=RingkasanImpor)
    # Kesalahan per baris (dibatasi) untuk ditampilkan ulang di riwayat impor.
    kesalahan: List[Dict[str, Any]] = Field(default_factory=list)
    dibuat_oleh: Optional[str] = None
    dibuat_oleh_nama: Optional[str] = None
    dibuat_pada: datetime = Field(default_factory=datetime.utcnow)
    diperbarui_pada: datetime = Field(default_factory=datetime.utcnow)


class PerubahanModel(BaseModel):
    """Jejak audit satu field yang diubah oleh impor."""
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    impor_id: str
    jenis: JenisImpor
    target_id: str            # users.id siswa / GTK
    baris: int                # nomor baris Excel sumber
    kolom: str                # key kolom skema (mis. 'ayah_nama')
    path: str                 # lokasi tersimpan (mis. 'detail.ayah.nama' atau 'user.phone')
    aksi: Literal['isi', 'timpa']
    nilai_lama: Any = None
    nilai_baru: Any = None
    oleh: Optional[str] = None
    pada: datetime = Field(default_factory=datetime.utcnow)


class BarisImporModel(BaseModel):
    """Jejak hasil satu baris berkas impor."""
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    impor_id: str
    jenis: JenisImpor
    baris: int
    identitas: str = '-'
    target_id: Optional[str] = None
    status: StatusBaris
    kolom: Optional[str] = None
    pesan: Optional[str] = None
    field_diisi: int = 0
    field_ditimpa: int = 0
    oleh: Optional[str] = None
    pada: datetime = Field(default_factory=datetime.utcnow)


MAKS_KESALAHAN_TERSIMPAN = 500


async def ensure_data_master_impor_indexes(database) -> List[str]:
    """Index riwayat & jejak impor (idempotent). Mengembalikan daftar index yang gagal."""
    specs = [
        (KOLEKSI_IMPOR, 'id', {'unique': True}),
        (KOLEKSI_IMPOR, [('jenis', 1), ('dibuat_pada', -1)], {}),
        (KOLEKSI_PERUBAHAN, 'id', {'unique': True}),
        (KOLEKSI_PERUBAHAN, [('impor_id', 1), ('baris', 1)], {}),
        (KOLEKSI_PERUBAHAN, [('target_id', 1), ('pada', -1)], {}),
        (KOLEKSI_BARIS, 'id', {'unique': True}),
        (KOLEKSI_BARIS, [('impor_id', 1), ('status', 1), ('baris', 1)], {}),
    ]
    gagal = []
    for coll, keys, kwargs in specs:
        try:
            await database[coll].create_index(keys, **kwargs)
        except Exception as e:  # noqa: BLE001 - dilaporkan ke pemanggil
            gagal.append(f'{coll} {keys}: {e}')
    return gagal


def dokumen(model: BaseModel) -> Dict[str, Any]:
    """Model -> dokumen Mongo (datetime disimpan ISO agar seragam dengan koleksi lain)."""
    d = model.model_dump()
    for k, v in list(d.items()):
        if isinstance(v, datetime):
            d[k] = v.isoformat()
    return d


# ------------------------------------------------------------
# PEMBACA BERKAS IMPOR (sisi server)
# ------------------------------------------------------------
MAKS_UKURAN_BERKAS = 5 * 1024 * 1024  # 5 MB, sama dengan frontend
MAKS_BARIS = 5000


def baca_berkas_impor(content: bytes, jenis: str) -> Dict[str, Any]:
    """Baca berkas .xlsx impor: cek judul kolom terhadap skema, lewati baris keterangan
    & baris kosong. Mengembalikan {sheet, masalah_header, baris: [{baris, data}], dilewati}.
    Melempar ValueError bila berkas bukan Excel yang valid."""
    import io
    from openpyxl import load_workbook
    from data_master_excel import SKEMA, adalah_baris_keterangan, cek_header, nilai_sel

    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001 - zip rusak, bukan xlsx, dll.
        raise ValueError('Berkas bukan Excel .xlsx yang valid') from e
    sk = SKEMA[jenis]
    kolom = sk['kolom']
    ws = wb[sk['sheet']] if sk['sheet'] in wb.sheetnames else wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = list(next(rows, []) or [])
    hasil: Dict[str, Any] = {'sheet': ws.title, 'masalah_header': cek_header(header, kolom), 'baris': [],
                             'dilewati': {'kosong': 0, 'keterangan': 0}}
    if hasil['masalah_header']:
        wb.close()
        return hasil
    for nomor, row in enumerate(rows, 2):
        row = list(row or [])
        if adalah_baris_keterangan(row):
            hasil['dilewati']['keterangan'] += 1
            continue
        sel = [nilai_sel(row[i]).strip() if i < len(row) else '' for i in range(len(kolom))]
        if not any(sel):
            hasil['dilewati']['kosong'] += 1
            continue
        hasil['baris'].append({'baris': nomor, 'data': {k['key']: sel[i] for i, k in enumerate(kolom)}})
    wb.close()
    return hasil


# ------------------------------------------------------------
# PENCOCOKAN BARIS VIA KOLOM IDENTITAS
# ------------------------------------------------------------
# Kolom identitas per jenis (sama dengan kolom 'kunci' skema): ID dulu, lalu NISN/NIP.
KUNCI_IDENTITAS = {'siswa': ('id', 'nisn'), 'gtk': ('id', 'nip')}
LABEL_IDENTITAS = {'id': 'ID', 'nisn': 'NISN', 'nip': 'NIP'}


def normalisasi_identitas(field: str, nilai: Any) -> str:
    """Rapikan nilai identitas dari Excel: buang spasi; NISN yang berupa angka dan
    kehilangan nol di depan (diubah Excel menjadi angka) dikembalikan ke 10 digit."""
    t = re.sub(r'\s+', '', str(nilai or ''))
    if t.endswith('.0') and t[:-2].isdigit():
        t = t[:-2]
    if field == 'nisn' and t.isdigit() and len(t) < 10:
        t = t.zfill(10)
    return t


def _query_target(jenis: str) -> Dict[str, Any]:
    from data_master_service import GURU_ROLES, TENDIK_ROLES
    roles = ['siswa'] if jenis == 'siswa' else list(GURU_ROLES + TENDIK_ROLES)
    return {'roles': {'$in': roles}, 'mutation_type': {'$ne': 'keluar'}}


async def cocokkan_baris(db, jenis: str, baris: List[Dict[str, Any]], sudah: Optional[Dict[str, int]] = None) -> List[Dict[str, Any]]:
    """Cocokkan tiap baris impor ke pengguna yang ada.

    Aturan:
    - ID diisi -> dicari berdasarkan ID; bila NISN/NIP di berkas juga diisi dan berbeda dengan
      yang tersimpan, baris ditolak (mencegah data tertukar).
    - ID kosong -> dicari berdasarkan NISN (siswa) / NIP (GTK); harus tepat satu pengguna.
    - Hanya siswa / GTK aktif (bukan mutasi keluar) yang bisa dicocokkan.
    - Satu pengguna hanya boleh muncul sekali per berkas; `sudah` = {target_id: baris pertama}
      dari batch sebelumnya agar duplikat lintas batch juga terdeteksi (diperbarui di tempat).

    Mengembalikan per baris: {baris, identitas, target (dok users | None), pesan, kolom}.
    """
    kunci_id, kunci_no = KUNCI_IDENTITAS[jenis]
    sudah = {} if sudah is None else sudah
    q = _query_target(jenis)
    ids = {normalisasi_identitas('id', b['data'].get(kunci_id)) for b in baris} - {''}
    nos = {normalisasi_identitas(kunci_no, b['data'].get(kunci_no)) for b in baris} - {''}
    proj = {'_id': 0, 'password_hash': 0}
    by_id = {u['id']: u for u in await db.users.find({**q, 'id': {'$in': list(ids)}}, proj).to_list(len(ids) or 1)} if ids else {}
    by_no: Dict[str, List[Dict]] = {}
    if nos:
        for u in await db.users.find({**q, kunci_no: {'$in': list(nos)}}, proj).to_list(len(nos) * 2 or 1):
            by_no.setdefault(normalisasi_identitas(kunci_no, u.get(kunci_no)), []).append(u)

    hasil = []
    lab_id, lab_no = LABEL_IDENTITAS[kunci_id], LABEL_IDENTITAS[kunci_no]
    for b in baris:
        uid = normalisasi_identitas('id', b['data'].get(kunci_id))
        no = normalisasi_identitas(kunci_no, b['data'].get(kunci_no))
        identitas = ' / '.join(x for x in (uid, no) if x) or '-'
        r = {'baris': b['baris'], 'identitas': identitas, 'target': None, 'pesan': None, 'kolom': None}
        if not uid and not no:
            r.update(pesan=f'Kolom identitas kosong; isi {lab_id} atau {lab_no}', kolom=f'{lab_id} / {lab_no}')
        elif uid:
            u = by_id.get(uid)
            tersimpan = normalisasi_identitas(kunci_no, (u or {}).get(kunci_no))
            if not u:
                r.update(pesan=f'{lab_id} {uid} tidak ditemukan (atau bukan {jenis} aktif)', kolom=lab_id)
            elif no and tersimpan and no != tersimpan:
                r.update(pesan=f'{lab_no} {no} tidak sesuai dengan data {lab_id} ini ({lab_no} tersimpan {tersimpan}); periksa apakah baris tertukar',
                         kolom=lab_no)
            else:
                r['target'] = u
        else:
            cocok = by_no.get(no, [])
            if not cocok:
                r.update(pesan=f'{lab_no} {no} tidak ditemukan (atau bukan {jenis} aktif)', kolom=lab_no)
            elif len(cocok) > 1:
                r.update(pesan=f'{lab_no} {no} dimiliki {len(cocok)} pengguna; isi kolom {lab_id} agar tidak tertukar', kolom=lab_no)
            else:
                r['target'] = cocok[0]
        if r['target']:
            tid = r['target']['id']
            if tid in sudah and sudah[tid] != b['baris']:
                r.update(target=None, pesan=f"Data yang sama sudah ada di baris {sudah[tid]}; baris ini dilewati", kolom=lab_id)
            else:
                sudah.setdefault(tid, b['baris'])
        hasil.append(r)
    return hasil


# ------------------------------------------------------------
# RENCANA & PENERAPAN PERUBAHAN FIELD
# ------------------------------------------------------------
# Kolom yang tidak pernah ditulis impor: penanda baris & nilai turunan.
KOLOM_TIDAK_DITULIS = {'id', 'nisn', 'nip', 'kelas', 'jenis'}
FIELD_16_DIGIT = {'nik', 'nomor_kk'}
FIELD_TANGGAL = {'birth_date', 'tgl_lahir', 'tmt_pns', 'tanggal_sk_pns', 'tmt_pegawai', 'tmt_guru'}
_POLA_TANGGAL = re.compile(r'^\d{4}-\d{2}-\d{2}$')


def _kosong(v: Any) -> bool:
    return v is None or v == '' or v == [] or (isinstance(v, str) and not v.strip())


def _setara(a: Any, b: Any) -> bool:
    if isinstance(a, list) or isinstance(b, list):
        return [str(x) for x in (a or [])] == [str(x) for x in (b or [])]
    if isinstance(a, bool) or isinstance(b, bool):
        return a == b
    return str(a if a is not None else '').strip() == str(b if b is not None else '').strip()


def _nilai_tersimpan(kolom: Dict[str, Any], user: Dict[str, Any], detail: Optional[Dict[str, Any]]) -> Any:
    from data_master_excel import FALLBACK_GTK, _ambil
    sumber, *rest = kolom['path'].split('.')
    v = _ambil(user if sumber == 'user' else (detail or {}), rest)
    if _kosong(v) and sumber == 'user' and rest[-1] in FALLBACK_GTK:
        v = user.get(FALLBACK_GTK[rest[-1]])
    return v


def ubah_nilai_impor(kolom: Dict[str, Any], teks: Any) -> Any:
    """Teks sel -> nilai tersimpan, dengan validasi. None = sel kosong (tidak mengubah apa pun).
    Melempar ValueError berisi pesan berbahasa Indonesia bila isi tidak valid."""
    from data_master_excel import PILIHAN_GTK, normalisasi_pilihan, parse_nilai
    nilai = parse_nilai(kolom, teks)
    if nilai is None:
        return None
    field = kolom['path'].split('.')[-1]
    # Pilihan berkode GTK hanya untuk field akun (users). Field senama di detail siswa punya isian lain,
    # mis. alamat_siswa.status_tempat_tinggal = "Tinggal dengan Ayah Kandung" (bukan milik_sendiri/sewa/...).
    if field in PILIHAN_GTK and kolom['path'].startswith('user.') and isinstance(nilai, str):
        try:
            nilai = normalisasi_pilihan(field, nilai)
        except ValueError as e:
            raise ValueError(f"{kolom['label']}: {e}")
    if field in FIELD_16_DIGIT and isinstance(nilai, str):
        digit = re.sub(r'\D', '', nilai)
        if len(digit) != 16:
            raise ValueError(f"{kolom['label']}: harus 16 digit angka")
        nilai = digit
    if field in FIELD_TANGGAL and isinstance(nilai, str):
        nilai = nilai[:10]
        if not _POLA_TANGGAL.match(nilai):
            raise ValueError(f"{kolom['label']}: format tanggal harus TAHUN-BULAN-TANGGAL (mis. 2012-05-10)")
    return nilai


def rencanakan_perubahan(jenis: str, data: Dict[str, str], user: Dict[str, Any], detail: Optional[Dict[str, Any]],
                         mode: str = 'isi_kosong') -> Dict[str, Any]:
    """Bandingkan isi satu baris berkas dengan data tersimpan.

    - Sel kosong di berkas tidak pernah mengubah/menghapus data.
    - mode 'isi_kosong': hanya field yang tersimpan kosong yang diisi.
    - mode 'timpa'     : field kosong diisi + field berisi yang nilainya berbeda diganti.
    Mengembalikan {perubahan: [{kolom, label, path, aksi, lama, baru}], kesalahan: [{kolom, pesan}]}.
    """
    from data_master_excel import SKEMA
    perubahan, kesalahan = [], []
    for k in SKEMA[jenis]['kolom']:
        if k['key'] in KOLOM_TIDAK_DITULIS:
            continue
        try:
            baru = ubah_nilai_impor(k, data.get(k['key']))
        except ValueError as e:
            kesalahan.append({'kolom': k['label'], 'pesan': str(e)})
            continue
        if baru is None:
            continue
        lama = _nilai_tersimpan(k, user, detail)
        if _kosong(lama):
            aksi = 'isi'
        elif mode == 'timpa' and not _setara(lama, baru):
            aksi = 'timpa'
        else:
            continue
        perubahan.append({'kolom': k['key'], 'label': k['label'], 'path': k['path'], 'aksi': aksi, 'lama': lama, 'baru': baru})
    return {'perubahan': perubahan, 'kesalahan': kesalahan}


def _set_detail(perubahan: List[Dict[str, Any]], detail: Optional[Dict[str, Any]]) -> Dict[str, Any]:
    """Susun $set student_details. Path bersarang (mis. ayah.nama) ditulis langsung bila objek
    induknya sudah berupa dict; bila induk kosong/None (data lama), seluruh objek induk ditulis
    agar MongoDB tidak gagal membuat field di dalam nilai null."""
    detail = detail or {}
    hasil: Dict[str, Any] = {}
    for p in perubahan:
        if not p['path'].startswith('detail.'):
            continue
        parts = p['path'][7:].split('.')
        if len(parts) == 1:
            hasil[parts[0]] = p['baru']
        elif isinstance(detail.get(parts[0]), dict):
            hasil['.'.join(parts)] = p['baru']
        else:
            induk = hasil.setdefault(parts[0], {})
            induk[parts[1]] = p['baru']
    return hasil


async def terapkan_perubahan(db, jenis: str, target: Dict[str, Any], perubahan: List[Dict[str, Any]], oleh: Optional[str],
                             detail: Optional[Dict[str, Any]] = None) -> Optional[str]:
    """Tulis perubahan ke users ($set field) dan student_details ($set path bersarang, upsert).
    Nama wilayah alamat yang berubah dicocokkan ke master wilayah (kode wilayah + nama resmi ikut disimpan);
    mengembalikan status pencocokan wilayah baris ini ('cocok'/'sebagian'/'tidak_cocok') atau None."""
    sekarang = datetime.utcnow()
    set_user = {p['path'][5:]: p['baru'] for p in perubahan if p['path'].startswith('user.')}
    if jenis == 'gtk':
        # Nama & gelar GTK: full_name disusun ulang dari nama tanpa gelar + gelar (gabung data tersimpan);
        # mengalahkan kolom Nama Lengkap di berkas agar susunan nama selalu konsisten.
        from nama_gelar import lengkapi_nama
        lengkapi_nama(set_user, target)
    set_detail = _set_detail(perubahan, detail)
    from wilayah_cocok import cocokkan_alamat_impor
    status_wilayah = await cocokkan_alamat_impor(db, jenis, target, detail, set_user, set_detail)
    if set_user:
        set_user['updated_at'] = sekarang.isoformat()
        await db.users.update_one({'id': target['id']}, {'$set': set_user})
    if set_detail:
        if jenis != 'siswa':
            raise ValueError('field detail hanya untuk siswa')
        await db.student_details.update_one(
            {'student_id': target['id']},
            {'$set': {**set_detail, 'updated_at': sekarang, 'updated_by': oleh},
             '$setOnInsert': {'id': str(uuid.uuid4()), 'student_id': target['id'], 'created_at': sekarang}},
            upsert=True,
        )
    return status_wilayah
