"""Susunan kolom baku & pembuat berkas Excel Data Siswa / Data GTK (unduh, template, impor).

Urutan, key, dan judul kolom HARUS identik dengan frontend
(frontend/src/lib/dataSiswaKolom.js dan dataGtkKolom.js) agar ekspor–impor akurat.
"""
import io
import re
from typing import Any, Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter

# ------------------------------------------------------------
# KOLOM DATA SISWA
# ------------------------------------------------------------
_ORTU_FIELDS = [
    ('nama', 'Nama'), ('status', 'Status'), ('citizenship', 'Kewarganegaraan'), ('nik', 'NIK'),
    ('asal_negara', 'Asal Negara (WNA)'), ('nomor_izin_tinggal', 'Nomor Izin Tinggal (WNA)'),
    ('tempat_lahir', 'Tempat Lahir'), ('tgl_lahir', 'Tanggal Lahir'), ('pendidikan', 'Pendidikan'),
    ('pekerjaan', 'Pekerjaan'), ('penghasilan', 'Penghasilan'), ('no_hp', 'No. HP'),
    ('no_hp_unavailable', 'Tidak Punya Nomor HP (Ya/Tidak)'),
]


def _k(key: str, label: str, grup: str, path: str, **extra) -> Dict[str, Any]:
    return {'key': key, 'label': label, 'grup': grup, 'path': path, **extra}


def _ortu(obj: str, label: str, extra=()) -> List[Dict[str, Any]]:
    return [_k(f'{obj}_{f}', f'{label} - {l}', 'Data Orang Tua', f'detail.{obj}.{f}')
            for f, l in list(_ORTU_FIELDS) + list(extra)]


_ALAMAT_FIELDS = [
    ('tinggal_luar_negeri', 'Tinggal di Luar Negeri (Ya/Tidak)'), ('status_kepemilikan', 'Status Kepemilikan Rumah'),
    ('provinsi', 'Provinsi'), ('kabupaten', 'Kabupaten/Kota'), ('kecamatan', 'Kecamatan'), ('kelurahan', 'Kelurahan/Desa'),
    ('rt', 'RT'), ('rw', 'RW'), ('kode_pos', 'Kode Pos'), ('alamat', 'Alamat Lengkap'),
]


def _alamat(obj: str, label: str, extra=()) -> List[Dict[str, Any]]:
    return [_k(f'{obj}_{f}', f'{label} - {l}', 'Data Alamat', f'detail.{obj}.{f}')
            for f, l in list(extra) + list(_ALAMAT_FIELDS)]


KOLOM_DATA_SISWA: List[Dict[str, Any]] = [
    _k('id', 'ID Siswa', 'Identitas', 'user.id', kunci=True),
    _k('nisn', 'NISN', 'Identitas', 'user.nisn', kunci=True),
    _k('nis', 'NIS', 'Identitas', 'user.nis'),
    _k('full_name', 'Nama Lengkap', 'Identitas', 'user.full_name'),
    _k('kelas', 'Kelas', 'Identitas', 'kelas'),
    _k('gender', 'Jenis Kelamin (L/P)', 'Data Siswa', 'user.gender'),
    _k('birth_place', 'Tempat Lahir', 'Data Siswa', 'user.birth_place'),
    _k('birth_date', 'Tanggal Lahir (YYYY-MM-DD)', 'Data Siswa', 'user.birth_date'),
    _k('email', 'Email Siswa', 'Data Siswa', 'user.email'),
    _k('phone', 'Nomor HP', 'Data Siswa', 'user.phone'),
    _k('no_hp_unavailable', 'Tidak Punya Nomor HP (Ya/Tidak)', 'Data Siswa', 'detail.no_hp_unavailable'),
    _k('citizenship', 'Kewarganegaraan', 'Data Siswa', 'detail.citizenship'),
    _k('nik', 'NIK', 'Data Siswa', 'detail.nik'),
    _k('asal_negara', 'Asal Negara (WNA)', 'Data Siswa', 'detail.asal_negara'),
    _k('nomor_izin_tinggal', 'Nomor Izin Tinggal (WNA)', 'Data Siswa', 'detail.nomor_izin_tinggal'),
    _k('jumlah_saudara', 'Jumlah Saudara', 'Data Siswa', 'detail.jumlah_saudara'),
    _k('anak_ke', 'Anak Ke-', 'Data Siswa', 'detail.anak_ke'),
    _k('agama', 'Agama', 'Data Siswa', 'detail.agama'),
    _k('cita_cita', 'Cita-cita', 'Data Siswa', 'detail.cita_cita'),
    _k('hobi', 'Hobi', 'Data Siswa', 'detail.hobi'),
    _k('pembiaya_sekolah', 'Yang Membiayai Sekolah', 'Data Siswa', 'detail.pembiaya_sekolah'),
    _k('pra_sekolah', 'Pra-Sekolah', 'Data Siswa', 'detail.pra_sekolah', daftar=True),
    _k('imunisasi', 'Imunisasi', 'Data Siswa', 'detail.imunisasi', daftar=True),
    _k('nomor_kip', 'Nomor KIP', 'Data Siswa', 'detail.nomor_kip'),
    _k('nomor_kk', 'Nomor KK', 'Data Siswa', 'detail.nomor_kk'),
    _k('nama_kepala_keluarga', 'Nama Kepala Keluarga', 'Data Siswa', 'detail.nama_kepala_keluarga'),
    _k('santri_mahad', 'Santri Mahad (Ya/Tidak)', 'Data Siswa', 'detail.santri_mahad'),
    _k('kamar_mahad', 'Kamar Mahad', 'Data Siswa', 'detail.kamar_mahad'),
    *_ortu('ayah', 'Ayah'),
    *_ortu('ibu', 'Ibu'),
    *_ortu('wali', 'Wali', [('hubungan_wali', 'Hubungan'), ('nomor_kks', 'Nomor KKS'), ('nomor_pkh', 'Nomor PKH')]),
    *_alamat('alamat_ayah', 'Alamat Ayah'),
    *_alamat('alamat_ibu', 'Alamat Ibu', [('sama_dengan_ayah', 'Sama dengan Alamat Ayah (Ya/Tidak)')]),
    *_alamat('alamat_wali', 'Alamat Wali', [('status_wali', 'Status Wali'), ('sama_dengan_ayah', 'Sama dengan Alamat Ayah (Ya/Tidak)')]),
    *[_k(f'alamat_siswa_{f}', f'Siswa - {l}', 'Data Alamat', f'detail.alamat_siswa.{f}') for f, l in [
        ('status_tempat_tinggal', 'Status Tempat Tinggal'), ('jarak_tempuh', 'Jarak Tempuh'),
        ('transportasi', 'Transportasi'), ('waktu_tempuh', 'Waktu Tempuh'),
    ]],
]


# ------------------------------------------------------------
# KOLOM DATA GTK (semua field di dokumen users)
# ------------------------------------------------------------
def _gtk(grup: str, daftar) -> List[Dict[str, Any]]:
    out = []
    for item in daftar:
        key, label = item[0], item[1]
        extra = item[2] if len(item) > 2 else {}
        out.append(_k(key, label, grup, f'user.{key}', **extra))
    return out


KOLOM_DATA_GTK: List[Dict[str, Any]] = [
    *_gtk('Identitas', [('id', 'ID GTK', {'kunci': True}), ('nip', 'NIP', {'kunci': True}), ('full_name', 'Nama Lengkap')]),
    _k('jenis', 'Jenis GTK', 'Identitas', 'jenis'),
    *_gtk('Nama & Gelar', [('gelar_depan', 'Gelar Depan'), ('nama_tanpa_gelar', 'Nama Tanpa Gelar'), ('gelar_belakang', 'Gelar Belakang')]),
    *_gtk('Data Diri', [
        ('gender', 'Jenis Kelamin (L/P)'), ('birth_place', 'Tempat Lahir'), ('birth_date', 'Tanggal Lahir (YYYY-MM-DD)'),
        ('nik', 'NIK'), ('nomor_kk', 'Nomor KK'), ('nama_ibu_kandung', 'Nama Ibu Kandung'), ('agama', 'Agama'),
    ]),
    *_gtk('Kepegawaian', [
        ('status_kepegawaian', 'Status Kepegawaian (pns/pppk/non_asn)'), ('peg_id', 'Peg ID'), ('nuptk', 'NUPTK'),
        ('npk', 'NPK'), ('nrg', 'NRG'), ('tmt_pns', 'TMT PNS (YYYY-MM-DD)'), ('no_sk_pns', 'No SK PNS'),
        ('tanggal_sk_pns', 'Tanggal SK (YYYY-MM-DD)'),
    ]),
    *_gtk('Informasi Lain', [
        ('phone', 'Nomor HP'), ('email', 'Email Pribadi'), ('email_madrasah', 'Email Madrasah Hebat'),
        ('bpjs_kesehatan', 'BPJS Kesehatan'), ('bpjs_ketenagakerjaan', 'BPJS Ketenagakerjaan'), ('npwp', 'NPWP'),
        ('golongan_darah', 'Golongan Darah'), ('nama_rekening', 'Nama Pemilik Rekening'),
        ('nomor_rekening', 'Nomor Rekening'), ('bank', 'Bank'),
    ]),
    *_gtk('Tempat Tinggal', [
        ('status_tempat_tinggal', 'Status Tempat Tinggal (milik_sendiri/sewa/menumpang/dinas)'), ('provinsi', 'Provinsi'),
        ('kab_kota', 'Kabupaten/Kota'),
        ('kecamatan', 'Kecamatan'), ('kelurahan', 'Kelurahan/Desa'), ('rt', 'RT'), ('rw', 'RW'), ('kode_pos', 'Kode Pos'),
        ('jarak_ke_sekolah', 'Jarak ke Sekolah (km)'), ('transportasi', 'Transportasi'),
        ('waktu_tempuh', 'Waktu Tempuh (menit)'), ('lintang', 'Lintang'), ('bujur', 'Bujur'),
    ]),
    *_gtk('Status Perkawinan', [
        ('status_perkawinan', 'Status Perkawinan (belum_menikah/menikah/duda/janda)'), ('nama_pasangan', 'Nama Suami/Istri'),
    ]),
    *_gtk('Penugasan', [
        ('jenis_ptk', 'Jenis PTK (guru_mapel/guru_bk/kepala_sekolah/tenaga_administrasi/pustakawan/laboran)'),
        ('tmt_pegawai', 'TMT Pegawai (YYYY-MM-DD)'), ('tmt_guru', 'TMT Guru/Tanggal SK PTK (YYYY-MM-DD)'),
        ('tugas_utama', 'Tugas Utama'), ('tugas_tambahan', 'Tugas Tambahan'),
    ]),
]

# Tipe nilai saat impor per NAMA FIELD (segmen terakhir path; default teks): mengubah teks
# Excel kembali ke tipe data tersimpan. Berlaku juga untuk field orang tua (mis. ayah.no_hp_unavailable).
TIPE_KOLOM: Dict[str, str] = {
    'jumlah_saudara': 'int', 'anak_ke': 'int', 'santri_mahad': 'bool', 'no_hp_unavailable': 'bool',
    'tinggal_luar_negeri': 'bool', 'sama_dengan_ayah': 'bool',
}

# Nilai pilihan GTK sesuai isian form detail GTK: kode tersimpan -> label tampilan.
# Saat impor, kode maupun label (tanpa beda huruf besar/kecil) diterima lalu disimpan sebagai kode.
PILIHAN_GTK: Dict[str, Dict[str, str]] = {
    'status_kepegawaian': {'pns': 'PNS', 'pppk': 'PPPK', 'non_asn': 'Non ASN'},
    'status_tempat_tinggal': {'milik_sendiri': 'Milik Sendiri', 'sewa': 'Sewa/Kontrak', 'menumpang': 'Menumpang', 'dinas': 'Rumah Dinas'},
    'status_perkawinan': {'belum_menikah': 'Belum Menikah', 'menikah': 'Menikah', 'duda': 'Duda', 'janda': 'Janda'},
    'gender': {'L': 'Laki-laki', 'P': 'Perempuan'},
    'jenis_ptk': {'guru_mapel': 'Guru Mata Pelajaran', 'guru_bk': 'Guru BK', 'kepala_sekolah': 'Kepala Sekolah',
                  'tenaga_administrasi': 'Tenaga Administrasi', 'pustakawan': 'Pustakawan', 'laboran': 'Laboran'},
    'agama': {'islam': 'Islam', 'kristen': 'Kristen', 'katolik': 'Katolik', 'hindu': 'Hindu', 'buddha': 'Buddha', 'konghucu': 'Konghucu'},
    'golongan_darah': {'A': 'A', 'B': 'B', 'AB': 'AB', 'O': 'O'},
    'transportasi': {'jalan_kaki': 'Jalan Kaki', 'sepeda': 'Sepeda', 'motor': 'Sepeda Motor', 'mobil': 'Mobil Pribadi',
                     'angkutan_umum': 'Angkutan Umum'},
}

# Pilihan isian Data Siswa (field student_details, per NAMA FIELD) — HARUS sama dengan pilihan form detail siswa
# (frontend/src/lib/pilihanDataMaster.js, dipakai StudentDetailDialog). Disimpan apa adanya (label).
PILIHAN_SISWA: Dict[str, List[str]] = {
    'citizenship': ['WNI', 'WNA'],
    'agama': ['Islam', 'Kristen Protestan', 'Katolik', 'Hindu', 'Buddha', 'Kong hu cu'],
    'cita_cita': ['PNS', 'TNI/Polri', 'Guru/Dosen', 'Dokter', 'Politikus', 'Wiraswasta', 'Seniman/Artis', 'Ilmuwan', 'Agamawan', 'Lainnya'],
    'hobi': ['Olahraga', 'Kesenian', 'Membaca', 'Menulis', 'Jalan-jalan', 'Lainnya'],
    'pembiaya_sekolah': ['Orang Tua', 'Wali/Orang Tua Asuh', 'Tanggungan Sendiri', 'Lainnya'],
    'pra_sekolah': ['Pernah TK/RA', 'Pernah PAUD'],
    'imunisasi': ['Hepatitis B', 'BCG', 'DPT', 'Polio', 'Campak', 'Covid'],
    'status': ['Masih Hidup', 'Sudah Meninggal', 'Tidak Diketahui'],
    'pendidikan': ['SD/Sederajat', 'SMP/Sederajat', 'SMA/Sederajat', 'D1', 'D2', 'D3', 'D4/S1', 'S2', 'S3', 'Tidak Sekolah', 'Lainnya'],
    'pekerjaan': ['Tidak Bekerja', 'Pensiunan', 'PNS', 'TNI/Polri', 'Guru/Dosen', 'Pegawai Swasta', 'Wiraswasta',
                  'Pengacara/Jaksa/Hakim/Notaris', 'Seniman/Pelukis/Artis/Sejenis', 'Dokter/Bidan/Perawat', 'Pilot/Pramugara',
                  'Pedagang', 'Petani/Peternak', 'Nelayan', 'Buruh (Tani/Pabrik/Bangunan)', 'Sopir/Masinis/Kondektur/Tukang Ojek',
                  'Politikus', 'Lainnya'],
    'penghasilan': ['Di bawah 800.000', '800.001-1.200.000', '1.200.001-1.800.000', '1.800.001-2.500.000', '2.500.001-3.500.000',
                    '3.500.001-4.800.000', '4.800.001-6.500.000', '6.500.001-10.000.000', '10.000.001-20.000.000', 'Lebih dari 20.000.000'],
    'hubungan_wali': ['Sama dengan ayah kandung', 'Sama dengan ibu kandung', 'Lainnya'],
    'status_wali': ['Sama dengan ayah kandung', 'Sama dengan ibu kandung', 'Lainnya'],
    'status_kepemilikan': ['Milik Sendiri', 'Rumah Orang Tua', 'Rumah Saudara/Kerabat', 'Rumah Dinas', 'Sewa/Kontrak', 'Lainnya'],
    'status_tempat_tinggal': ['Tinggal dengan Ayah Kandung', 'Tinggal dengan Ibu Kandung', 'Tinggal dengan Wali', 'Ikut Saudara/Kerabat',
                              'Asrama Madrasah', 'Kontrak/Kost', 'Tinggal di Asrama Pesantren', 'Panti Asuhan', 'Rumah Singgah', 'Lainnya'],
    'jarak_tempuh': ['Kurang dari 5 km', '5-10 km', '11-20 km', '21-30 km', 'Lebih dari 30 km'],
    'transportasi': ['Jalan Kaki', 'Sepeda', 'Sepeda Motor', 'Mobil Pribadi', 'Antar Jemput Sekolah', 'Angkutan Umum', 'Perahu/Sampan',
                     'Kendaraan Pribadi', 'Kereta Api', 'Ojek', 'Andong/Bendi/Sado/Dokar/Delman/Becak', 'Lainnya'],
    'waktu_tempuh': ['1-10 menit', '10-19 menit', '20-29 menit', '30-39 menit', '1-2 jam', 'Lebih dari 2 jam'],
}


def opsi_kolom(kolom: Dict[str, Any]) -> List[str]:
    """Isian yang tersedia di aplikasi untuk satu kolom (kosong = teks bebas).
    Field akun (user.*) memakai kode PILIHAN_GTK; field detail siswa (detail.*) memakai PILIHAN_SISWA."""
    path = kolom.get('path') or ''
    field = path.split('.')[-1]
    if path.startswith('user.') and field in PILIHAN_GTK:
        return list(PILIHAN_GTK[field])
    if path.startswith('detail.') and field in PILIHAN_SISWA:
        return list(PILIHAN_SISWA[field])
    return []


def label_opsi(kolom: Dict[str, Any]) -> List[str]:
    """Opsi untuk dibaca manusia: kode GTK beserta labelnya ('islam = Islam'), pilihan siswa apa adanya."""
    path = kolom.get('path') or ''
    field = path.split('.')[-1]
    if path.startswith('user.') and field in PILIHAN_GTK:
        return [kd if kd == lb else f'{kd} = {lb}' for kd, lb in PILIHAN_GTK[field].items()]
    return opsi_kolom(kolom)


# Variasi isian yang lazim di data pendaftaran/EMIS -> pilihan baku aplikasi (kunci = teks dinormalkan).
ALIAS_PILIHAN: Dict[str, Dict[str, str]] = {
    'agama': {'kristen': 'Kristen Protestan', 'protestan': 'Kristen Protestan', 'kristenprotestan': 'Kristen Protestan',
              'katholik': 'Katolik', 'katolikroma': 'Katolik', 'budha': 'Buddha', 'konghucu': 'Kong hu cu',
              'khonghucu': 'Kong hu cu', 'khonghuchu': 'Kong hu cu', 'konghuchu': 'Kong hu cu'},
    'citizenship': {'indonesia': 'WNI', 'warganeganindonesia': 'WNI', 'warganegaraindonesia': 'WNI',
                    'asing': 'WNA', 'warganegaraasing': 'WNA'},
    'status': {'hidup': 'Masih Hidup', 'masihada': 'Masih Hidup', 'meninggal': 'Sudah Meninggal', 'meninggaldunia': 'Sudah Meninggal',
               'almarhum': 'Sudah Meninggal', 'almarhumah': 'Sudah Meninggal', 'wafat': 'Sudah Meninggal',
               'tidaktahu': 'Tidak Diketahui', 'tidakdiketahui': 'Tidak Diketahui'},
    'pendidikan': {'sd': 'SD/Sederajat', 'mi': 'SD/Sederajat', 'sdmi': 'SD/Sederajat', 'sdsederajat': 'SD/Sederajat', 'paketa': 'SD/Sederajat',
                   'smp': 'SMP/Sederajat', 'mts': 'SMP/Sederajat', 'smpmts': 'SMP/Sederajat', 'sltp': 'SMP/Sederajat', 'paketb': 'SMP/Sederajat',
                   'sma': 'SMA/Sederajat', 'ma': 'SMA/Sederajat', 'smk': 'SMA/Sederajat', 'smamasmk': 'SMA/Sederajat', 'smama': 'SMA/Sederajat',
                   'slta': 'SMA/Sederajat', 'paketc': 'SMA/Sederajat', 'mak': 'SMA/Sederajat',
                   's1': 'D4/S1', 'd4': 'D4/S1', 'sarjana': 'D4/S1', 's1d4': 'D4/S1', 'd4s1': 'D4/S1', 'diploma4': 'D4/S1',
                   'diploma1': 'D1', 'diploma2': 'D2', 'diploma3': 'D3', 'magister': 'S2', 'doktor': 'S3',
                   'tidaksekolah': 'Tidak Sekolah', 'tidakbersekolah': 'Tidak Sekolah', 'tidaktamatsd': 'Tidak Sekolah', 'putussd': 'Tidak Sekolah'},
    'pekerjaan': {'tidakbekerja': 'Tidak Bekerja', 'ibumahtangga': 'Tidak Bekerja', 'iburumahtangga': 'Tidak Bekerja', 'irt': 'Tidak Bekerja',
                  'pensiun': 'Pensiunan', 'pnstnipolri': 'PNS', 'tni': 'TNI/Polri', 'polri': 'TNI/Polri', 'polisi': 'TNI/Polri',
                  'guru': 'Guru/Dosen', 'dosen': 'Guru/Dosen', 'karyawanswasta': 'Pegawai Swasta', 'pegawaiswasta': 'Pegawai Swasta',
                  'swasta': 'Pegawai Swasta', 'karyawan': 'Pegawai Swasta', 'wirausaha': 'Wiraswasta', 'wiraswasta': 'Wiraswasta',
                  'pengusaha': 'Wiraswasta', 'dagang': 'Pedagang', 'pedagangkecil': 'Pedagang', 'pedagangbesar': 'Pedagang',
                  'petani': 'Petani/Peternak', 'peternak': 'Petani/Peternak', 'buruh': 'Buruh (Tani/Pabrik/Bangunan)',
                  'buruhtani': 'Buruh (Tani/Pabrik/Bangunan)', 'buruhpabrik': 'Buruh (Tani/Pabrik/Bangunan)',
                  'buruhbangunan': 'Buruh (Tani/Pabrik/Bangunan)', 'sopir': 'Sopir/Masinis/Kondektur/Tukang Ojek',
                  'supir': 'Sopir/Masinis/Kondektur/Tukang Ojek', 'ojek': 'Sopir/Masinis/Kondektur/Tukang Ojek',
                  'tukangojek': 'Sopir/Masinis/Kondektur/Tukang Ojek', 'dokter': 'Dokter/Bidan/Perawat', 'bidan': 'Dokter/Bidan/Perawat',
                  'perawat': 'Dokter/Bidan/Perawat', 'tenagakesehatan': 'Dokter/Bidan/Perawat', 'nelayanperikanan': 'Nelayan',
                  'seniman': 'Seniman/Pelukis/Artis/Sejenis', 'artis': 'Seniman/Pelukis/Artis/Sejenis'},
    'pembiaya_sekolah': {'orangtua': 'Orang Tua', 'ortu': 'Orang Tua', 'wali': 'Wali/Orang Tua Asuh', 'orangtuaasuh': 'Wali/Orang Tua Asuh',
                         'sendiri': 'Tanggungan Sendiri'},
    'pra_sekolah': {'tk': 'Pernah TK/RA', 'ra': 'Pernah TK/RA', 'tkra': 'Pernah TK/RA', 'paud': 'Pernah PAUD', 'kb': 'Pernah PAUD'},
    'imunisasi': {'hepatitis': 'Hepatitis B', 'hepb': 'Hepatitis B', 'hb': 'Hepatitis B', 'covid19': 'Covid', 'vaksincovid': 'Covid'},
    'hubungan_wali': {'ayah': 'Sama dengan ayah kandung', 'ayahkandung': 'Sama dengan ayah kandung',
                      'ibu': 'Sama dengan ibu kandung', 'ibukandung': 'Sama dengan ibu kandung'},
    'status_wali': {'ayah': 'Sama dengan ayah kandung', 'ayahkandung': 'Sama dengan ayah kandung',
                    'ibu': 'Sama dengan ibu kandung', 'ibukandung': 'Sama dengan ibu kandung'},
    'status_kepemilikan': {'milikpribadi': 'Milik Sendiri', 'sendiri': 'Milik Sendiri', 'milikorangtua': 'Rumah Orang Tua',
                           'orangtua': 'Rumah Orang Tua', 'rumahortu': 'Rumah Orang Tua', 'saudara': 'Rumah Saudara/Kerabat',
                           'kerabat': 'Rumah Saudara/Kerabat', 'dinas': 'Rumah Dinas', 'sewa': 'Sewa/Kontrak', 'kontrak': 'Sewa/Kontrak',
                           'kost': 'Sewa/Kontrak', 'kos': 'Sewa/Kontrak'},
    'status_tempat_tinggal': {'denganayah': 'Tinggal dengan Ayah Kandung',
                              'bersamaayah': 'Tinggal dengan Ayah Kandung', 'ayahkandung': 'Tinggal dengan Ayah Kandung',
                              'denganibu': 'Tinggal dengan Ibu Kandung', 'bersamaibu': 'Tinggal dengan Ibu Kandung',
                              'ibukandung': 'Tinggal dengan Ibu Kandung', 'denganwali': 'Tinggal dengan Wali', 'bersamawali': 'Tinggal dengan Wali',
                              'wali': 'Tinggal dengan Wali', 'saudara': 'Ikut Saudara/Kerabat', 'kerabat': 'Ikut Saudara/Kerabat',
                              'asrama': 'Asrama Madrasah', 'pesantren': 'Tinggal di Asrama Pesantren', 'pondokpesantren': 'Tinggal di Asrama Pesantren',
                              'pondok': 'Tinggal di Asrama Pesantren', 'kost': 'Kontrak/Kost', 'kos': 'Kontrak/Kost', 'kontrak': 'Kontrak/Kost',
                              'panti': 'Panti Asuhan'},
    # GTK (label form detail GTK)
    'status_perkawinan': {'kawin': 'Menikah', 'sudahmenikah': 'Menikah', 'sudahkawin': 'Menikah', 'belumkawin': 'Belum Menikah',
                          'lajang': 'Belum Menikah'},
    'transportasi': {'jalan': 'Jalan Kaki', 'motor': 'Sepeda Motor', 'sepedamotorpribadi': 'Sepeda Motor', 'mobil': 'Mobil Pribadi',
                     'mobilpribadi': 'Mobil Pribadi', 'antarjemput': 'Antar Jemput Sekolah', 'jemputan': 'Antar Jemput Sekolah',
                     'angkot': 'Angkutan Umum', 'angkutanumumbusbus': 'Angkutan Umum', 'bus': 'Angkutan Umum', 'ojekonline': 'Ojek',
                     'perahu': 'Perahu/Sampan', 'sampan': 'Perahu/Sampan', 'becak': 'Andong/Bendi/Sado/Dokar/Delman/Becak',
                     'delman': 'Andong/Bendi/Sado/Dokar/Delman/Becak', 'keretaapi': 'Kereta Api', 'kereta': 'Kereta Api'},
}

# Batas bawah (rupiah) tiap pilihan penghasilan aplikasi, urut naik; batas atas = batas bawah berikutnya - 1.
_BATAS_PENGHASILAN = [(0, 'Di bawah 800.000'), (800001, '800.001-1.200.000'), (1200001, '1.200.001-1.800.000'),
                      (1800001, '1.800.001-2.500.000'), (2500001, '2.500.001-3.500.000'), (3500001, '3.500.001-4.800.000'),
                      (4800001, '4.800.001-6.500.000'), (6500001, '6.500.001-10.000.000'), (10000001, '10.000.001-20.000.000'),
                      (20000001, 'Lebih dari 20.000.000')]


def _kunci_pilihan(teks: str) -> str:
    return re.sub(r'[^a-z0-9]', '', str(teks or '').lower())


def _kelompok_penghasilan(teks: str) -> Optional[str]:
    """Penghasilan berupa angka/rentang rupiah -> pilihan aplikasi bila seluruh rentang masuk satu pilihan."""
    t = teks.lower().replace('rp', ' ')
    if re.search(r'\b(juta|jt)\b', t):
        angka = [float(x.replace(',', '.')) * 1_000_000 for x in re.findall(r'(\d+(?:[.,]\d+)?)\s*(?:juta|jt)', t)]
    else:
        angka = [int(re.sub(r'\D', '', x)) for x in re.findall(r'\d[\d.,]*', t) if re.sub(r'\D', '', x)]
    if not angka:
        if re.search(r'tidak\s+(ada|berpenghasilan)', t):
            return 'Di bawah 800.000'
        return None
    if re.search(r'kurang|di ?bawah|<', t) and len(angka) == 1:
        rendah, tinggi = 0, angka[0] - 1
    elif re.search(r'lebih|di ?atas|>', t) and len(angka) == 1:
        rendah, tinggi = angka[0] + 1, float('inf')
    else:
        rendah, tinggi = min(angka[:2]), max(angka[:2])

    def kelompok(n):
        hasil = _BATAS_PENGHASILAN[0][1]
        for batas, nama in _BATAS_PENGHASILAN:
            if n >= batas:
                hasil = nama
        return hasil

    a, b = kelompok(rendah), kelompok(min(tinggi, 10 ** 12))
    return a if a == b else None


def kenali_pilihan(field: str, teks: str, opsi: List[str]) -> Optional[str]:
    """Kenali isian bebas sebagai salah satu `opsi`: sama persis (tanpa beda huruf/spasi/tanda baca), alias lazim
    data pendaftaran, rentang penghasilan, lalu kemiripan tinggi yang tidak ambigu. None bila tidak dikenali."""
    from difflib import SequenceMatcher
    kunci = _kunci_pilihan(teks)
    if not kunci:
        return None
    for o in opsi:
        if _kunci_pilihan(o) == kunci:
            return o
    alias = ALIAS_PILIHAN.get(field, {}).get(kunci)
    if alias in opsi:
        return alias
    if field == 'penghasilan':
        k = _kelompok_penghasilan(teks)
        if k in opsi:
            return k
    # isian memuat persis satu pilihan (mis. "Tinggal bersama Ayah Kandung")
    for k_alias, tujuan in ALIAS_PILIHAN.get(field, {}).items():
        if len(k_alias) >= 4 and k_alias in kunci and tujuan in opsi:
            calon = {t for ka, t in ALIAS_PILIHAN[field].items() if len(ka) >= 4 and ka in kunci}
            if len(calon) == 1:
                return tujuan
    skor = sorted(((SequenceMatcher(None, kunci, _kunci_pilihan(o)).ratio(), o) for o in opsi), reverse=True)
    if skor and skor[0][0] >= 0.85 and (len(skor) == 1 or skor[0][0] - skor[1][0] >= 0.08):
        return skor[0][1]
    return None


def normalisasi_pilihan_siswa(field: str, teks: str) -> str:
    """Isian pilihan siswa -> ejaan baku di aplikasi (dikenali dari variasi penulisan, lihat kenali_pilihan);
    ValueError berisi daftar pilihan bila tidak dikenali."""
    hasil = kenali_pilihan(field, teks, PILIHAN_SISWA[field])
    if hasil:
        return hasil
    raise ValueError('pilihan yang tersedia: ' + ' / '.join(PILIHAN_SISWA[field]))


def normalisasi_pilihan(field: str, teks: str) -> str:
    """Ubah isian pilihan (kode atau label) menjadi kode tersimpan; ValueError bila tidak dikenal."""
    opsi = PILIHAN_GTK[field]
    t = teks.strip().lower()
    for kode, label in opsi.items():
        if t in (kode.lower(), label.lower(), kode.lower().replace('_', ' ')):
            return kode
    label_ke_kode = {label: kode for kode, label in opsi.items()}
    dikenali = kenali_pilihan(field, teks, list(label_ke_kode)) or kenali_pilihan(field, teks, list(opsi))
    if dikenali:
        return label_ke_kode.get(dikenali, dikenali)
    raise ValueError(f"isi salah satu: {', '.join(opsi)}")


SKEMA = {
    'siswa': {'sheet': 'Data Siswa', 'kolom': KOLOM_DATA_SISWA},
    'gtk': {'sheet': 'Data GTK', 'kolom': KOLOM_DATA_GTK},
}


def validasi_skema(kolom: List[Dict[str, Any]]) -> None:
    """Pastikan key & judul unik, path dikenal, dan kolom penanda berada di depan."""
    keys = [k['key'] for k in kolom]
    labels = [k['label'] for k in kolom]
    if len(set(keys)) != len(keys) or len(set(labels)) != len(labels):
        raise ValueError('key atau judul kolom ganda')
    if not all(k['path'] in ('kelas', 'jenis') or k['path'].split('.')[0] in ('user', 'detail') for k in kolom):
        raise ValueError('path kolom tidak dikenal')
    kunci = [i for i, k in enumerate(kolom) if k.get('kunci')]
    if not kunci or kunci != list(range(len(kunci))):
        raise ValueError('kolom penanda harus di awal')


for _s in SKEMA.values():
    validasi_skema(_s['kolom'])


def cek_header(header: List[Any], kolom: List[Dict[str, Any]]) -> List[str]:
    """Bandingkan baris judul berkas dengan skema. Kosong = cocok persis; selain itu
    daftar masalah: kolom hilang, kolom asing, atau urutan berbeda."""
    got = [str(h).strip() if h is not None else '' for h in header]
    while got and got[-1] == '':
        got.pop()
    want = [k['label'] for k in kolom]
    if got == want:
        return []
    masalah = []
    hilang = [x for x in want if x not in got]
    asing = [x for x in got if x and x not in want]
    if hilang:
        masalah.append('Kolom tidak ditemukan: ' + ', '.join(hilang[:10]) + (' dst.' if len(hilang) > 10 else ''))
    if asing:
        masalah.append('Kolom tidak dikenal: ' + ', '.join(asing[:10]) + (' dst.' if len(asing) > 10 else ''))
    if not hilang and not asing:
        masalah.append('Urutan kolom berbeda dari template; gunakan template atau unduhan terbaru')
    return masalah


def parse_nilai(kolom: Dict[str, Any], v: Any) -> Any:
    """Kebalikan nilai_sel: teks sel Excel -> nilai tersimpan (None bila kosong).
    Melempar ValueError dengan pesan berbahasa Indonesia bila isi tidak valid."""
    teks = nilai_sel(v).strip()
    if teks == '':
        return None
    if kolom.get('daftar'):
        return [x.strip() for x in teks.split(';') if x.strip()]
    tipe = TIPE_KOLOM.get(kolom['path'].split('.')[-1], 'teks')
    if tipe == 'bool':
        t = teks.lower()
        if t in ('ya', 'y', 'true', '1'):
            return True
        if t in ('tidak', 'false', '0'):
            return False
        raise ValueError(f"{kolom['label']}: isi 'Ya' atau 'Tidak'")
    if tipe == 'int':
        try:
            return int(float(teks))
        except ValueError:
            raise ValueError(f"{kolom['label']}: harus berupa angka")
    return teks


def _ambil(obj: Optional[Dict], parts: List[str]) -> Any:
    for p in parts:
        if not isinstance(obj, dict):
            return None
        obj = obj.get(p)
    return obj


def nilai_sel(v: Any) -> str:
    """Nilai sel Excel sebagai teks (agar NISN/NIK/no HP tidak kehilangan nol di depan)."""
    if v is None:
        return ''
    if isinstance(v, bool):
        return 'Ya' if v else 'Tidak'
    if isinstance(v, (list, tuple)):
        return '; '.join(str(x) for x in v if x not in (None, ''))
    if hasattr(v, 'isoformat'):
        return v.isoformat()[:10]
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


# Field lama yang dipakai bila field baku kosong (impor GTK lama menyimpan NUPTK di nip_nuptk).
FALLBACK_GTK = {'nuptk': 'nip_nuptk'}
LABEL_JENIS_GTK = {'guru': 'Guru', 'tendik': 'Tenaga Kependidikan'}


def nilai_kolom_gtk(kolom: Dict, user: Dict, jenis: Optional[str]) -> str:
    if kolom['path'] == 'jenis':
        return LABEL_JENIS_GTK.get(jenis or '', '')
    field = kolom['path'].split('.', 1)[1]
    v = user.get(field)
    if v in (None, '') and field in FALLBACK_GTK:
        v = user.get(FALLBACK_GTK[field])
    return nilai_sel(v)


def nilai_kolom_siswa(kolom: Dict, user: Dict, detail: Optional[Dict], kelas: Optional[str]) -> str:
    sumber, *rest = kolom['path'].split('.')
    if sumber == 'kelas':
        return nilai_sel(kelas)
    return nilai_sel(_ambil(user if sumber == 'user' else (detail or {}), rest))


# ------------------------------------------------------------
# PEMBUAT WORKBOOK
# ------------------------------------------------------------
_HEADER_FILL = PatternFill(start_color='006837', end_color='006837', fill_type='solid')
_KUNCI_FILL = PatternFill(start_color='B45309', end_color='B45309', fill_type='solid')
_BORDER = Border(left=Side(style='thin', color='CBD5E1'), right=Side(style='thin', color='CBD5E1'),
                 top=Side(style='thin', color='CBD5E1'), bottom=Side(style='thin', color='CBD5E1'))


def buat_workbook(sheet: str, kolom: List[Dict], rows: List[List[str]]) -> Workbook:
    """Workbook satu sheet: baris 1 judul kolom (kolom penanda berwarna oranye),
    baris berikutnya data. Semua sel berformat teks."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    for c, k in enumerate(kolom, 1):
        cell = ws.cell(row=1, column=c, value=k['label'])
        cell.fill = _KUNCI_FILL if k.get('kunci') else _HEADER_FILL
        cell.font = Font(color='FFFFFF', bold=True)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = _BORDER
        ws.column_dimensions[get_column_letter(c)].width = max(12, min(40, len(k['label']) + 4))
    for r, row in enumerate(rows, 2):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.number_format = '@'
            cell.border = _BORDER
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = 'C2'
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(kolom))}{len(rows) + 1}"
    kunci_kolom_identitas(ws, kolom, terisi=bool(rows))
    return wb


TERBUKA = Protection(locked=False)
TERKUNCI = Protection(locked=True)


def kunci_kolom_identitas(ws, kolom: List[Dict], terisi: bool) -> None:
    """Proteksi sheet (tanpa kata sandi — mencegah salah ubah, bukan pengamanan):
    - baris judul selalu terkunci agar susunan kolom tidak berubah;
    - berkas berisi data (unduhan): sel kolom identitas (ID, NISN/NIP) terkunci karena
      dipakai mencocokkan baris saat impor; kolom lain bebas diisi;
    - template kosong: kolom identitas tetap bisa diisi (penanda baris baru diketik pengguna).
    Menambah/menghapus baris, menyortir, memfilter, dan memformat tetap diizinkan."""
    for c, k in enumerate(kolom, 1):
        huruf = get_column_letter(c)
        kunci = terisi and k.get('kunci')
        # Gaya kolom berlaku untuk sel kosong di bawah data (baris baru yang ditambahkan pengguna).
        ws.column_dimensions[huruf].protection = TERBUKA
        ws.column_dimensions[huruf].number_format = '@'
        for (cell,) in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=c, max_col=c):
            cell.protection = TERKUNCI if kunci else TERBUKA
        ws.cell(row=1, column=c).protection = TERKUNCI
    pr = ws.protection
    pr.sheet = True
    for opsi in ('formatCells', 'formatColumns', 'formatRows', 'insertRows', 'deleteRows', 'sort', 'autoFilter'):
        setattr(pr, opsi, False)  # False = tindakan ini TIDAK diproteksi (diizinkan)


def tambah_keterangan_kosong(wb: Workbook, pesan: List[str]) -> None:
    """Tambahkan sheet 'Keterangan' berisi alasan berkas tidak memuat baris data.
    Sheet data tetap berisi baris judul agar bisa langsung dipakai sebagai acuan isian."""
    ws = wb.create_sheet('Keterangan')
    ws['A1'] = 'Tidak ada data untuk diunduh'
    ws['A1'].font = Font(bold=True, size=12, color='B45309')
    for i, p in enumerate(pesan, 3):
        ws.cell(row=i, column=1, value=p)
    ws.column_dimensions['A'].width = 100


# Penanda di sel pertama baris keterangan pengisian (baris 2 template). Impor melewati
# baris yang diawali penanda ini sehingga baris keterangan tidak pernah dianggap data.
PENANDA_KETERANGAN = '[Keterangan]'


def adalah_baris_keterangan(row) -> bool:
    if not row:
        return False
    pertama = row[0]
    return isinstance(pertama, str) and pertama.strip().startswith(PENANDA_KETERANGAN)


def buat_berkas(jenis: str, rows: List[List[str]]) -> Workbook:
    """Satu-satunya pembuat berkas data master: dipakai unduhan (rows berisi data) dan
    template impor (rows kosong) sehingga nama sheet, urutan, dan judul kolom selalu identik
    dengan SKEMA[jenis] — yang juga dipakai cek_header saat impor."""
    if jenis not in SKEMA:
        raise ValueError(f'jenis berkas tidak dikenal: {jenis}')
    sk = SKEMA[jenis]
    lebar = len(sk['kolom'])
    if any(len(r) != lebar for r in rows):
        raise ValueError('jumlah sel baris tidak sama dengan jumlah kolom skema')
    return buat_workbook(sk['sheet'], sk['kolom'], rows)


def buat_template(jenis: str) -> Workbook:
    """Template impor kosong: sheet data hanya berisi baris judul skema yang sama dengan unduhan."""
    return buat_berkas(jenis, [])


def baca_header(wb: Workbook, jenis: str) -> List[Any]:
    """Baris judul dari sheet data berkas (sheet bernama sesuai skema, atau sheet pertama)."""
    nama = SKEMA[jenis]['sheet']
    ws = wb[nama] if nama in wb.sheetnames else wb.worksheets[0]
    return [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]


def workbook_bytes(wb: Workbook) -> bytes:
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
