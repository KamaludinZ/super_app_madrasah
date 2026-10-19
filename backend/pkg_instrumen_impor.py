"""Impor instrumen PKG dari berkas Excel berformat template resmi (pkg_template_v2026.xlsx).

Struktur yang dibaca dari tiap sheet kompetensi (pkg_template.SHEET_KOMPETENSI, berurutan kode 1–4):
    kolom B  "N. NAMA KOMPETENSI"            -> kompetensi (baris pertama yang cocok)
    kolom B  "N.M Nama sub-kompetensi" / "N.M. ..." -> sub-kompetensi
    kolom C  "L.M" + kolom D teks            -> indikator Level L untuk sub-kompetensi ke-M dalam kompetensi
Tanda centang contoh di kolom B (huruf "v") diabaikan. Bobot tiap sub = 1 (template tidak memuat bobot; rumus resmi
menghitung skor 0–5 per sub). Hasil divalidasi dengan pkg_skema.StrukturInstrumen.

Berkas hanya dibaca dari bytes di memori (pola ckg_impor/excel_io); berkas sumber tidak pernah ditulis.
Kegagalan dilaporkan sebagai GalatImporInstrumen dengan daftar {lokasi ('sheet!sel'), pesan}.
"""
import hashlib
import io
import re
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import load_workbook
from pydantic import ValidationError

from pkg_skema import StrukturInstrumen
from pkg_template import SHEET_KOMPETENSI, TEMPLATE_PKG, baca_bytes_template

BATAS_UKURAN_BERKAS = 10 * 1024 * 1024
POLA_KOMPETENSI = re.compile(r'^(\d{1,2})\.\s+(.+)$')
POLA_SUB = re.compile(r'^(\d{1,2})\.(\d{1,2})\.?\s+(.+)$')
POLA_INDIKATOR = re.compile(r'^(\d)\.(\d{1,2})$')
BARIS_MAKS = 200


class GalatImporInstrumen(Exception):
    def __init__(self, pesan: str, galat: Optional[List[Dict[str, str]]] = None):
        super().__init__(pesan)
        self.pesan = pesan
        self.galat = galat or []


def _teks(v: Any) -> str:
    return re.sub(r'\s+', ' ', str(v)).strip() if v is not None else ''


def _nama_kompetensi(nama: str) -> str:
    """'PEDAGOGIK' / 'KOMPETENSI KEPRIBADIAN' -> 'Kompetensi Pedagogik' / 'Kompetensi Kepribadian'."""
    nama = re.sub(r'(?i)^kompetensi\s+', '', _teks(nama)).strip()
    return f'Kompetensi {nama.title()}' if nama else ''


def _baca_sheet(ws, nama_sheet: str, kode_harapan: str, galat: List[Dict[str, str]]) -> Optional[Dict[str, Any]]:
    kompetensi: Optional[Dict[str, Any]] = None
    subs: List[Dict[str, Any]] = []
    for nomor, (b, c, d) in enumerate(ws.iter_rows(min_row=1, max_row=BARIS_MAKS, min_col=2, max_col=4, values_only=True), 1):
        b, c, d = _teks(b), _teks(c), _teks(d)
        if kompetensi is None:
            m = POLA_KOMPETENSI.match(b)
            if m and not POLA_SUB.match(b):
                if m.group(1) != kode_harapan:
                    galat.append({'lokasi': f'{nama_sheet}!B{nomor}',
                                  'pesan': f'Kode kompetensi {m.group(1)} tidak sesuai urutan sheet (seharusnya {kode_harapan}).'})
                kompetensi = {'kode': kode_harapan, 'nama': _nama_kompetensi(m.group(2)), 'sub': []}
            continue
        m = POLA_SUB.match(b)
        if m:
            if m.group(1) != kode_harapan:
                galat.append({'lokasi': f'{nama_sheet}!B{nomor}', 'pesan': f'Sub-kompetensi {m.group(1)}.{m.group(2)} '
                              f'berada di sheet kompetensi {kode_harapan}.'})
            subs.append({'kode': f'{kode_harapan}.{int(m.group(2))}', 'nama': m.group(3).strip().rstrip('.'),
                         'bobot': 1, 'indikator': [], '_baris': nomor})
            continue
        m = POLA_INDIKATOR.match(c)
        if m and subs:
            sub = subs[-1]
            level, urutan = int(m.group(1)), int(m.group(2))
            if f'{kode_harapan}.{urutan}' != sub['kode']:
                galat.append({'lokasi': f'{nama_sheet}!C{nomor}',
                              'pesan': f'Kode indikator {c} tidak sesuai sub-kompetensi {sub["kode"]}.'})
                continue
            if not d:
                galat.append({'lokasi': f'{nama_sheet}!D{nomor}', 'pesan': f'Teks indikator Level {level} sub-kompetensi {sub["kode"]} kosong.'})
                continue
            sub['indikator'].append({'level': level, 'teks': d.rstrip(';').strip()})
    if kompetensi is None:
        galat.append({'lokasi': nama_sheet, 'pesan': 'Judul kompetensi (mis. "1. PEDAGOGIK") tidak ditemukan di kolom B.'})
        return None
    if not subs:
        galat.append({'lokasi': nama_sheet, 'pesan': 'Tidak ada sub-kompetensi yang terbaca.'})
    for sub in subs:
        jumlah = len(sub['indikator'])
        if jumlah != 5:
            galat.append({'lokasi': f"{nama_sheet}!B{sub['_baris']}",
                          'pesan': f"Sub-kompetensi {sub['kode']} memuat {jumlah} indikator; seharusnya 5 (Level 1–5)."})
        sub.pop('_baris')
    kompetensi['sub'] = subs
    return kompetensi


def _galat_validasi(e: ValidationError) -> List[Dict[str, str]]:
    hasil = []
    for g in e.errors():
        jalur = '.'.join(str(x) for x in g.get('loc', ()))
        hasil.append({'lokasi': jalur or 'struktur', 'pesan': g.get('msg', 'tidak valid').replace('Value error, ', '')})
    return hasil


def baca_instrumen_excel(isi: bytes) -> Tuple[StrukturInstrumen, List[str]]:
    """Baca & validasi struktur instrumen dari bytes .xlsx. Hasil: (struktur, peringatan)."""
    if not isi:
        raise GalatImporInstrumen('Berkas kosong.')
    if len(isi) > BATAS_UKURAN_BERKAS:
        raise GalatImporInstrumen('Ukuran berkas melebihi 10 MB.')
    try:
        wb = load_workbook(io.BytesIO(isi), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        raise GalatImporInstrumen('Berkas bukan Excel .xlsx yang valid.') from e
    galat: List[Dict[str, str]] = []
    kompetensi = []
    try:
        for urutan, nama_sheet in enumerate(SHEET_KOMPETENSI, 1):
            if nama_sheet not in wb.sheetnames:
                galat.append({'lokasi': nama_sheet, 'pesan': f'Sheet "{nama_sheet}" tidak ditemukan.'})
                continue
            k = _baca_sheet(wb[nama_sheet], nama_sheet, str(urutan), galat)
            if k:
                kompetensi.append(k)
    finally:
        wb.close()
    if galat:
        raise GalatImporInstrumen('Berkas tidak sesuai struktur instrumen PKG.', galat)
    try:
        struktur = StrukturInstrumen(kompetensi=kompetensi)
    except ValidationError as e:
        raise GalatImporInstrumen('Isi instrumen tidak valid.', _galat_validasi(e)) from e
    return struktur, []


def pratinjau_dari_bytes(isi: bytes, sumber: str, nama_berkas: str) -> Dict[str, Any]:
    """Baca berkas lalu kembalikan pratinjau impor: struktur, checksum, kecocokan dengan template resmi, peringatan."""
    struktur, peringatan = baca_instrumen_excel(isi)
    checksum = hashlib.sha256(isi).hexdigest()
    resmi = TEMPLATE_PKG['instrumen']
    cocok = checksum == resmi['checksum_resmi']
    if not cocok:
        peringatan = [*peringatan, 'Berkas bukan salinan identik template resmi; pastikan isinya sudah disahkan.']
    return {
        'sumber': sumber, 'nama_berkas': nama_berkas, 'versi_template': resmi['versi'], 'checksum': checksum,
        'cocok_resmi': cocok, 'kompetensi': struktur.model_dump()['kompetensi'], 'peringatan': peringatan,
    }


def pratinjau_dari_template() -> Dict[str, Any]:
    """Pratinjau impor dari berkas template resmi di folder internal (tanpa unggah)."""
    return pratinjau_dari_bytes(baca_bytes_template('instrumen'), 'template', TEMPLATE_PKG['instrumen']['nama_berkas'])
