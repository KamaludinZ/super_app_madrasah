"""Cetak dokumen resmi PKG dari template Excel (salinan di memori; berkas master tidak pernah ditulis).

Rekap hasil (rekap_hasil_pkg_v2026.xlsx, sheet 'Rekap hasil PKG'), posisi sel dari pkg_template.PETA_CETAK['rekap']:
    A1/A2   judul & sub-judul (tahun, semester, tahun pelajaran periode)
    baris 6 dst.  No | Nama Guru | Pangkat/Gol | Mapel | Tugas tambahan | NKG (capaian %) | Predikat (rumus) | Rekomendasi
    Template menyediakan 10 baris; baris disisipkan bila guru lebih banyak, baris sisa dikosongkan.
    Jumlah & rata-rata ditulis ulang dengan rumus yang benar (rumus asli template menjumlah sel label & membagi 10).
    G19 tempat & tanggal, G24 nama kepala madrasah.
Predikat rekap memakai rumus template (Sangat Baik >= 91, Baik >= 76, Cukup >= 61, Sedang >= 60, Kurang).
"""
import re
from copy import copy
from datetime import date
from typing import Any, Dict, List, Optional

from openpyxl.styles import Alignment
from openpyxl.worksheet.worksheet import Worksheet

from pkg_ambang import AMBANG_PREDIKAT_REKAP, terapkan
from pkg_template import PETA_CETAK, isi_template_ke_bytes, tampilkan_hanya

BULAN = ('Januari', 'Februari', 'Maret', 'April', 'Mei', 'Juni', 'Juli', 'Agustus', 'September', 'Oktober', 'November',
         'Desember')


def tanggal_indonesia(tgl: date) -> str:
    return f'{tgl.day} {BULAN[tgl.month - 1]} {tgl.year}'


def predikat_rekap(nkg: Optional[float]) -> str:
    """Predikat skala rekap (rumus kolom G template rekap); dipakai juga oleh PDF cadangan."""
    return '' if nkg is None else terapkan(AMBANG_PREDIKAT_REKAP, nkg, kosong_bila_nol=False) or ''


def _rumus_predikat(sel: str) -> str:
    return (f'=IF({sel}="","",IF({sel}>=91,"Sangat Baik",IF({sel}>=76,"Baik",IF({sel}>=61,"Cukup",'
            f'IF({sel}>=60,"Sedang","Kurang")))))')


def _salin_gaya_baris(ws: Worksheet, sumber: int, tujuan: int, kolom_maks: int) -> None:
    for c in range(1, kolom_maks + 1):
        a, b = ws.cell(sumber, c), ws.cell(tujuan, c)
        if a.has_style:
            b._style = copy(a._style)
    ws.row_dimensions[tujuan].height = ws.row_dimensions[sumber].height


def isi_rekap(ws: Worksheet, guru: List[Dict[str, Any]], periode: Dict[str, Any], kepala: Optional[str] = None,
              tanggal_cetak: Optional[date] = None) -> None:
    """Isi sheet rekap dengan baris guru (terurut seperti diberikan). `guru` = baris rekap terkunci:
    {nama, pangkat?, mapel, tugas_tambahan?, capaian_persen, rekomendasi?}."""
    peta = PETA_CETAK['rekap']
    k = peta['kolom']
    awal, tersedia = peta['baris_awal'], peta['baris_tersedia']
    kolom_maks = ws.max_column
    tambahan = max(0, len(guru) - tersedia)

    # Sisipkan baris sebelum baris Jumlah; sel gabungan di bawahnya dipindah manual (openpyxl tidak menggesernya).
    if tambahan:
        gabung_bawah = [r for r in list(ws.merged_cells.ranges) if r.min_row >= peta['baris_jumlah']]
        for r in gabung_bawah:
            ws.unmerge_cells(str(r))
        ws.insert_rows(peta['baris_jumlah'], tambahan)
        for r in gabung_bawah:
            r.shift(0, tambahan)
            ws.merge_cells(str(r))
        for i in range(tambahan):
            _salin_gaya_baris(ws, awal + tersedia - 1, awal + tersedia + i, kolom_maks)
    akhir = awal + max(len(guru), tersedia) - 1
    geser = tambahan

    # Judul: tahun & semester/tahun pelajaran periode.
    tahun = str((periode.get('tanggal_tutup') or periode.get('tanggal_buka') or '')[:4] or '')
    if tahun and isinstance(ws[peta['judul']].value, str):
        ws[peta['judul']] = re.sub(r'TAHUN \d{4}', f'TAHUN {tahun}', ws[peta['judul']].value)
    semester = (periode.get('semester') or '').upper()
    tp = periode.get('tahun_pelajaran') or ''
    ws[peta['subjudul']] = f'PADA SEMESTER {semester} TAHUN PELAJARAN {tp}'.strip()

    for i in range(awal, akhir + 1):
        g = guru[i - awal] if i - awal < len(guru) else None
        nilai = {
            'no': (i - awal + 1) if g else None,
            'nama': g.get('nama') if g else None,
            'pangkat': g.get('pangkat') if g else None,
            'mapel': g.get('mapel') if g else None,
            'tugas_tambahan': g.get('tugas_tambahan') if g else None,
            'nkg': round(float(g['capaian_persen']), 2) if g and g.get('capaian_persen') is not None else None,
            'predikat': _rumus_predikat(f"{k['nkg']}{i}") if g else None,
            'rekomendasi': g.get('rekomendasi') if g else None,
        }
        for kunci, kol in k.items():
            ws[f'{kol}{i}'] = nilai[kunci]
        if nilai['rekomendasi']:
            sel = ws[f"{k['rekomendasi']}{i}"]
            sel.alignment = Alignment(horizontal=sel.alignment.horizontal, vertical='center', wrap_text=True)
        ws[f"{k['nkg']}{i}"].number_format = '0.00'

    rentang = f"{k['nkg']}{awal}:{k['nkg']}{akhir}"
    baris_jumlah, baris_rata = peta['baris_jumlah'] + geser, peta['baris_rata'] + geser
    ws[f"{k['nkg']}{baris_jumlah}"] = f'=SUM({rentang})'
    ws[f"{k['nkg']}{baris_rata}"] = f'=IF(COUNT({rentang})=0,"",ROUND(AVERAGE({rentang}),2))'
    ws[f"{k['nkg']}{baris_rata}"].number_format = '0.00'
    ws[f"{k['predikat']}{baris_rata}"] = _rumus_predikat(f"{k['nkg']}{baris_rata}")

    sel_tanggal = _geser_sel(peta['tempat_tanggal'], geser)
    teks = str(ws[sel_tanggal].value or '')
    kota = teks.split(',')[0].strip() if ',' in teks else ''
    tgl = tanggal_indonesia(tanggal_cetak or date.today())
    ws[sel_tanggal] = f'{kota}, {tgl}' if kota else tgl
    if kepala:
        ws[_geser_sel(peta['nama_kepala'], geser)] = kepala


def _geser_sel(sel: str, baris: int) -> str:
    m = re.match(r'^([A-Z]+)(\d+)$', sel)
    return f'{m.group(1)}{int(m.group(2)) + baris}'


def rekap_excel(guru: List[Dict[str, Any]], periode: Dict[str, Any], kepala: Optional[str] = None,
                tanggal_cetak: Optional[date] = None) -> bytes:
    """Berkas .xlsx rekap hasil PKG dari template resmi (template harus utuh sesuai checksum resmi)."""
    sheet = PETA_CETAK['rekap']['sheet']
    return isi_template_ke_bytes('rekap', lambda wb: isi_rekap(wb[sheet], guru, periode, kepala, tanggal_cetak))


# ---------------------------------------------------------------- laporan per guru (template instrumen)

class StrukturTidakSesuai(Exception):
    """Instrumen lembar tidak dapat dipetakan ke format cetak template resmi (4 kompetensi x 3 sub x 5 level)."""


def _tanggal_teks(nilai: Optional[str]) -> str:
    try:
        return tanggal_indonesia(date.fromisoformat(str(nilai)[:10]))
    except (TypeError, ValueError):
        return '-'


def isi_laporan_guru(wb, laporan: Dict[str, Any]) -> None:
    """Isi workbook template instrumen untuk satu guru: identitas di '00-DATA' & centang 'v' indikator di KOMP-n
    (rumus resmi 04-RKAP-PKG menghitung tingkat, jumlah, capaian, dan predikat).
    `laporan` = {guru:{nama,nip}, periode, sk:{nomor_sk,tanggal_sk}, penilai:{nama,nip}, kepala:{nama,nip},
                 nama_sekolah, instrumen:{kompetensi}, isian, terkunci_pada, kehadiran:[baris bulan pkg_kehadiran],
                 rekomendasi: teks rekomendasi PKB}."""
    from pkg_template import SHEET_KOMPETENSI
    peta = PETA_CETAK['instrumen']
    kompetensi = (laporan.get('instrumen') or {}).get('kompetensi') or []
    if len(kompetensi) != len(SHEET_KOMPETENSI) or any(
            len(k.get('sub') or []) != len(peta['baris_sub']) or any(len(s.get('indikator') or []) != 5 for s in k['sub'])
            for k in kompetensi):
        raise StrukturTidakSesuai('Instrumen lembar tidak sesuai format cetak template resmi.')

    data = wb[peta['sheet_data']]
    p = laporan.get('periode') or {}
    tahun = str((p.get('tanggal_tutup') or p.get('tanggal_buka') or '')[:4] or '-')
    isian_sel = {
        'nama_guru': (laporan.get('guru') or {}).get('nama'), 'nip_guru': (laporan.get('guru') or {}).get('nip'),
        'tahun': tahun, 'nama_kepala': (laporan.get('kepala') or {}).get('nama'), 'nip_kepala': (laporan.get('kepala') or {}).get('nip'),
        'nama_sekolah': laporan.get('nama_sekolah'),
        'penilai': (laporan.get('penilai') or {}).get('nama'), 'nip_penilai': (laporan.get('penilai') or {}).get('nip'),
        'nomor_sk': (laporan.get('sk') or {}).get('nomor_sk'), 'tanggal_sk': _tanggal_teks((laporan.get('sk') or {}).get('tanggal_sk')),
        'periode_awal': _tanggal_teks(p.get('tanggal_buka')), 'periode_akhir': _tanggal_teks(p.get('tanggal_tutup')),
        'tanggal_nilai': _tanggal_teks(laporan.get('terkunci_pada')),
    }
    for kunci, sel in peta['data'].items():
        data[sel] = isian_sel.get(kunci) or '-'

    isi_kehadiran(wb, laporan.get('kehadiran'))
    isi_rekomendasi(wb, laporan.get('rekomendasi'))

    isian = laporan.get('isian') or {}
    kol = peta['kolom_centang']
    for k, sheet in zip(kompetensi, SHEET_KOMPETENSI):
        ws = wb[sheet]
        for s, baris in zip(k['sub'], peta['baris_sub']):
            centang = set((isian.get(s['kode']) or {}).get('centang') or [])
            for level in range(1, 6):
                ws[f'{kol}{baris + level - 1}'] = 'v' if level in centang else None


def isi_kehadiran(wb, kehadiran: Optional[List[Dict[str, Any]]], tugas_pokok: str = 'Guru') -> None:
    """Isi sheet INDEK-HDIR: label bulan periode & nilai terlambat/pulang cepat (menit) serta tanpa keterangan (hari)
    per bulan; baris bulan sisa dikosongkan (nilai 0). Contoh isian template ditimpa seluruhnya; rumus konversi,
    jumlah, dan indeks kehadiran tetap dari template. Tugas pokok diisi 'Guru' (template berisi contoh kepala sekolah)."""
    peta = PETA_CETAK['instrumen']['kehadiran']
    ws = wb[peta['sheet']]
    k = peta['kolom']
    ws[peta['tugas_pokok']] = tugas_pokok
    baris = list(kehadiran or [])[:peta['baris_maks']]
    for i in range(peta['baris_maks']):
        r = peta['baris_awal'] + i
        b = baris[i] if i < len(baris) else None
        ws[f"{k['label']}{r}"] = (b.get('label') or b.get('bulan')) if b else None
        for kunci in ('terlambat_menit', 'pulang_cepat_menit', 'tanpa_keterangan'):
            ws[f'{k[kunci]}{r}'] = int((b or {}).get(kunci) or 0)


def isi_rekomendasi(wb, teks: Optional[str]) -> None:
    """Tulis rekomendasi PKB di baris kosong 01-LAPORN (di bawah Deskripsi Capaian) pada salinan cetak, dengan gaya
    baris Deskripsi Capaian; tinggi baris menyesuaikan panjang teks. Berkas master tidak berubah."""
    if not teks:
        return
    peta = PETA_CETAK['instrumen']['rekomendasi']
    ws = wb[peta['sheet']]
    r, asal = peta['baris'], peta['gaya_dari']
    for kol, nilai in ((peta['label'], 'Rekomendasi PKB'), (peta['titik_dua'], ':'), (peta['isi'], teks)):
        sel = ws[f'{kol}{r}']
        sel.value = nilai
        if ws[f'{kol}{asal}'].has_style:
            sel._style = copy(ws[f'{kol}{asal}']._style)
    ws.merge_cells(peta['gabung'])
    baris_teks = max(1, -(-len(teks) // 48))
    ws.row_dimensions[r].height = max(15.0, 13.0 * baris_teks + 2)


def laporan_guru_excel(laporan: Dict[str, Any], hanya_rekap: bool = False) -> bytes:
    """Berkas .xlsx laporan PKG satu guru dari template instrumen resmi (seluruh sheet resmi ikut).
    `hanya_rekap=True` (untuk PDF): sheet selain 04-RKAP-PKG disembunyikan agar hanya lembar rekapitulasi yang
    dicetak; rumusnya tetap menghitung dari sheet tersembunyi."""
    def pengisi(wb):
        isi_laporan_guru(wb, laporan)
        if hanya_rekap:
            tampilkan_hanya(wb, [PETA_CETAK['instrumen']['sheet_rekap']])
    return isi_template_ke_bytes('instrumen', pengisi)
