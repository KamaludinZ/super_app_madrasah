"""API Input Tata Tertib terpadu (routers/tatib.py) di atas MongoDB sungguhan (fixture `mongo_db`)."""

import os
import sys

import pytest
from fastapi import HTTPException

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'backend'))

from routers import tatib as api  # noqa: E402
from migrations.migrate_tatib_struktur_baru import migrate  # noqa: E402

PETUGAS = {'id': 'u1', 'full_name': 'Guru Tatib', 'roles': ['guru_tata_tertib'], 'active_role': 'guru_tata_tertib'}


@pytest.fixture
async def db(mongo_db, monkeypatch):
    async def _tanpa_audit(*_a, **_k):
        return None
    monkeypatch.setattr(api, 'db', mongo_db)
    monkeypatch.setattr(api, 'log_audit', _tanpa_audit)
    await migrate(mongo_db)  # index unik seperti di server
    return mongo_db


def _req(model, **isian):
    return model(**isian)


async def test_kategori_tambah_rapikan_nama_dan_tolak_ganda(db):
    k = await api.create_kategori(_req(api.KategoriTatibRequest, nama='  Kedisiplinan  ', deskripsi=' '), PETUGAS)
    assert k['nama'] == 'Kedisiplinan' and k['nama_key'] == 'kedisiplinan' and k['deskripsi'] is None
    with pytest.raises(HTTPException) as e:
        await api.create_kategori(_req(api.KategoriTatibRequest, nama='KEDISIPLINAN'), PETUGAS)
    assert e.value.status_code == 409
    with pytest.raises(HTTPException) as e:
        await api.create_kategori(_req(api.KategoriTatibRequest, nama='   '), PETUGAS)
    assert e.value.status_code == 400
    daftar = await api.list_kategori(PETUGAS)
    assert [x['nama'] for x in daftar] == ['Kedisiplinan']


async def test_kategori_ubah_nama_ikut_ke_jenis_tapi_tidak_ke_riwayat(db):
    k = await api.create_kategori(_req(api.KategoriTatibRequest, nama='Kerapian'), PETUGAS)
    lain = await api.create_kategori(_req(api.KategoriTatibRequest, nama='Akhlak'), PETUGAS)
    await db.tatib_aturan.insert_one({'id': 'a1', 'kategori_id': k['id'], 'kategori_nama': 'Kerapian', 'kode_aturan': 'KER-01'})
    await db.tatib_penanganan.insert_one({'id': 'p1', 'kategori_id': k['id'], 'kategori_nama': 'Kerapian'})

    with pytest.raises(HTTPException) as e:
        await api.update_kategori(k['id'], _req(api.KategoriTatibRequest, nama='akhlak'), PETUGAS)
    assert e.value.status_code == 409

    u = await api.update_kategori(k['id'], _req(api.KategoriTatibRequest, nama='Kerapian Seragam', deskripsi='Atribut'), PETUGAS)
    assert (u['nama'], u['nama_key'], u['deskripsi']) == ('Kerapian Seragam', 'kerapian seragam', 'Atribut')
    assert (await db.tatib_aturan.find_one({'id': 'a1'}))['kategori_nama'] == 'Kerapian Seragam'
    assert (await db.tatib_penanganan.find_one({'id': 'p1'}))['kategori_nama'] == 'Kerapian'  # riwayat tetap
    # nama sama milik sendiri boleh disimpan ulang
    await api.update_kategori(lain['id'], _req(api.KategoriTatibRequest, nama='Akhlak'), PETUGAS)


async def test_kategori_hapus_ditolak_bila_dipakai(db):
    k = await api.create_kategori(_req(api.KategoriTatibRequest, nama='Ibadah'), PETUGAS)
    await db.tatib_penanganan.insert_one({'id': 'p1', 'kategori_id': k['id']})
    with pytest.raises(HTTPException) as e:
        await api.delete_kategori(k['id'], PETUGAS)
    assert e.value.status_code == 400
    await db.tatib_penanganan.delete_many({})
    assert (await api.delete_kategori(k['id'], PETUGAS))['message']
    with pytest.raises(HTTPException) as e:
        await api.get_kategori(k['id'], PETUGAS)
    assert e.value.status_code == 404


async def test_jenis_aturan_per_kategori_bentuk_kanonis(db):
    k1 = await api.create_kategori(_req(api.KategoriTatibRequest, nama='Kedisiplinan', urutan=1), PETUGAS)
    k2 = await api.create_kategori(_req(api.KategoriTatibRequest, nama='Akhlak', urutan=2), PETUGAS)
    await db.tatib_aturan.insert_many([
        # format lama (belum dimigrasi) tetap diterjemahkan
        {'id': 'a2', 'kode': 'DIS-02', 'nama_aturan': 'Bolos', 'kategori_id': k1['id'], 'kategori_nama': 'lama', 'poin': -15, 'is_active': True},
        {'id': 'a1', 'kode_aturan': 'DIS-01', 'nama': 'Terlambat', 'jalur_poin': 'minus', 'besar_poin': 5, 'kategori_id': k1['id'], 'aktif': True},
        {'id': 'a3', 'kode_aturan': 'AKH-01', 'nama': 'Imam salat', 'jalur_poin': 'plus', 'besar_poin': 5, 'kategori_id': k2['id'], 'aktif': True, 'deskripsi': 'x'},
        {'id': 'a4', 'kode_aturan': 'AKH-02', 'nama': 'Kasar', 'jalur_poin': 'minus', 'besar_poin': 10, 'kategori_id': k2['id'], 'aktif': False},
    ])
    semua = await api.list_jenis_aturan(user=PETUGAS)
    assert [j['id'] for j in semua] == ['a1', 'a2', 'a3', 'a4']  # urut kategori lalu kode
    assert semua[1] == {'id': 'a2', 'kode_aturan': 'DIS-02', 'nama': 'Bolos', 'jalur_poin': 'minus', 'besar_poin': 15, 'aktif': True,
                        'poin': -15, 'kategori_id': k1['id'], 'kategori_nama': 'Kedisiplinan', 'deskripsi': '', 'jumlah_catatan': 0}
    assert semua[2]['poin'] == 5 and semua[2]['deskripsi'] == 'x'

    per_kategori = await api.list_jenis_aturan(kategori_id=k2['id'], user=PETUGAS)
    assert [j['id'] for j in per_kategori] == ['a3', 'a4']
    aktif = await api.list_jenis_aturan(kategori_id=k2['id'], aktif=True, user=PETUGAS)
    assert [j['id'] for j in aktif] == ['a3']
    assert [j['id'] for j in await api.list_jenis_aturan(aktif=False, user=PETUGAS)] == ['a4']


@pytest.fixture
async def data_input(db, monkeypatch):
    """Kategori, jenis aturan (aktif/nonaktif), dan satu siswa ber-kelas."""
    def _tanpa_notifikasi(coro):
        coro.close()
    monkeypatch.setattr(api, 'spawn', _tanpa_notifikasi)
    k = await api.create_kategori(_req(api.KategoriTatibRequest, nama='Kedisiplinan'), PETUGAS)
    await db.classes.insert_one({'id': 'c7a', 'name': 'VII A'})
    await db.users.insert_one({'id': 's1', 'roles': ['siswa'], 'full_name': 'Ahmad', 'nis': '2401', 'nisn': '001', 'student_class_id': 'c7a'})
    await db.tatib_aturan.insert_many([
        {'id': 'j1', 'kode_aturan': 'DIS-01', 'nama': 'Terlambat', 'jalur_poin': 'minus', 'besar_poin': 5, 'kategori_id': k['id'], 'aktif': True},
        {'id': 'j2', 'kode_aturan': 'AKH-01', 'nama': 'Imam salat', 'jalur_poin': 'plus', 'besar_poin': 10, 'kategori_id': k['id'], 'aktif': True},
        {'id': 'j3', 'kode_aturan': 'DIS-09', 'nama': 'Lama', 'jalur_poin': 'minus', 'besar_poin': 3, 'kategori_id': k['id'], 'aktif': False},
    ])
    return {'kategori': k}


def _input(**isian):
    dasar = {'siswa_id': 's1', 'jenis_id': 'j1', 'tanggal': '2026-10-01', 'asal_input': 'Razia'}
    return api.InputTatibRequest(**{**dasar, **isian})


async def test_simpan_input_poin_dari_master_jenis(db, data_input):
    c = await api.simpan_input_tatib(_input(keterangan='  Telat 10 menit  ', semester_id='sem1'), PETUGAS)
    assert (c['poin'], c['jenis_poin'], c['kode_aturan'], c['jenis_nama']) == (-5, 'pelanggaran', 'DIS-01', 'Terlambat')
    assert (c['siswa_nama'], c['siswa_kelas'], c['kategori_nama']) == ('Ahmad', 'VII A', 'Kedisiplinan')
    assert (c['asal_input'], c['keterangan'], c['petugas_nama'], c['semester_id']) == ('Razia', 'Telat 10 menit', 'Guru Tatib', 'sem1')
    tersimpan = await db.tatib_penanganan.find_one({'id': c['id']}, {'_id': 0})
    # kompatibel dengan endpoint & rekap lama
    assert (tersimpan['tatib_id'], tersimpan['tatib_poin'], tersimpan['catatan'], tersimpan['tindak_lanjut']) == ('j1', -5, 'Telat 10 menit', [])

    plus = await api.simpan_input_tatib(_input(jenis_id='j2', asal_input='Ibadah'), PETUGAS)
    assert plus['poin'] == 10 and plus['jenis_poin'] == 'kebaikan' and plus['keterangan'] is None


@pytest.mark.parametrize('isian,status,pesan', [
    ({'asal_input': 'Lainnya'}, 400, 'Asal input'),
    ({'asal_input': ''}, 400, 'Asal input'),
    ({'tanggal': '2999-01-01'}, 400, 'melewati hari ini'),
    ({'tanggal': '01-10-2026'}, 400, 'tidak valid'),
    ({'keterangan': 'x' * 501}, 400, 'maksimal'),
    ({'siswa_id': 'tidak-ada'}, 404, 'Siswa'),
    ({'jenis_id': 'tidak-ada'}, 404, 'Jenis'),
    ({'jenis_id': 'j3'}, 400, 'tidak aktif'),
])
async def test_simpan_input_validasi(db, data_input, isian, status, pesan):
    with pytest.raises(HTTPException) as e:
        await api.simpan_input_tatib(_input(**isian), PETUGAS)
    assert e.value.status_code == status and pesan in e.value.detail
    assert await db.tatib_penanganan.count_documents({}) == 0


async def test_daftar_terbaru_urut_waktu_input_dan_filter_siswa(db, data_input):
    await db.users.insert_one({'id': 's2', 'roles': ['siswa'], 'full_name': 'Budi', 'student_class_id': 'c7a'})
    await db.tatib_penanganan.insert_one({'id': 'lama', 'siswa_id': 's2', 'tatib_id': 'j1', 'tatib_kode': 'DIS-01', 'tatib_nama': 'Terlambat',
                                          'tatib_poin': -5, 'catatan': 'catatan lama', 'tanggal': '2026-09-01', 'created_at': '2026-09-01T00:00:00'})
    a = await api.simpan_input_tatib(_input(tanggal='2026-09-05'), PETUGAS)
    b = await api.simpan_input_tatib(_input(siswa_id='s2', jenis_id='j2', asal_input='BIPH'), PETUGAS)

    hasil = await api.daftar_input_terbaru(batas=2, user=PETUGAS)
    assert [i['id'] for i in hasil['items']] == [b['id'], a['id']] and hasil['ada_lagi'] is True
    semua = await api.daftar_input_terbaru(batas=10, user=PETUGAS)
    lama = semua['items'][-1]
    # catatan format lama diterjemahkan ke bentuk daftar
    assert (lama['jenis_id'], lama['kode_aturan'], lama['keterangan'], lama['asal_input'], lama['poin']) == ('j1', 'DIS-01', 'catatan lama', None, -5)
    milik_s2 = await api.daftar_input_terbaru(batas=10, siswa_id='s2', user=PETUGAS)
    assert {i['siswa_id'] for i in milik_s2['items']} == {'s2'} and len(milik_s2['items']) == 2


async def test_ubah_input_hitung_ulang_poin_dan_pertahankan_jejak(db, data_input):
    c = await api.simpan_input_tatib(_input(), PETUGAS)
    await db.tatib_penanganan.update_one({'id': c['id']}, {'$set': {'tindak_lanjut': [{'uraian': 'dipanggil'}]}})
    admin = {'id': 'u9', 'full_name': 'Admin', 'roles': ['admin'], 'active_role': 'admin'}
    u = await api.ubah_input_tatib(c['id'], _input(jenis_id='j2', asal_input='Ibadah', keterangan='salah pilih'), admin)
    assert (u['poin'], u['jenis_poin'], u['kode_aturan'], u['asal_input'], u['keterangan']) == (10, 'kebaikan', 'AKH-01', 'Ibadah', 'salah pilih')
    assert (u['petugas_nama'], u['diubah_oleh_nama'], u['created_at']) == ('Guru Tatib', 'Admin', c['created_at'])
    assert u['tindak_lanjut'] == [{'uraian': 'dipanggil'}]

    with pytest.raises(HTTPException) as e:
        await api.ubah_input_tatib('tidak-ada', _input(), PETUGAS)
    assert e.value.status_code == 404
    with pytest.raises(HTTPException) as e:
        await api.ubah_input_tatib(c['id'], _input(asal_input='Lainnya'), PETUGAS)
    assert e.value.status_code == 400


async def test_ubah_input_jenis_lama_nonaktif_boleh_dipertahankan(db, data_input):
    await db.tatib_penanganan.insert_one({'id': 'x', 'siswa_id': 's1', 'tatib_id': 'j3', 'tanggal': '2026-10-01', 'created_at': '2026-10-01T00:00:00'})
    u = await api.ubah_input_tatib('x', _input(jenis_id='j3', asal_input='BAP'), PETUGAS)
    assert u['poin'] == -3 and u['asal_input'] == 'BAP'
    c = await api.simpan_input_tatib(_input(), PETUGAS)
    with pytest.raises(HTTPException) as e:  # pindah KE jenis nonaktif tetap ditolak
        await api.ubah_input_tatib(c['id'], _input(jenis_id='j3'), PETUGAS)
    assert e.value.status_code == 400


async def test_hapus_input(db, data_input):
    c = await api.simpan_input_tatib(_input(), PETUGAS)
    admin = {'id': 'u9', 'roles': ['admin'], 'active_role': 'admin'}
    assert (await api.hapus_input_tatib(c['id'], admin))['message']
    assert await db.tatib_penanganan.count_documents({}) == 0
    with pytest.raises(HTTPException) as e:
        await api.hapus_input_tatib(c['id'], admin)
    assert e.value.status_code == 404


async def test_hak_akses_endpoint_input():
    """Dependency: input (admin/guru tatib/waka kesiswaan), daftar (input + lihat), hapus hanya admin."""
    from routers.tatib_poin import wajib_input_tatib, wajib_lihat_tatib
    peran = lambda r: {'id': 'x', 'roles': [r], 'active_role': r}  # noqa: E731
    for r in ('admin', 'guru_tata_tertib', 'waka_kesiswaan'):
        assert await wajib_input_tatib(peran(r))
    for r in ('guru_bk', 'kepala_sekolah', 'wali_kelas', 'siswa'):
        with pytest.raises(HTTPException):
            await wajib_input_tatib(peran(r))
    assert await wajib_lihat_tatib(peran('kepala_sekolah'))
    for r in ('wali_kelas', 'siswa', 'guru_mapel'):
        with pytest.raises(HTTPException):
            await wajib_lihat_tatib(peran(r))
    hanya_admin = api.require_role('admin')
    assert await hanya_admin(peran('admin'))
    with pytest.raises(HTTPException):
        await hanya_admin(peran('guru_tata_tertib'))


def test_semua_endpoint_tatib_berpenjaga_peran():
    """Tiap endpoint routers/tatib.py memakai penjaga peran yang sesuai (bukan sekadar login)."""
    import inspect
    from routers.tatib_poin import wajib_akses_tatib, wajib_input_tatib, wajib_lihat_tatib

    def penjaga(route):
        return [d.call for d in route.dependant.dependencies]

    def hanya_admin(fn):
        return 'require_role' in fn.__qualname__ and fn.__closure__ and \
            any(c.cell_contents == ('admin',) for c in fn.__closure__)

    for route in api.router.routes:
        method = next(iter(route.methods))
        calls = penjaga(route)
        kunci = f'{method} {route.path}'
        if method == 'GET' and route.path in ('/tatib/kategori', '/tatib/kategori/{kategori_id}', '/tatib/jenis',
                                              '/tatib/jenis-aturan', '/tatib/aturan', '/tatib/aturan/{aturan_id}'):
            assert wajib_akses_tatib in calls, kunci
        elif method == 'GET' and route.path in ('/tatib/input/terbaru', '/tatib/stats/summary'):
            assert wajib_lihat_tatib in calls, kunci
        elif method == 'GET':
            # data per siswa: dibatasi di dalam endpoint (siswa: miliknya, wali kelas: kelasnya, lainnya 403)
            assert 'batasi_query_tatib' in inspect.getsource(route.endpoint), kunci
        elif method == 'DELETE' and route.path in ('/tatib/input/{catatan_id}', '/tatib/penanganan/{penanganan_id}'):
            assert any(hanya_admin(c) for c in calls), kunci
        else:
            assert wajib_input_tatib in calls, kunci


async def test_master_tatib_ditolak_untuk_peran_tanpa_akses_tatib():
    from routers.tatib_poin import wajib_akses_tatib
    for r in ('admin', 'guru_tata_tertib', 'guru_bk', 'kepala_sekolah', 'wali_kelas', 'siswa'):
        assert await wajib_akses_tatib({'id': 'x', 'roles': [r], 'active_role': r})
    for r in ('guru_mapel', 'guru_piket', 'staf_gtk'):
        with pytest.raises(HTTPException) as e:
            await wajib_akses_tatib({'id': 'x', 'roles': [r], 'active_role': r})
        assert e.value.status_code == 403


async def test_jalur_catat_lama_menerima_dan_menyimpan_asal_input(db, data_input):
    # aturan format lama berkondisi (dipakai Catat Poin Tatib)
    await db.tatib_aturan.insert_one({'id': 'lama1', 'kode': 'L-1', 'nama_aturan': 'Lama', 'poin': -4, 'jenis_poin': 'pelanggaran',
                                      'kategori_id': data_input['kategori']['id'], 'kategori_nama': 'Kedisiplinan'})
    req = api.PenangananRequest(siswa_id='s1', tatib_id='lama1', tanggal='2026-10-01', catatan='x', asal_input='Whatsapp')
    c = await api.catat_poin_jalur('pelanggaran', req, PETUGAS)
    assert (c['asal_input'], c['jenis_id'], c['kategori_id'], c['keterangan'], c['poin']) == ('Whatsapp', 'lama1', data_input['kategori']['id'], 'x', -4)

    # asal input wajib untuk catatan baru di semua endpoint catat
    for catat in (lambda r: api.catat_poin_jalur('pelanggaran', r, PETUGAS), lambda r: api.create_penanganan(r, PETUGAS)):
        with pytest.raises(HTTPException) as e:
            await catat(api.PenangananRequest(siswa_id='s1', tatib_id='lama1', tanggal='2026-10-01'))
        assert e.value.status_code == 400 and 'Asal input' in e.value.detail
    assert await db.tatib_penanganan.count_documents({}) == 1

    with pytest.raises(HTTPException) as e:
        await api.catat_poin_jalur('pelanggaran', api.PenangananRequest(siswa_id='s1', tatib_id='lama1', tanggal='2026-10-01', asal_input='SMS'), PETUGAS)
    assert e.value.status_code == 400

    u = await api.update_penanganan(c['id'], api.PenangananRequest(siswa_id='s1', tatib_id='lama1', tanggal='2026-10-01', asal_input='BAP', catatan='y'), PETUGAS)
    assert (u['asal_input'], u['keterangan']) == ('BAP', 'y')
    with pytest.raises(HTTPException):
        await api.update_penanganan(c['id'], api.PenangananRequest(siswa_id='s1', tatib_id='lama1', tanggal='2026-10-01', asal_input='x'), PETUGAS)
    # catatan lama tanpa asal input tetap bisa diubah tanpa mengisinya
    await db.tatib_penanganan.insert_one({'id': 'old', 'siswa_id': 's1', 'tatib_id': 'lama1', 'tanggal': '2026-09-01'})
    o = await api.update_penanganan('old', api.PenangananRequest(siswa_id='s1', tatib_id='lama1', tanggal='2026-09-02'), PETUGAS)
    assert o['tanggal'] == '2026-09-02' and o.get('asal_input') is None


def test_cek_asal_input():
    assert api.cek_asal_input(None) is None and api.cek_asal_input('') is None
    assert api.cek_asal_input('Razia', wajib=True) == 'Razia'
    for salah, wajib in (('razia', False), (None, True), ('', True)):
        with pytest.raises(HTTPException):
            api.cek_asal_input(salah, wajib=wajib)


async def test_daftar_tatib_selalu_memuat_asal_input_dan_bisa_difilter(db, data_input):
    admin = {'id': 'u9', 'full_name': 'Admin', 'roles': ['admin'], 'active_role': 'admin'}
    await api.simpan_input_tatib(_input(asal_input='Razia'), PETUGAS)
    await api.simpan_input_tatib(_input(jenis_id='j2', asal_input='Ibadah', tanggal='2026-10-02'), PETUGAS)
    await db.tatib_penanganan.insert_one({'id': 'lama', 'siswa_id': 's1', 'tatib_id': 'j1', 'tatib_poin': -5, 'catatan': 'c', 'tanggal': '2026-09-01'})

    semua = await api.list_penanganan(user=admin)
    assert [c['asal_input'] for c in semua] == ['Ibadah', 'Razia', None]
    lama = semua[-1]
    assert (lama['poin'], lama['jenis_poin'], lama['keterangan']) == (-5, 'pelanggaran', 'c')

    razia = await api.list_penanganan(asal_input='Razia', user=admin)
    assert [c['asal_input'] for c in razia] == ['Razia']
    with pytest.raises(HTTPException) as e:
        await api.list_penanganan(asal_input='razia', user=admin)
    assert e.value.status_code == 400

    detail = await api.get_penanganan('lama', admin)
    assert 'asal_input' in detail and detail['asal_input'] is None


def test_rapikan_catatan_memuat_asal_input():
    from routers.tatib_poin import rapikan_catatan
    r = rapikan_catatan({'id': 'x', 'tatib_poin': -3, 'catatan': 'lama'})
    assert (r['asal_input'], r['keterangan'], r['poin']) == (None, 'lama', -3)
    assert rapikan_catatan({'id': 'y', 'poin': 2, 'asal_input': 'BAP', 'keterangan': 'k', 'catatan': 'c'})['keterangan'] == 'k'


async def test_rekap_bisa_difilter_asal_input(db, data_input, monkeypatch):
    from routers import tatib_poin as rekap_api
    monkeypatch.setattr(rekap_api, 'db', db)
    pimpinan = {'id': 'k1', 'roles': ['kepala_sekolah'], 'active_role': 'kepala_sekolah'}
    await api.simpan_input_tatib(_input(asal_input='Razia'), PETUGAS)
    await api.simpan_input_tatib(_input(asal_input='Razia', tanggal='2026-10-02'), PETUGAS)
    await api.simpan_input_tatib(_input(jenis_id='j2', asal_input='Ibadah'), PETUGAS)

    semua = await rekap_api.rekap_poin(user=pimpinan)
    assert (semua['ringkasan']['jumlah_pelanggaran'], semua['ringkasan']['jumlah_kebaikan']) == (2, 1)
    razia = await rekap_api.rekap_poin(asal_input='Razia', user=pimpinan)
    assert (razia['ringkasan']['jumlah_pelanggaran'], razia['ringkasan']['jumlah_kebaikan'], razia['ringkasan']['total_minus']) == (2, 0, -10)
    ibadah = await rekap_api.rekap_poin(asal_input='Ibadah', user=pimpinan)
    assert ibadah['ringkasan']['jumlah_kebaikan'] == 1 and ibadah['ringkasan']['jumlah_pelanggaran'] == 0
    with pytest.raises(HTTPException) as e:
        await rekap_api.rekap_poin(asal_input='SMS', user=pimpinan)
    assert e.value.status_code == 400

    unduh = await rekap_api.ekspor_rekap_poin(asal_input='Screening pagi', user=pimpinan)
    assert 'screening-pagi' in unduh.headers['content-disposition']


async def test_jenis_aturan_pencarian_dan_jumlah_catatan(db):
    k = await api.create_kategori(_req(api.KategoriTatibRequest, nama='Kedisiplinan'), PETUGAS)
    await db.tatib_aturan.insert_many([
        {'id': 'a1', 'kode_aturan': 'DIS-01', 'nama': 'Terlambat masuk', 'jalur_poin': 'minus', 'besar_poin': 5, 'kategori_id': k['id'], 'aktif': True},
        {'id': 'a2', 'kode_aturan': 'DIS-02', 'nama': 'Membolos pelajaran', 'jalur_poin': 'minus', 'besar_poin': 15, 'kategori_id': k['id'], 'aktif': True},
        {'id': 'a3', 'kode_aturan': 'AKH-01', 'nama': 'Imam salat', 'jalur_poin': 'plus', 'besar_poin': 5, 'kategori_id': k['id'], 'aktif': False},
    ])
    await db.tatib_penanganan.insert_many([
        {'id': 'p1', 'jenis_id': 'a1', 'tatib_id': 'a1'}, {'id': 'p2', 'jenis_id': 'a1', 'tatib_id': 'a1'},
        {'id': 'p3', 'tatib_id': 'a3'},  # catatan lama tanpa jenis_id
    ])
    cari = lambda q, **kw: api.list_jenis_aturan(cari=q, user=PETUGAS, **kw)  # noqa: E731
    assert [j['id'] for j in await cari('dis')] == ['a1', 'a2']
    assert [j['id'] for j in await cari('  SALAT imam ')] == ['a3']
    assert [j['id'] for j in await cari('dis-02 bolos')] == ['a2']
    assert await cari('tidak-ada') == []
    assert [j['id'] for j in await cari('', aktif=True)] == ['a1', 'a2']
    jumlah = {j['id']: j['jumlah_catatan'] for j in await api.list_jenis_aturan(user=PETUGAS)}
    assert jumlah == {'a1': 2, 'a2': 0, 'a3': 1}


def _jenis_req(**kw):
    return api.JenisAturanRequest(**{'kode_aturan': 'dis-01', 'nama': 'Terlambat', 'jalur_poin': 'minus', 'besar_poin': 5, 'kategori_id': 'k?', **kw})


async def test_simpan_jenis_aturan_kode_unik_dan_field_kompatibel(db):
    k = await api.create_kategori(_req(api.KategoriTatibRequest, nama='Kedisiplinan'), PETUGAS)
    j = await api.tambah_jenis_aturan(_jenis_req(kategori_id=k['id'], deskripsi=' Datang telat '), PETUGAS)
    assert (j['kode_aturan'], j['poin'], j['aktif'], j['kategori_nama'], j['deskripsi'], j['jumlah_catatan']) == ('DIS-01', -5, True, 'Kedisiplinan', 'Datang telat', 0)
    doc = await db.tatib_aturan.find_one({'id': j['id']}, {'_id': 0})
    # field lama ikut terisi: alur Catat Poin Tatib & endpoint /tatib/aturan tetap bisa memakai jenis ini
    assert (doc['kode'], doc['nama_aturan'], doc['jenis_poin'], doc['poin'], doc['is_active']) == ('DIS-01', 'Terlambat', 'pelanggaran', -5, True)
    assert doc['kondisi'] == [{'id': 'umum', 'label': 'Setiap kejadian', 'poin': -5}]

    for kode in ('DIS-01', ' dis-01 '):
        with pytest.raises(HTTPException) as e:
            await api.tambah_jenis_aturan(_jenis_req(kode_aturan=kode, kategori_id=k['id']), PETUGAS)
        assert e.value.status_code == 409 and 'Terlambat' in e.value.detail
    await db.tatib_aturan.insert_one({'id': 'lama', 'kode': 'OLD-1', 'nama_aturan': 'Format lama', 'poin': -1})
    with pytest.raises(HTTPException) as e:  # kode di field lama juga dicek
        await api.tambah_jenis_aturan(_jenis_req(kode_aturan='old-1', kategori_id=k['id']), PETUGAS)
    assert e.value.status_code == 409
    with pytest.raises(HTTPException) as e:
        await api.tambah_jenis_aturan(_jenis_req(kode_aturan='X-1', kategori_id='tidak-ada'), PETUGAS)
    assert e.value.status_code == 404


async def test_ubah_jenis_aturan(db, data_input):
    k2 = await api.create_kategori(_req(api.KategoriTatibRequest, nama='Akhlak'), PETUGAS)
    # j1 dipakai satu catatan: riwayatnya tidak berubah saat jenis diubah
    c = await api.simpan_input_tatib(_input(), PETUGAS)
    u = await api.ubah_jenis_aturan('j1', _jenis_req(kode_aturan='dis-01', nama='Terlambat datang', jalur_poin='minus', besar_poin=7,
                                                      kategori_id=k2['id']), PETUGAS)
    assert (u['nama'], u['poin'], u['kategori_nama'], u['jumlah_catatan']) == ('Terlambat datang', -7, 'Akhlak', 1)
    doc = await db.tatib_aturan.find_one({'id': 'j1'}, {'_id': 0})
    assert doc['kondisi'][0]['poin'] == -7 and doc['poin'] == -7
    catatan = await db.tatib_penanganan.find_one({'id': c['id']}, {'_id': 0})
    assert (catatan['poin'], catatan['tatib_nama']) == (-5, 'Terlambat')

    with pytest.raises(HTTPException) as e:  # pakai kode jenis lain
        await api.ubah_jenis_aturan('j1', _jenis_req(kode_aturan='akh-01', kategori_id=k2['id']), PETUGAS)
    assert e.value.status_code == 409
    with pytest.raises(HTTPException) as e:
        await api.ubah_jenis_aturan('tidak-ada', _jenis_req(kategori_id=k2['id']), PETUGAS)
    assert e.value.status_code == 404


async def test_toggle_status_jenis_tanpa_menghapus_riwayat(db, data_input):
    c = await api.simpan_input_tatib(_input(), PETUGAS)
    j = await api.ubah_status_jenis_aturan('j1', api.StatusJenisRequest(aktif=False), PETUGAS)
    assert (j['aktif'], j['jumlah_catatan']) == (False, 1)
    doc = await db.tatib_aturan.find_one({'id': 'j1'}, {'_id': 0})
    assert doc['aktif'] is False and doc['is_active'] is False
    assert await db.tatib_penanganan.count_documents({'id': c['id']}) == 1  # riwayat tetap
    assert [x['id'] for x in await api.list_jenis_aturan(aktif=True, user=PETUGAS)] == ['j2']
    with pytest.raises(HTTPException) as e:  # tidak bisa dicatat lagi
        await api.simpan_input_tatib(_input(), PETUGAS)
    assert e.value.status_code == 400
    u = await api.ubah_input_tatib(c['id'], _input(asal_input='BAP'), PETUGAS)  # catatan lamanya tetap bisa diperbaiki
    assert u['asal_input'] == 'BAP'

    j = await api.ubah_status_jenis_aturan('j1', api.StatusJenisRequest(aktif=True), PETUGAS)
    assert j['aktif'] is True
    assert (await api.simpan_input_tatib(_input(), PETUGAS))['poin'] == -5
    with pytest.raises(HTTPException) as e:
        await api.ubah_status_jenis_aturan('tidak-ada', api.StatusJenisRequest(aktif=True), PETUGAS)
    assert e.value.status_code == 404


async def test_jenis_nonaktif_tersaring_dari_semua_jalur_pencatatan(db, data_input):
    kid = data_input['kategori']['id']
    # aturan lama: hanya `aktif` (dari Master Jenis) / hanya `is_active` / tanpa keduanya
    await db.tatib_aturan.insert_many([
        {'id': 'x1', 'kode': 'X-1', 'nama_aturan': 'Nonaktif baru', 'poin': -2, 'jenis_poin': 'pelanggaran', 'kategori_id': kid, 'aktif': False},
        {'id': 'x2', 'kode': 'X-2', 'nama_aturan': 'Nonaktif lama', 'poin': -2, 'jenis_poin': 'pelanggaran', 'kategori_id': kid, 'is_active': False},
        {'id': 'x3', 'kode': 'X-3', 'nama_aturan': 'Lama aktif', 'poin': -2, 'jenis_poin': 'pelanggaran', 'kategori_id': kid},
    ])
    pilihan = await api.list_aturan(jenis_poin='pelanggaran', is_active=True, user=PETUGAS)
    ids = {a['id'] for a in pilihan}
    assert 'x3' in ids and 'j1' in ids and not ids & {'x1', 'x2', 'j3'}
    nonaktif = {a['id'] for a in await api.list_aturan(is_active=False, user=PETUGAS)}
    assert {'x1', 'x2', 'j3'} <= nonaktif

    for tatib_id in ('x1', 'x2', 'j3'):
        with pytest.raises(HTTPException) as e:
            await api.catat_poin_jalur('pelanggaran', api.PenangananRequest(siswa_id='s1', tatib_id=tatib_id, tanggal='2026-10-01',
                                                                          asal_input='Razia'), PETUGAS)
        assert e.value.status_code == 400
    c = await api.catat_poin_jalur('pelanggaran', api.PenangananRequest(siswa_id='s1', tatib_id='x3', tanggal='2026-10-01',
                                                                      asal_input='Razia'), PETUGAS)
    with pytest.raises(HTTPException) as e:  # memindahkan catatan KE aturan nonaktif ditolak
        await api.update_penanganan(c['id'], api.PenangananRequest(siswa_id='s1', tatib_id='x1', tanggal='2026-10-01'), PETUGAS)
    assert e.value.status_code == 400


async def test_endpoint_aturan_lama_sinkron_dengan_master_jenis(db, data_input, monkeypatch):
    kid = data_input['kategori']['id']
    await db.tatib_jenis.insert_one({'id': 'sj', 'kategori_id': kid, 'nama': 'Sub'})
    req = api.TatibRequest(kategori_id=kid, jenis_id='sj', kode='ker-5', nama_aturan='Kuku', poin=-3)
    a = await api.create_aturan(req, PETUGAS)
    doc = await db.tatib_aturan.find_one({'id': a['id']}, {'_id': 0})
    assert (doc['kode'], doc['kode_aturan'], doc['nama'], doc['jalur_poin'], doc['besar_poin'], doc['aktif']) == ('KER-5', 'KER-5', 'Kuku', 'minus', 3, True)
    with pytest.raises(HTTPException) as e:  # kode ganda tanpa beda huruf, juga terhadap master jenis
        await api.create_aturan(api.TatibRequest(kategori_id=kid, jenis_id='sj', kode='dis-01', nama_aturan='X', poin=-1), PETUGAS)
    assert e.value.status_code == 409

    await api.update_aturan(a['id'], api.TatibRequest(kategori_id=kid, jenis_id='sj', kode='KER-5', nama_aturan='Kuku panjang',
                                                      poin=-4, is_active=False), PETUGAS)
    doc = await db.tatib_aturan.find_one({'id': a['id']}, {'_id': 0})
    assert (doc['nama'], doc['besar_poin'], doc['aktif'], doc['is_active']) == ('Kuku panjang', 4, False, False)
    master = {j['id']: j for j in await api.list_jenis_aturan(user=PETUGAS)}
    assert master[a['id']]['aktif'] is False


async def test_jenis_hanya_field_kanonis_bisa_dicatat_lewat_jalur_lama(db, data_input):
    """Jenis j1/j2 (fixture) hanya punya field Master Jenis: Catat Poin Tatib lama harus menghitungnya benar."""
    pilihan = {a['id']: a for a in await api.list_aturan(jenis_poin='kebaikan', is_active=True, user=PETUGAS)}
    assert 'j2' in pilihan and (pilihan['j2']['poin'], pilihan['j2']['jenis_poin'], pilihan['j2']['kode']) == (10, 'kebaikan', 'AKH-01')
    c = await api.catat_poin_jalur('pelanggaran', api.PenangananRequest(siswa_id='s1', tatib_id='j1', tanggal='2026-10-01',
                                                                      asal_input='Razia'), PETUGAS)
    assert (c['poin'], c['jenis_poin'], c['tatib_kode'], c['tatib_nama']) == (-5, 'pelanggaran', 'DIS-01', 'Terlambat')
    with pytest.raises(HTTPException):  # jalur salah
        await api.catat_poin_jalur('kebaikan', api.PenangananRequest(siswa_id='s1', tatib_id='j1', tanggal='2026-10-01',
                                                                   asal_input='Razia'), PETUGAS)


async def test_setiap_jalur_pencatatan_menulis_asal_input(db, data_input):
    """Model pencatatan: kunci asal_input selalu ada di dokumen yang ditulis jalur mana pun."""
    await db.tatib_aturan.insert_one({'id': 'lama1', 'kode': 'L-1', 'nama_aturan': 'Lama', 'poin': -4, 'jenis_poin': 'pelanggaran',
                                      'kategori_id': data_input['kategori']['id']})
    a = await api.simpan_input_tatib(_input(asal_input='BIPH'), PETUGAS)
    b = await api.catat_poin_jalur('pelanggaran', api.PenangananRequest(siswa_id='s1', tatib_id='lama1', tanggal='2026-10-01',
                                                                      asal_input='Keputrian'), PETUGAS)
    c = await api.create_penanganan(api.PenangananRequest(siswa_id='s1', tatib_id='lama1', tanggal='2026-10-01',
                                                          asal_input='Whatsapp'), PETUGAS)
    docs = {d['id']: d async for d in db.tatib_penanganan.find({}, {'_id': 0})}
    assert {docs[x['id']]['asal_input'] for x in (a, b, c)} == {'BIPH', 'Keputrian', 'Whatsapp'}
    await api.ubah_input_tatib(a['id'], _input(asal_input='Screening pagi'), PETUGAS)
    assert (await db.tatib_penanganan.find_one({'id': a['id']}))['asal_input'] == 'Screening pagi'
    assert all('asal_input' in d for d in docs.values())


async def test_semua_endpoint_daftar_mengirim_asal_input(db, data_input, monkeypatch):
    from routers import tatib_poin as poin_api
    monkeypatch.setattr(poin_api, 'db', db)
    admin = {'id': 'u9', 'full_name': 'Admin', 'roles': ['admin'], 'active_role': 'admin'}
    await api.simpan_input_tatib(_input(asal_input='Razia'), PETUGAS)
    await db.tatib_penanganan.insert_one({'id': 'lama', 'siswa_id': 's1', 'tatib_id': 'j1', 'tatib_poin': -5, 'kategori_nama': 'Kedisiplinan',
                                          'tanggal': '2026-09-01', 'catatan': 'c'})
    daftar = await api.list_penanganan(user=admin)
    terbaru = (await api.daftar_input_terbaru(batas=10, user=admin))['items']
    stats = await api.get_siswa_stats('s1', user=admin)
    poin_siswa = await poin_api.paket_poin_siswa('s1')
    for nama, items in (('penanganan', daftar), ('input terbaru', terbaru), ('stats siswa', stats['records']), ('poin siswa', poin_siswa['records'])):
        asal = {i['id']: i['asal_input'] for i in items}
        assert asal == {next(i['id'] for i in items if i['id'] != 'lama'): 'Razia', 'lama': None}, nama
    assert (stats['total_poin'], stats['total_pelanggaran'], stats['by_asal_input']) == (-10, 2, {'Razia': 1, '-': 1})
    assert stats['by_kategori'] == {'Kedisiplinan': {'count': 2, 'total_poin': -10}}


async def test_filter_asal_input_termasuk_catatan_lama(db, data_input, monkeypatch):
    from routers import tatib_poin as poin_api
    monkeypatch.setattr(poin_api, 'db', db)
    admin = {'id': 'u9', 'full_name': 'Admin', 'roles': ['admin'], 'active_role': 'admin'}
    pimpinan = {'id': 'k1', 'roles': ['kepala_sekolah'], 'active_role': 'kepala_sekolah'}
    baru = await api.simpan_input_tatib(_input(asal_input='Razia'), PETUGAS)
    await db.tatib_penanganan.insert_many([
        {'id': 'lama1', 'siswa_id': 's1', 'tatib_id': 'j1', 'tatib_poin': -5, 'tanggal': '2026-09-01', 'created_at': '2026-09-01T00:00:00'},
        {'id': 'lama2', 'siswa_id': 's1', 'tatib_id': 'j1', 'tatib_poin': -5, 'tanggal': '2026-09-02', 'asal_input': None, 'created_at': '2026-09-02T00:00:00'},
    ])
    ids = lambda items: sorted(i['id'] for i in items)  # noqa: E731
    assert ids(await api.list_penanganan(asal_input='-', user=admin)) == ['lama1', 'lama2']
    assert ids(await api.list_penanganan(asal_input='Razia', user=admin)) == [baru['id']]
    assert ids((await api.daftar_input_terbaru(asal_input='-', user=admin))['items']) == ['lama1', 'lama2']
    assert ids((await api.daftar_input_terbaru(asal_input='Razia', user=admin))['items']) == [baru['id']]
    for salah in ('razia', 'SMS'):
        with pytest.raises(HTTPException) as e:
            await api.daftar_input_terbaru(asal_input=salah, user=admin)
        assert e.value.status_code == 400
    kosong = await poin_api.rekap_poin(asal_input='-', user=pimpinan)
    assert kosong['ringkasan']['jumlah_pelanggaran'] == 2
    with pytest.raises(HTTPException):
        await poin_api.rekap_poin(asal_input='x', user=pimpinan)


async def test_api_data_tatib_filter_kategori_jenis(db, data_input):
    admin = {'id': 'u9', 'full_name': 'Admin', 'roles': ['admin'], 'active_role': 'admin'}
    kid = data_input['kategori']['id']
    a = await api.simpan_input_tatib(_input(), PETUGAS)                           # j1 (kategori kid)
    b = await api.simpan_input_tatib(_input(jenis_id='j2', asal_input='BAP'), PETUGAS)  # j2 (kategori kid)
    await db.tatib_penanganan.insert_one({'id': 'lama', 'siswa_id': 's1', 'tatib_id': 'j1', 'tatib_kode': 'DIS-01', 'tatib_poin': -5,
                                          'kategori_nama': 'Kedisiplinan', 'tanggal': '2026-09-01'})  # tanpa jenis_id & kategori_id
    ids = lambda items: sorted(i['id'] for i in items)  # noqa: E731
    assert ids(await api.list_penanganan(jenis_id='j1', user=admin)) == sorted([a['id'], 'lama'])
    assert ids(await api.list_penanganan(jenis_id='j2', user=admin)) == [b['id']]
    assert ids(await api.list_penanganan(kategori_id=kid, user=admin)) == sorted([a['id'], b['id']])
    assert ids(await api.list_penanganan(kategori='Kedisiplinan', user=admin)) == sorted([a['id'], b['id'], 'lama'])
    assert ids(await api.list_penanganan(kategori='Kedisiplinan', jenis_id='j1', asal_input='-', user=admin)) == ['lama']
    lama = next(i for i in await api.list_penanganan(user=admin) if i['id'] == 'lama')
    assert (lama['jenis_id'], lama['kode_aturan'], lama['kategori_id']) == ('j1', 'DIS-01', None)


def test_susun_rekap_per_kategori_dan_asal_input():
    from tatib_poin import susun_rekap
    rec = [
        {'siswa_id': 's1', 'poin': -5, 'kategori_nama': 'Kedisiplinan', 'asal_input': 'Razia', 'tatib_kode': 'DIS-01', 'tatib_nama': 'Telat'},
        {'siswa_id': 's1', 'poin': -10, 'kategori_nama': 'Kedisiplinan', 'asal_input': 'BIPH', 'tatib_kode': 'DIS-02', 'tatib_nama': 'Bolos'},
        {'siswa_id': 's2', 'poin': 7, 'kategori_nama': 'Akhlak', 'asal_input': 'Razia', 'tatib_kode': 'AKH-01', 'tatib_nama': 'Imam'},
        {'siswa_id': 's2', 'tatib_poin': -3, 'kategori_nama': None, 'asal_input': '  '},  # catatan lama
    ]
    r = susun_rekap(rec, {'VII A': ['s1', 's2']})
    assert r['per_kategori'] == [
        {'kategori': 'Kedisiplinan', 'jumlah_pelanggaran': 2, 'jumlah_kebaikan': 0, 'total_minus': -15, 'total_plus': 0},
        {'kategori': 'Lainnya', 'jumlah_pelanggaran': 1, 'jumlah_kebaikan': 0, 'total_minus': -3, 'total_plus': 0},
        {'kategori': 'Akhlak', 'jumlah_pelanggaran': 0, 'jumlah_kebaikan': 1, 'total_minus': 0, 'total_plus': 7},
    ]
    assert r['per_asal_input'] == [
        {'asal_input': 'BIPH', 'jumlah': 1, 'total_minus': -10, 'total_plus': 0},
        {'asal_input': 'Razia', 'jumlah': 2, 'total_minus': -5, 'total_plus': 7},
        {'asal_input': None, 'jumlah': 1, 'total_minus': -3, 'total_plus': 0},
    ]
    assert {'kode': 'DIS-01', 'nama': 'Telat', 'kategori_nama': 'Kedisiplinan', 'jumlah': 1} in r['teratas_pelanggaran']


def test_ekspor_rekap_memuat_sheet_kategori_dan_asal():
    import io
    from openpyxl import load_workbook
    from excel_io import export_rekap_tatib_xlsx
    from tatib_poin import susun_rekap
    rec = [{'siswa_id': 's1', 'poin': -5, 'kategori_nama': 'Kedisiplinan', 'asal_input': 'Razia', 'tatib_kode': 'D', 'tatib_nama': 'T'},
           {'siswa_id': 's1', 'tatib_poin': -3}]
    wb = load_workbook(io.BytesIO(export_rekap_tatib_xlsx(susun_rekap(rec, {'VII A': ['s1']}))))
    assert {'Per Kategori', 'Per Asal Input'} <= set(wb.sheetnames)
    nilai = [c.value for row in wb['Per Asal Input'].iter_rows() for c in row]
    assert 'Razia' in nilai and '- (belum diisi)' in nilai
    assert 'Kedisiplinan' in [c.value for row in wb['Per Kategori'].iter_rows() for c in row]


@pytest.mark.parametrize('peran,harapan', [
    *[(p, {'a1', 'a2', 'b1'}) for p in ('admin', 'guru_tata_tertib', 'waka_kesiswaan', 'guru_bk', 'kepala_sekolah',
                                         'penjamin_mutu', 'kepala_tata_usaha', 'waka_kurikulum')],
    ('wali_kelas', {'a1', 'a2'}),   # hanya siswa kelas yang diampu (VII A: s1)
    ('siswa', {'a1', 'a2'}),        # hanya miliknya (user id s1)
    ('guru_mapel', None),           # tanpa akses tatib: 403
])
async def test_daftar_tatib_dapat_diakses_semua_peran_sesuai_cakupan(db, monkeypatch, peran, harapan):
    from routers import tatib_poin as poin_api
    import rombel
    monkeypatch.setattr(poin_api, 'db', db)

    async def _perwalian(user, *_a, **_k):
        return {'id': 'c7a', 'name': 'VII A'} if user.get('id') == 'wali1' else None
    monkeypatch.setattr(rombel, 'kelas_perwalian', _perwalian)
    await db.classes.insert_many([{'id': 'c7a', 'name': 'VII A'}, {'id': 'c8b', 'name': 'VIII B'}])
    await db.users.insert_many([
        {'id': 's1', 'roles': ['siswa'], 'full_name': 'Ahmad', 'student_class_id': 'c7a'},
        {'id': 's2', 'roles': ['siswa'], 'full_name': 'Budi', 'student_class_id': 'c8b'},
    ])
    await db.tatib_penanganan.insert_many([
        {'id': 'a1', 'siswa_id': 's1', 'tatib_poin': -5, 'asal_input': 'Razia', 'tanggal': '2026-10-01'},
        {'id': 'a2', 'siswa_id': 's1', 'tatib_poin': 3, 'tanggal': '2026-10-02'},
        {'id': 'b1', 'siswa_id': 's2', 'tatib_poin': -2, 'asal_input': 'Razia', 'tanggal': '2026-10-03'},
    ])
    uid = {'wali_kelas': 'wali1', 'siswa': 's1'}.get(peran, 'u-' + peran)
    user = {'id': uid, 'full_name': peran, 'roles': [peran], 'active_role': peran}
    if harapan is None:
        with pytest.raises(HTTPException) as e:
            await api.list_penanganan(user=user)
        assert e.value.status_code == 403
        return
    semua = {c['id'] for c in await api.list_penanganan(user=user)}
    assert semua == harapan
    assert all('asal_input' in c for c in await api.list_penanganan(user=user))
    razia = {c['id'] for c in await api.list_penanganan(asal_input='Razia', user=user)}
    assert razia == harapan & {'a1', 'b1'}
    # detail & statistik siswa ikut cakupan: siswa s2 hanya untuk peran input/lihat
    if 'b1' in harapan:
        assert (await api.get_penanganan('b1', user))['id'] == 'b1'
        assert (await api.get_siswa_stats('s2', user=user))['total_records'] == 1
    else:
        for panggil in (lambda: api.get_penanganan('b1', user), lambda: api.get_siswa_stats('s2', user=user)):
            with pytest.raises(HTTPException) as e:
                await panggil()
            assert e.value.status_code == 403
    assert (await api.get_siswa_stats('s1', user=user))['total_records'] == 2


async def test_rekap_walikelas_per_kategori_dan_asal(db, monkeypatch):
    from routers import tatib_poin as poin_api
    import rombel
    monkeypatch.setattr(poin_api, 'db', db)

    async def _perwalian(user, *_a, **_k):
        return {'id': 'c7a', 'name': 'VII A'}
    monkeypatch.setattr(rombel, 'kelas_perwalian', _perwalian)
    await db.users.insert_many([
        {'id': 's1', 'roles': ['siswa'], 'full_name': 'Ahmad', 'student_class_id': 'c7a', 'is_active': True},
        {'id': 's2', 'roles': ['siswa'], 'full_name': 'Budi', 'student_class_id': 'c8b', 'is_active': True},
    ])
    await db.tatib_penanganan.insert_many([
        {'id': 'a1', 'siswa_id': 's1', 'poin': -5, 'kategori_nama': 'Kedisiplinan', 'asal_input': 'Razia'},
        {'id': 'a2', 'siswa_id': 's1', 'poin': 4, 'kategori_nama': 'Akhlak', 'asal_input': None},
        {'id': 'b1', 'siswa_id': 's2', 'poin': -9, 'kategori_nama': 'Kedisiplinan', 'asal_input': 'Razia'},  # kelas lain
    ])
    wali = {'id': 'w1', 'roles': ['wali_kelas'], 'active_role': 'wali_kelas'}
    data = await poin_api.daftar_siswa_walikelas(wali)
    assert [s['id'] for s in data['siswa']] == ['s1']
    assert data['per_kategori'] == [
        {'kategori': 'Kedisiplinan', 'jumlah_pelanggaran': 1, 'jumlah_kebaikan': 0, 'total_minus': -5, 'total_plus': 0},
        {'kategori': 'Akhlak', 'jumlah_pelanggaran': 0, 'jumlah_kebaikan': 1, 'total_minus': 0, 'total_plus': 4},
    ]
    assert data['per_asal_input'] == [{'asal_input': 'Razia', 'jumlah': 1, 'total_minus': -5, 'total_plus': 0},
                                      {'asal_input': None, 'jumlah': 1, 'total_minus': 0, 'total_plus': 4}]


async def test_saldo_dan_riwayat_poin_siswa_struktur_baru(db, data_input, monkeypatch):
    from routers import tatib_poin as poin_api
    monkeypatch.setattr(poin_api, 'db', db)
    await api.simpan_input_tatib(_input(tanggal='2026-10-01'), PETUGAS)                                  # -5 Kedisiplinan, Razia
    await api.simpan_input_tatib(_input(jenis_id='j2', asal_input='Ibadah', tanggal='2026-10-03'), PETUGAS)  # +10
    await db.tatib_penanganan.insert_one({'id': 'lama', 'siswa_id': 's1', 'tatib_id': 'x', 'tatib_kode': 'L-1', 'tatib_poin': -3,
                                          'tanggal': '2026-09-01', 'catatan': 'lama'})
    siswa = {'id': 's1', 'full_name': 'Ahmad', 'roles': ['siswa'], 'active_role': 'siswa'}
    p = await poin_api.poin_saya(user=siswa)
    assert (p['total_plus'], p['total_minus'], p['saldo'], p['jumlah_kebaikan'], p['jumlah_pelanggaran']) == (10, -8, 2, 1, 2)
    assert p['siswa']['kelas'] == 'VII A'
    assert [r['tanggal'] for r in p['records']] == ['2026-10-03', '2026-10-01', '2026-09-01']  # terbaru dulu
    terbaru = p['records'][0]
    assert (terbaru['kode_aturan'], terbaru['kategori_nama'], terbaru['asal_input'], terbaru['jenis_id']) == ('AKH-01', 'Kedisiplinan', 'Ibadah', 'j2')
    lama = p['records'][-1]
    assert (lama['poin'], lama['keterangan'], lama['asal_input'], lama['kode_aturan']) == (-3, 'lama', None, 'L-1')
    assert {k['kategori']: k['total_plus'] + k['total_minus'] for k in p['per_kategori']} == {'Kedisiplinan': 5, 'Lainnya': -3}
    assert [a['asal_input'] for a in p['per_asal_input']] == ['Ibadah', 'Razia', None]
    # pimpinan melihat poin siswa yang sama lewat /tatib/poin/siswa/{id}
    lihat = await poin_api.poin_siswa('s1', user={'id': 'k', 'roles': ['kepala_sekolah'], 'active_role': 'kepala_sekolah'})
    assert lihat['saldo'] == 2 and len(lihat['records']) == 3


# ------------------------------------------------------------
# Impor input tata tertib dari berkas (POST /tatib/input/impor)
# ------------------------------------------------------------

def _impor(*baris, **isian):
    return api.ImporInputTatibRequest(baris=[api.BarisImporTatib(baris=i + 2, data=d) for i, d in enumerate(baris)], **isian)


def _baris(**isian):
    return {'nama': 'Ahmad', 'siswa_id': 's1', 'kode_aturan': 'DIS-01', 'tanggal': '2026-10-01', 'asal_input': 'Razia', **isian}


@pytest.mark.parametrize('teks,hasil', [
    ('2026-10-05', '2026-10-05'), ('05/10/2026', '2026-10-05'), ('5-10-2026', '2026-10-05'), ('05.10.26', '2026-10-05'),
    ('5 Oktober 2026', '2026-10-05'), ('Senin, 5 Okt 2026', '2026-10-05'), ('5 Agustus 2026', '2026-08-05'),
    ('31/02/2026', None), ('10/2026', None), ('kemarin', None), ('', None),
])
def test_tanggal_impor(teks, hasil):
    assert api.tanggal_impor(teks) == hasil


async def test_impor_baris_dirapikan_dan_tersimpan(db, data_input):
    r = await api.impor_input_tatib(_impor(
        _baris(kode_aturan=' dis-01 ', tanggal='01/10/2026', asal_input='  razia ', keterangan=' telat apel '),
        _baris(kode_aturan='AKH-01', asal_input='IBADAH', semester='x'),
    ), PETUGAS)
    a, b = r['hasil']
    assert (a['baris'], a['status'], a['identitas'], a['kolom']) == (2, 'berhasil', 'Ahmad (VII A)', '-')
    assert 'DIS-01' in a['pesan'] and '(-5)' in a['pesan'] and a['catatan_id']
    assert b['status'] == 'berhasil' and '(+10)' in b['pesan']
    c = await db.tatib_penanganan.find_one({'id': a['catatan_id']}, {'_id': 0})
    assert (c['tanggal'], c['asal_input'], c['keterangan'], c['poin'], c['petugas_nama']) == ('2026-10-01', 'Razia', 'telat apel', -5, 'Guru Tatib')


async def test_impor_cari_siswa_dari_nama_bila_tanpa_siswa_id(db, data_input):
    await db.users.insert_many([
        {'id': 's2', 'roles': ['siswa'], 'full_name': 'Budi Santoso'},
        {'id': 's3', 'roles': ['siswa'], 'full_name': 'Citra Dewi'},
        {'id': 's4', 'roles': ['siswa'], 'full_name': 'citra  dewi'},
        {'id': 's5', 'roles': ['siswa'], 'full_name': 'Dodi Keluar', 'mutation_type': 'keluar'},
    ])
    r = await api.impor_input_tatib(_impor(
        _baris(siswa_id=None, nama='  budi   SANTOSO '),
        _baris(siswa_id='', nama='Citra Dewi'),
        _baris(siswa_id=None, nama='Dodi Keluar'),
        _baris(siswa_id=None, nama='Eka'),
        _baris(siswa_id=None, nama=' '),
    ), PETUGAS)
    st = [(h['status'], h['kolom']) for h in r['hasil']]
    assert st == [('berhasil', '-'), ('gagal', 'Nama'), ('gagal', 'Nama'), ('gagal', 'Nama'), ('gagal', 'Nama')]
    assert r['hasil'][0]['identitas'] == 'Budi Santoso'
    assert '2 siswa' in r['hasil'][1]['pesan'] and 'tidak ditemukan' in r['hasil'][3]['pesan'] and 'kosong' in r['hasil'][4]['pesan']
    assert await db.tatib_penanganan.count_documents({}) == 1


@pytest.mark.parametrize('isian,kolom,pesan', [
    ({'kode_aturan': ''}, 'Kode Aturan', 'kosong'),
    ({'kode_aturan': 'XYZ-99'}, 'Kode Aturan', 'tidak terdaftar'),
    ({'kode_aturan': 'dis-09'}, 'Kode Aturan', 'tidak aktif'),
    ({'tanggal': ''}, 'Tanggal Kejadian', 'kosong'),
    ({'tanggal': '31/02/2026'}, 'Tanggal Kejadian', 'tidak dikenali'),
    ({'tanggal': '01/01/2999'}, 'Tanggal Kejadian', 'melewati hari ini'),
    ({'asal_input': 'Lainnya'}, 'Asal Input', 'Asal input'),
    ({'asal_input': None}, 'Asal Input', 'Asal input'),
    ({'keterangan': 'x' * 501}, 'Keterangan', 'maksimal'),
    ({'siswa_id': 'tidak-ada'}, 'Nama', 'Siswa tidak ditemukan'),
])
async def test_impor_baris_gagal_dengan_keterangan_tanpa_menghentikan_baris_lain(db, data_input, isian, kolom, pesan):
    r = await api.impor_input_tatib(_impor(_baris(**isian), _baris()), PETUGAS)
    gagal, lain = r['hasil']
    assert (gagal['status'], gagal['kolom']) == ('gagal', kolom) and pesan in gagal['pesan']
    assert lain['status'] == 'berhasil'
    assert await db.tatib_penanganan.count_documents({}) == 1


async def test_impor_ulang_tidak_menggandakan_catatan(db, data_input):
    pertama = await api.impor_input_tatib(_impor(_baris(keterangan='x'), _baris()), PETUGAS)
    assert [h['status'] for h in pertama['hasil']] == ['berhasil', 'berhasil']
    ulang = await api.impor_input_tatib(_impor(_baris(keterangan='x'), _baris(), _baris(tanggal='2026-10-02')), PETUGAS)
    assert [h['status'] for h in ulang['hasil']] == ['tanpa_perubahan', 'tanpa_perubahan', 'berhasil']
    assert 'Sudah tercatat' in ulang['hasil'][0]['pesan']
    assert await db.tatib_penanganan.count_documents({}) == 3


async def test_impor_batas_baris_per_kiriman(db, data_input):
    with pytest.raises(HTTPException) as e:
        await api.impor_input_tatib(_impor(*[_baris()] * (api.MAKS_BARIS_IMPOR + 1)), PETUGAS)
    assert e.value.status_code == 413
    assert await db.tatib_penanganan.count_documents({}) == 0
