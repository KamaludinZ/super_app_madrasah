"""Cetak rekap Excel dari template resmi: isi baris, rumus, sisip baris >10 guru, master tidak berubah."""

import io
import os
import sys
from datetime import date

from openpyxl import load_workbook

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'backend'))

import pkg_cetak  # noqa: E402
import pkg_template  # noqa: E402

PERIODE = {'id': 'p1', 'semester': 'Ganjil', 'tahun_pelajaran': '2026/2027', 'tanggal_buka': '2026-10-01', 'tanggal_tutup': '2026-11-30'}


def _guru(n):
    return [{'nama': f'Guru {i:02d}', 'mapel': 'IPA', 'capaian_persen': 50 + i * 3.333} for i in range(1, n + 1)]


def _buka(isi):
    return load_workbook(io.BytesIO(isi))['Rekap hasil PKG']


def test_rekap_tiga_guru():
    asli = pkg_template.hitung_checksum(pkg_template.path_template('rekap'))
    ws = _buka(pkg_cetak.rekap_excel(_guru(3), PERIODE, kepala='Drs. H. Mahmud, M.Pd.', tanggal_cetak=date(2026, 12, 5)))
    assert ws['A1'].value == 'REKAP HASIL PENILAIAN KINERJA GURU (PKG) TAHUN 2026'
    assert ws['A2'].value == 'PADA SEMESTER GANJIL TAHUN PELAJARAN 2026/2027'
    assert [ws[f'A{r}'].value for r in range(6, 10)] == [1, 2, 3, None]
    assert [ws[f'B{r}'].value for r in (6, 8)] == ['Guru 01', 'Guru 03'] and ws['D6'].value == 'IPA'
    assert ws['F6'].value == 53.33 and ws['G6'].value.startswith('=IF(F6="","",IF(F6>=91,"Sangat Baik"')
    assert ws['G9'].value is None and ws['B15'].value is None  # baris sisa dikosongkan
    assert ws['F16'].value == '=SUM(F6:F15)' and 'AVERAGE(F6:F15)' in ws['F17'].value and ws['G17'].value.startswith('=IF(F17')
    assert ws['G19'].value == 'Malang, 5 Desember 2026' and ws['G24'].value == 'Drs. H. Mahmud, M.Pd.'
    assert pkg_template.hitung_checksum(pkg_template.path_template('rekap')) == asli  # master utuh


def test_rekap_lebih_dari_sepuluh_guru():
    ws = _buka(pkg_cetak.rekap_excel(_guru(12), PERIODE, tanggal_cetak=date(2026, 12, 5)))
    assert [ws[f'A{r}'].value for r in (6, 15, 16, 17)] == [1, 10, 11, 12]
    assert ws['B17'].value == 'Guru 12' and ws['A17'].border.left.style == 'thin'  # gaya baris ikut tersalin
    assert ws['A18'].value == 'Jumlah' and ws['F18'].value == '=SUM(F6:F17)' and ws['A19'].value.startswith('Rata')
    gabung = {str(r) for r in ws.merged_cells.ranges}
    assert {'A18:E18', 'A19:E19', 'A1:H1', 'A2:H2'} <= gabung and 'A16:E16' not in gabung
    assert ws['G21'].value == 'Malang, 5 Desember 2026' and ws['G22'].value.startswith('Kepala')


def test_predikat_rekap():
    assert [pkg_cetak.predikat_rekap(n) for n in (95, 91, 80, 61, 60, 59.9, None)] == \
        ['Sangat Baik', 'Sangat Baik', 'Baik', 'Cukup', 'Sedang', 'Kurang', '']
    assert pkg_cetak.tanggal_indonesia(date(2026, 1, 9)) == '9 Januari 2026'
