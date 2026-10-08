"""API poin tatib: akses per peran, pantauan wali kelas, poin siswa, rekap (db tiruan)."""

import asyncio
import os
import sys

import pytest
from fastapi import HTTPException

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'backend'))

from routers import tatib_poin as api  # noqa: E402
from tatib_poin import akses_tatib, rangkum_poin  # noqa: E402


def _cocok(d, q):
    for k, v in q.items():
        if k == '$or':
            if not any(_cocok(d, s) for s in v):
                return False
            continue
        if k == '$and':
            if not all(_cocok(d, s) for s in v):
                return False
            continue
        nilai = d.get(k)
        if isinstance(v, dict):
            for op, x in v.items():
                if op == '$in' and not (nilai in x or (isinstance(nilai, list) and set(nilai) & set(x))):
                    return False
                if op == '$nin' and nilai in x:
                    return False
                if op == '$ne' and nilai == x:
                    return False
                if op == '$lt' and (nilai is None or nilai >= x):
                    return False
                if op == '$gte' and (nilai is None or nilai < x):
                    return False
                if op == '$lte' and (nilai is None or nilai > x):
                    return False
                if op == '$exists' and (k in d) != x:
                    return False
        elif isinstance(nilai, list) and not isinstance(v, list):
            if v not in nilai:
                return False
        elif nilai != v:
            return False
    return True


class _Cursor:
    def __init__(self, docs):
        self.docs = docs

    def sort(self, key, arah=1):
        if isinstance(key, list):
            key, arah = key[0]
        self.docs = sorted(self.docs, key=lambda d: (d.get(key) is None, d.get(key) or ''), reverse=arah == -1)
        return self

    async def to_list(self, _n):
        return [dict(d) for d in self.docs]

    def __aiter__(self):
        self._iter = iter([dict(d) for d in self.docs])
        return self

    async def __anext__(self):
        try:
            return next(self._iter)
        except StopIteration:
            raise StopAsyncIteration


class _Koleksi:
    def __init__(self, docs=()):
        self.docs = [dict(d) for d in docs]

    def find(self, q=None, _proj=None):
        return _Cursor([d for d in self.docs if _cocok(d, q or {})])

    async def find_one(self, q, _proj=None):
        hasil = [d for d in self.docs if _cocok(d, q)]
        return dict(hasil[0]) if hasil else None

    async def insert_one(self, doc):
        self.docs.append(dict(doc))

    async def update_one(self, q, u):
        for d in self.docs:
            if _cocok(d, q):
                d.update(u.get('$set', {}))
                for k, v in u.get('$push', {}).items():
                    d.setdefault(k, []).append(v)
                return


class _Db(dict):
    def __getattr__(self, nama):
        if nama not in self:
            self[nama] = _Koleksi()
        return self[nama]

    def __getitem__(self, nama):
        return self.__getattr__(nama) if nama not in self.keys() else dict.__getitem__(self, nama)


def _user(peran, uid='u1', roles=None):
    return {'id': uid, 'active_role': peran, 'roles': roles or [peran], 'full_name': f'Pengguna {peran}'}


@pytest.fixture
def db(monkeypatch):
    d = _Db(
        classes=_Koleksi([{'id': 'k7a', 'name': '7A', 'homeroom_teacher_id': 'wali1'}, {'id': 'k7b', 'name': '7B'}]),
        academic_years=_Koleksi([]),
        users=_Koleksi([
            {'id': 's1', 'full_name': 'Ani', 'nis': '1', 'roles': ['siswa'], 'student_class_id': 'k7a'},
            {'id': 's2', 'full_name': 'Budi', 'nis': '2', 'roles': ['siswa'], 'student_class_id': 'k7a'},
            {'id': 's3', 'full_name': 'Caca', 'nis': '3', 'roles': ['siswa'], 'student_class_id': 'k7b'},
            {'id': 's4', 'full_name': 'Dedi', 'nis': '4', 'roles': ['siswa'], 'student_class_id': 'k7a', 'is_active': False},
        ]),
        tatib_penanganan=_Koleksi([
            {'id': 'p1', 'siswa_id': 's1', 'poin': 10, 'jenis_poin': 'kebaikan', 'tanggal': '2026-09-01', 'siswa_kelas': '7A'},
            {'id': 'p2', 'siswa_id': 's2', 'poin': -15, 'jenis_poin': 'pelanggaran', 'tanggal': '2026-09-02', 'siswa_kelas': '7A'},
            {'id': 'p3', 'siswa_id': 's2', 'tatib_poin': -10, 'tanggal': '2026-09-03', 'siswa_kelas': '7A'},
            {'id': 'p4', 'siswa_id': 's3', 'poin': -5, 'tanggal': '2026-09-04', 'siswa_kelas': '7B'},
        ]),
    )
    monkeypatch.setattr(api, 'db', d)
    return d


def test_akses_dan_rangkuman():
    assert akses_tatib(_user('guru_tata_tertib')) == 'input'
    assert akses_tatib(_user('guru', roles=['guru', 'admin'])) == 'input'
    assert akses_tatib(_user('kepala_tata_usaha')) == 'lihat'
    assert akses_tatib(_user('guru')) is None
    r = rangkum_poin([{'poin': 10}, {'poin': -15}, {'tatib_poin': -10}])
    assert r == {'total_plus': 10, 'total_minus': -25, 'saldo': -15, 'jumlah_kebaikan': 1, 'jumlah_pelanggaran': 2, 'perlu_perhatian': True}


def test_wajib_akses_menolak_peran_lain():
    cek = api.wajib_akses('kelas')
    with pytest.raises(HTTPException) as e:
        asyncio.run(cek(_user('guru')))
    assert e.value.status_code == 403
    assert asyncio.run(cek(_user('wali_kelas')))['id'] == 'u1'


def test_daftar_siswa_walikelas_hanya_kelasnya(db):
    hasil = asyncio.run(api.daftar_siswa_walikelas(_user('wali_kelas', 'wali1')))
    assert hasil['kelas'] == {'id': 'k7a', 'name': '7A'}
    siswa = {s['id']: s for s in hasil['siswa']}
    assert set(siswa) == {'s1', 's2'}  # s3 kelas lain, s4 nonaktif
    assert siswa['s1']['saldo'] == 10 and not siswa['s1']['perlu_perhatian']
    assert siswa['s2']['total_minus'] == -25 and siswa['s2']['perlu_perhatian']


def test_walikelas_tanpa_kelas(db):
    assert asyncio.run(api.daftar_siswa_walikelas(_user('wali_kelas', 'lain'))) == {'kelas': None, 'siswa': []}


def test_batasi_query_per_peran(db):
    assert asyncio.run(api.batasi_query_tatib(_user('kepala_sekolah'), {})) == {}
    assert asyncio.run(api.batasi_query_tatib(_user('siswa', 's1'), {})) == {'siswa_id': {'$in': ['s1']}}
    assert asyncio.run(api.batasi_query_tatib(_user('wali_kelas', 'wali1'), {})) == {'siswa_id': {'$in': ['s1', 's2']}}
    with pytest.raises(HTTPException) as e:
        asyncio.run(api.batasi_query_tatib(_user('siswa', 's1'), {'siswa_id': 's2'}))
    assert e.value.status_code == 403
    with pytest.raises(HTTPException):
        asyncio.run(api.batasi_query_tatib(_user('wali_kelas', 'wali1'), {'siswa_id': 's3'}))
    with pytest.raises(HTTPException):
        asyncio.run(api.batasi_query_tatib(_user('guru'), {}))


def test_poin_siswa_walikelas(db):
    hasil = asyncio.run(api.poin_siswa_walikelas('s2', _user('wali_kelas', 'wali1')))
    assert [r['id'] for r in hasil['records']] == ['p3', 'p2']  # terbaru dulu
    assert hasil['perlu_perhatian'] is True and hasil['siswa']['kelas'] == '7A'
    assert hasil['records'][0]['poin'] == -10 and hasil['records'][0]['jenis_poin'] == 'pelanggaran'  # catatan lama
    with pytest.raises(HTTPException) as e:
        asyncio.run(api.poin_siswa_walikelas('s3', _user('wali_kelas', 'wali1')))
    assert e.value.status_code == 403
    with pytest.raises(HTTPException):
        asyncio.run(api.wajib_wali_kelas(_user('guru_bk')))


def test_poin_saya_dan_poin_siswa(db):
    saya = asyncio.run(api.poin_saya(user=_user('siswa', 's1')))
    assert saya['siswa']['nama'] == 'Ani' and saya['saldo'] == 10 and len(saya['records']) == 1
    with pytest.raises(HTTPException):
        asyncio.run(api.poin_saya(user=_user('wali_kelas', 'wali1')))
    assert asyncio.run(api.poin_siswa('s3', user=_user('waka_kurikulum')))['total_minus'] == -5
    with pytest.raises(HTTPException):
        asyncio.run(api.poin_siswa('s2', user=_user('siswa', 's1')))


def test_detail_catatan_poin(db):
    c = asyncio.run(api.detail_catatan_poin('p3', _user('wali_kelas', 'wali1')))
    assert c['poin'] == -10 and c['tindak_lanjut'] == [] and c['jenis_poin'] == 'pelanggaran'
    for uid, peran, cid, kode in (('s1', 'siswa', 'p2', 403), ('wali1', 'wali_kelas', 'p4', 403), ('x', 'admin', 'tidak-ada', 404)):
        with pytest.raises(HTTPException) as e:
            asyncio.run(api.detail_catatan_poin(cid, _user(peran, uid)))
        assert e.value.status_code == kode


def test_ambang_perlu_perhatian(db):
    wali = _user('wali_kelas', 'wali1')
    assert asyncio.run(api.lihat_ambang(wali)) == {'batas_minus_perhatian': -20}
    hasil = asyncio.run(api.siswa_perlu_perhatian(wali))
    assert [s['id'] for s in hasil['siswa']] == ['s2']
    db['settings'] = _Koleksi([{'id': 'global_config', 'tatib_batas_minus_perhatian': -30}])
    assert asyncio.run(api.siswa_perlu_perhatian(wali))['siswa'] == []
    assert asyncio.run(api.poin_siswa('s2', user=_user('admin')))['perlu_perhatian'] is False
    with pytest.raises(HTTPException):
        asyncio.run(api.wajib_akses('input')(_user('kepala_sekolah')))
    with pytest.raises(Exception):
        api.AmbangRequest(batas_minus_perhatian=5)



def test_susun_rekap_murni():
    from tatib_poin import susun_rekap
    records = [
        {'siswa_id': 's1', 'poin': 10, 'tatib_kode': 'K1', 'tatib_nama': 'Petugas'},
        {'siswa_id': 's2', 'poin': -15, 'tatib_kode': 'P1', 'tatib_nama': 'Terlambat'},
        {'siswa_id': 's2', 'poin': -10, 'tatib_kode': 'P1', 'tatib_nama': 'Terlambat'},
        {'siswa_id': 'lama', 'poin': -3, 'siswa_kelas': '9C', 'tatib_kode': 'P2', 'tatib_nama': 'Seragam'},
    ]
    r = susun_rekap(records, {'7A': ['s1', 's2', 's5'], '7B': []})
    kelas = {k['kelas']: k for k in r['per_kelas']}
    assert set(kelas) == {'7A', '7B', '9C'}
    assert kelas['7A']['jumlah_siswa'] == 3 and kelas['7A']['saldo'] == -15 and kelas['7A']['siswa_perlu_perhatian'] == 1
    assert kelas['7B']['jumlah_siswa'] == 0 and kelas['9C']['total_minus'] == -3
    assert r['ringkasan']['siswa_perlu_perhatian'] == 1 and r['ringkasan']['jumlah_pelanggaran'] == 3
    assert r['teratas_pelanggaran'][0] == {'kode': 'P1', 'nama': 'Terlambat', 'kategori_nama': None, 'jumlah': 2}
    assert [t['kode'] for t in r['teratas_kebaikan']] == ['K1']


def test_rekap_endpoint_filter(db):
    db['semesters'] = _Koleksi([{'id': 'sem1', 'start_date': '2026-09-02', 'end_date': '2026-09-03'}])
    pimpinan = _user('penjamin_mutu')
    semua = asyncio.run(api.rekap_poin(user=pimpinan))
    assert semua['ringkasan']['jumlah_pelanggaran'] == 3
    hanya_7a = asyncio.run(api.rekap_poin(kelas='7A', user=pimpinan))
    assert [k['kelas'] for k in hanya_7a['per_kelas']] == ['7A'] and hanya_7a['per_kelas'][0]['jumlah_siswa'] == 2
    tanggal = asyncio.run(api.rekap_poin(start_date='2026-09-03', end_date='2026-09-04', user=pimpinan))
    assert tanggal['ringkasan']['jumlah_pelanggaran'] == 2
    semester = asyncio.run(api.rekap_poin(semester_id='sem1', user=pimpinan))  # catatan lama: lewat rentang tanggal
    assert semester['ringkasan']['total_minus'] == -25
    with pytest.raises(HTTPException) as e:
        asyncio.run(api.rekap_poin(start_date='2026-09-05', end_date='2026-09-01', user=pimpinan))
    assert e.value.status_code == 400
    with pytest.raises(HTTPException):
        asyncio.run(api.wajib_akses('input', 'lihat')(_user('wali_kelas')))


def test_rekap_validasi_tanggal_dan_semester_lama(db):
    db['tatib_penanganan'].docs[0]['semester'] = 'ganjil'
    pimpinan = _user('kepala_sekolah')
    assert asyncio.run(api.rekap_poin(semester='ganjil', user=pimpinan))['ringkasan']['jumlah_kebaikan'] == 1
    with pytest.raises(HTTPException) as e:
        asyncio.run(api.rekap_poin(start_date='08/10/2026', user=pimpinan))
    assert e.value.status_code == 400


def test_ekspor_rekap_xlsx(db):
    import io as _io
    from openpyxl import load_workbook
    resp = asyncio.run(api.ekspor_rekap_poin(kelas='7A', user=_user('waka_kesiswaan')))
    assert 'rekap_poin_tatib_kelas-7A' in resp.headers['content-disposition']

    async def _baca():
        return b''.join([c async for c in resp.body_iterator])
    wb = load_workbook(_io.BytesIO(asyncio.run(_baca())))
    assert wb.sheetnames == ['Rekap per Kelas', 'Aturan Terbanyak']
    ws = wb['Rekap per Kelas']
    assert [c.value for c in ws[2]][:2] == ['7A', 2] and ws[3][0].value == 'TOTAL'


def test_proteksi_read_only_endpoint_tatib():
    import inspect
    from routers import tatib
    tulis = ['create_kategori', 'update_kategori', 'delete_kategori', 'create_jenis', 'update_jenis', 'delete_jenis',
             'create_aturan', 'update_aturan', 'delete_aturan', 'import_aturan_excel', 'create_penanganan', 'update_penanganan']
    for fn in tulis:
        dep = inspect.signature(getattr(tatib, fn)).parameters['user'].default.dependency
        assert dep is api.wajib_input_tatib, fn
    assert inspect.signature(tatib.get_summary_stats).parameters['user'].default.dependency is api.wajib_lihat_tatib
    for peran in ('kepala_sekolah', 'penjamin_mutu', 'kepala_tata_usaha', 'waka_kurikulum', 'guru_bk', 'wali_kelas', 'siswa'):
        with pytest.raises(HTTPException) as e:
            asyncio.run(api.wajib_input_tatib(_user(peran)))
        assert e.value.status_code == 403
    for peran in ('admin', 'guru_tata_tertib', 'waka_kesiswaan'):
        assert asyncio.run(api.wajib_input_tatib(_user(peran)))
    for peran in ('kepala_tata_usaha', 'waka_kurikulum', 'waka_kesiswaan'):
        assert asyncio.run(api.wajib_lihat_tatib(_user(peran)))



@pytest.fixture
def tatib_db(db, monkeypatch):
    from routers import tatib

    async def _audit(*_a, **_k):
        return None
    monkeypatch.setattr(tatib, 'db', db)
    monkeypatch.setattr(tatib, 'spawn', lambda coro: coro.close())
    monkeypatch.setattr(tatib, 'log_audit', _audit)
    db['tatib_aturan'] = _Koleksi([
        {'id': 'a-telat', 'kode': 'P-12', 'nama_aturan': 'Terlambat', 'poin': -2, 'jenis_poin': 'pelanggaran',
         'kategori_nama': 'Kedisiplinan',
         'kondisi': [{'id': 'ringan', 'label': '< 15 menit', 'poin': -2}, {'id': 'berat', 'label': '> 15 menit', 'poin': -5}]},
        {'id': 'a-lomba', 'kode': 'K-07', 'nama_aturan': 'Juara lomba', 'poin': 20},  # aturan lama tanpa kondisi
        {'id': 'a-mati', 'kode': 'P-99', 'nama_aturan': 'Lama', 'poin': -1, 'is_active': False},
    ])
    return tatib


def _req(tatib, **kw):
    dasar = {'siswa_id': 's1', 'tatib_id': 'a-telat', 'tanggal': '2026-10-08'}
    return tatib.PenangananRequest(**{**dasar, **kw})


def test_catat_poin_nilai_dari_kondisi(tatib_db):
    t = tatib_db
    guru = _user('guru_tata_tertib', 'g1')
    doc = asyncio.run(t.catat_poin_jalur('pelanggaran', _req(t, kondisi_id='berat', poin=999), guru))
    assert doc['poin'] == -5 and doc['kondisi'] == '> 15 menit' and doc['jenis_poin'] == 'pelanggaran'
    assert doc['siswa_nama'] == 'Ani' and doc['siswa_kelas'] == '7A' and doc['tindak_lanjut'] == []
    kebaikan = asyncio.run(t.catat_poin_jalur('kebaikan', _req(t, tatib_id='a-lomba'), guru))
    assert kebaikan['poin'] == 20 and kebaikan['kondisi_id'] == 'umum'
    for jalur, kw, kode in (
        ('kebaikan', {}, 400),                                  # aturan pelanggaran di jalur kebaikan
        ('pelanggaran', {'kondisi_id': None}, 400),             # kondisi wajib dipilih (>1 kondisi)
        ('pelanggaran', {'kondisi_id': 'tak-ada'}, 400),
        ('pelanggaran', {'tatib_id': 'a-mati'}, 400),           # aturan nonaktif
        ('pelanggaran', {'siswa_id': 'tidak-ada', 'kondisi_id': 'ringan'}, 404),
        ('prestasi', {'kondisi_id': 'ringan'}, 404),
    ):
        with pytest.raises(HTTPException) as e:
            asyncio.run(t.catat_poin_jalur(jalur, _req(t, **kw), guru))
        assert e.value.status_code == kode, (jalur, kw)


def test_tindak_lanjut_hanya_pelanggaran(tatib_db):
    t = tatib_db
    guru = _user('waka_kesiswaan', 'w1')
    hasil = asyncio.run(t.tambah_tindak_lanjut('p2', t.TindakLanjutRequest(tanggal='2026-10-08', uraian=' Panggil orang tua '), guru))
    assert hasil['tindak_lanjut'][0]['uraian'] == 'Panggil orang tua' and hasil['tindak_lanjut'][0]['petugas_nama']
    for cid, uraian, kode in (('p1', 'x', 400), ('p2', '  ', 400), ('nihil', 'x', 404)):
        with pytest.raises(HTTPException) as e:
            asyncio.run(t.tambah_tindak_lanjut(cid, t.TindakLanjutRequest(tanggal='2026-10-08', uraian=uraian), guru))
        assert e.value.status_code == kode


def test_aturan_per_jalur_selalu_berkondisi(tatib_db):
    t = tatib_db
    pel = asyncio.run(t.list_aturan(jenis_poin='pelanggaran', user=_user('admin')))
    assert {a['id'] for a in pel} == {'a-telat', 'a-mati'}
    keb = asyncio.run(t.list_aturan(jenis_poin='kebaikan', user=_user('admin')))
    assert [a['id'] for a in keb] == ['a-lomba']
    assert keb[0]['kondisi'] == [{'id': 'umum', 'label': 'Setiap kejadian', 'poin': 20}] and keb[0]['jenis_poin'] == 'kebaikan'


def test_skema_aturan_dari_request(tatib_db):
    t = tatib_db
    req = t.TatibRequest(kategori_id='k', jenis_id='j', kode='P-1', nama_aturan='X', poin=3, jenis_poin='pelanggaran',
                         kondisi=[{'label': 'Pertama', 'poin': 4}, {'label': 'Ulang', 'poin': -8}])
    sk = t._skema_aturan(req)
    assert sk['poin'] == -4 and [k['poin'] for k in sk['kondisi']] == [-4, -8]
    with pytest.raises(HTTPException):
        t._skema_aturan(t.TatibRequest(kategori_id='k', jenis_id='j', kode='P', nama_aturan='X', poin=1, jenis_poin='prestasi'))
