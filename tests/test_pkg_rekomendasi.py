"""Rekomendasi PKB: opsi baku, teks cetak, dan API simpan oleh kepala/penilai."""

import os
import sys

import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'backend'))

import pkg_rekomendasi as r  # noqa: E402
from routers import pkg_angket_api, pkg_rekomendasi_api as api  # noqa: E402
from tests.pkg_palsu import Db, Koleksi  # noqa: E402


def test_teks_dan_saran():
    assert r.teks_rekomendasi(None) == ''
    teks = r.teks_rekomendasi({'pilihan_pkb': ['ptk', 'diklat'], 'fokus_kompetensi': ['pedagogik'], 'catatan': 'Prioritas semester depan'})
    assert teks.startswith(r.OPSI_PKB['diklat'])  # urutan mengikuti daftar opsi baku
    assert '(fokus kompetensi pedagogik)' in teks and teks.endswith('Prioritas semester depan')
    assert r.teks_rekomendasi({'catatan': 'Bebas'}) == 'Bebas'
    assert r.saran_pkb(5)[0] == 'pengimbasan' and r.saran_pkb(None) == []


@pytest.fixture
def klien(monkeypatch):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from core import get_current_user
    d = Db({
        'pkg_periode': Koleksi([{'id': 'p1', 'status': 'berjalan', 'nama_periode': 'Ganjil'}]),
        'pkg_sk': Koleksi([{'id': 's1', 'periode_id': 'p1', 'is_active': True, 'penilai_ids': ['pen'], 'dinilai_ids': ['g1', 'g2'],
                            'snapshot_anggota': {'g1': {'nama': 'Guru Satu'}, 'g2': {'nama': 'Guru Dua'}}}]),
        'pkg_penugasan': Koleksi([{'periode_id': 'p1', 'penilai_id': 'pen', 'guru_dinilai_ids': ['g1']}]),
        'pkg_penilaian': Koleksi([]),
        'users': Koleksi([]),
    })
    for m in (api, pkg_angket_api):
        monkeypatch.setattr(m, 'db', d)
    app = FastAPI()
    app.include_router(api.router, prefix='/api')
    pengguna = {'user': {'id': 'kep', 'roles': ['kepala_sekolah'], 'full_name': 'Kepala'}}
    app.dependency_overrides[get_current_user] = lambda: pengguna['user']
    return TestClient(app), d, pengguna


def test_simpan_rekomendasi(klien):
    c, d, pengguna = klien
    assert len(c.get('/api/pkg/rekomendasi-pkb/opsi').json()['opsi']) == len(r.OPSI_PKB)
    assert c.get('/api/pkg/rekomendasi-pkb/g1').json()['rekomendasi'] is None
    assert c.put('/api/pkg/rekomendasi-pkb/g1', json={}).status_code == 422  # kosong
    assert c.put('/api/pkg/rekomendasi-pkb/g1', json={'pilihan_pkb': ['liburan']}).status_code == 422
    hasil = c.put('/api/pkg/rekomendasi-pkb/g1', json={'pilihan_pkb': ['kolektif'], 'fokus_kompetensi': ['sosial'], 'catatan': ' Ikut MGMP '})
    assert hasil.status_code == 200 and hasil.json()['peran_penyusun'] == 'kepala' and hasil.json()['catatan'] == 'Ikut MGMP'
    pengguna['user'] = {'id': 'pen', 'roles': ['guru'], 'full_name': 'Penilai'}
    ubah = c.put('/api/pkg/rekomendasi-pkb/g1', json={'catatan': 'Catatan penilai'}).json()
    assert ubah['peran_penyusun'] == 'penilai' and ubah['pilihan_pkb'] == [] and len(d['pkg_rekomendasi_pkb'].docs) == 1
    assert c.put('/api/pkg/rekomendasi-pkb/g2', json={'catatan': 'x'}).status_code == 403
    ringkas = {g['guru_id']: g for g in c.get('/api/pkg/rekomendasi-pkb/ringkasan').json()['guru']}
    assert ringkas['g1']['rekomendasi']['teks'] == 'Catatan penilai' and 'g2' not in ringkas
    assert [a['aksi'] for a in d['pkg_audit_log'].docs].count('rekomendasi_simpan') == 2


def test_rekomendasi_di_lembar_cetak():
    import io
    from openpyxl import load_workbook
    import pkg_cetak
    import pkg_template
    teks = r.teks_rekomendasi({'pilihan_pkb': ['kolektif', 'ptk'], 'catatan': 'Didampingi kepala madrasah.'})
    wb = load_workbook(io.BytesIO(pkg_template.isi_template_ke_bytes('instrumen', lambda w: pkg_cetak.isi_rekomendasi(w, teks))))
    ws = wb['01-LAPORN']
    assert (ws['C21'].value, ws['D21'].value, ws['E21'].value) == ('Rekomendasi PKB', ':', teks)
    assert 'E21:H21' in [str(m) for m in ws.merged_cells.ranges] and ws.row_dimensions[21].height > 15
    assert ws['E23'].value == 'PERSETUJUAN' and ws['E20'].value.startswith('=IF(F19')  # baris lain tidak bergeser
    rekap = load_workbook(io.BytesIO(pkg_cetak.rekap_excel(
        [{'nama': 'A', 'capaian_persen': 80, 'rekomendasi': teks}, {'nama': 'B', 'capaian_persen': 70}], {'tanggal_tutup': '2026-12-01'})))
    ws = rekap['Rekap hasil PKG']
    assert ws['H6'].value == teks and ws['H6'].alignment.wrap_text and ws['H7'].value is None
