"""Angket 360°: butir dari template resmi, berkas cetak FORM-360-GR, dan API cetak angket (otorisasi & riwayat)."""

import io
import os
import sys

import pytest
from openpyxl import load_workbook

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'backend'))

import pkg_angket  # noqa: E402
from routers import pkg_angket_api as api  # noqa: E402
from routers import pkg_cetak_api  # noqa: E402
from tests.pkg_palsu import Db, Koleksi  # noqa: E402


def test_butir_sesuai_sheet_template():
    jumlah = {j['jenis']: (j['jumlah_butir'], j['skor_maks']) for j in pkg_angket.daftar_jenis()}
    assert jumlah == {'sejawat': (28, 56), 'siswa': (40, 80), 'ortu': (8, 16)}
    sejawat = pkg_angket.butir_angket('sejawat')
    assert sejawat[0]['komponen'] == 'Perilaku Guru sehari-hari' and sejawat[0]['pernyataan'].startswith('Guru mentaati')
    assert sejawat[10]['komponen'] == 'Hubungan Guru dengan Teman Sejawat' and sejawat[10]['nomor'] == 1
    assert [b['kode'] for b in sejawat[:2]] == ['sejawat.1', 'sejawat.2'] and sejawat[-1]['baris'] == 36


def test_berkas_angket_hanya_bagian_form():
    wb = load_workbook(io.BytesIO(pkg_angket.angket_excel('siswa', {'nama': 'Budi', 'nip': '123', 'tugas_tambahan': None})))
    ws = wb['FORM-360-GR']
    assert (ws['E56'].value, ws['E57'].value, ws['E58'].value) == (': Budi', ': 123', ': -')
    assert ws['E6'].value.startswith(': …')  # bagian lain tidak diubah
    assert ws.print_area == "'FORM-360-GR'!$B$52:$H$124" and wb.active.title == 'FORM-360-GR'
    lain = [w for w in wb.worksheets if w.title != 'FORM-360-GR']
    assert all(w.sheet_state == 'hidden' for w in lain)
    assert all(not w.print_area or w.print_area.endswith('!') for w in lain)  # area cetak sheet tersembunyi dihapus
    kosong = load_workbook(io.BytesIO(pkg_angket.angket_excel('ortu')))['FORM-360-GR']
    assert kosong['E159'].value.startswith(': …')


@pytest.fixture
def klien(monkeypatch):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from core import get_current_user
    d = Db({
        'pkg_periode': Koleksi([{'id': 'p1', 'status': 'berjalan', 'nama_periode': 'Ganjil', 'tanggal_buka': '2026-10-01'}]),
        'pkg_sk': Koleksi([{'id': 's1', 'periode_id': 'p1', 'is_active': True, 'penilai_ids': ['pen'], 'dinilai_ids': ['g1', 'g2'],
                            'snapshot_anggota': {'g1': {'nama': 'Guru Satu', 'nip': 'NIP-g1'}}}]),
        'pkg_penugasan': Koleksi([{'periode_id': 'p1', 'penilai_id': 'pen', 'guru_dinilai_ids': ['g1']}]),
        'pkg_penilaian': Koleksi([]),
        'users': Koleksi([{'id': 'g1', 'full_name': 'Guru Satu', 'tugas_tambahan': 'Wali Kelas 7A'}]),
    })
    monkeypatch.setattr(api, 'db', d)
    monkeypatch.setattr(pkg_cetak_api, 'kemampuan_pdf', lambda: {'libreoffice': False, 'mesin': 'reportlab', 'identik': False})
    app = FastAPI()
    app.include_router(api.router, prefix='/api')
    pengguna = {'user': {'id': 'kep', 'roles': ['kepala_sekolah'], 'full_name': 'Kepala'}}
    app.dependency_overrides[get_current_user] = lambda: pengguna['user']
    return TestClient(app), d, pengguna


def test_cetak_angket_excel_dan_pdf(klien):
    k, d, _ = klien
    r = k.post('/api/pkg/angket/cetak', json={'jenis_angket': 'sejawat', 'format': 'xlsx', 'guru_id': 'g1'})
    assert r.status_code == 200 and 'Angket_360_Teman_Sejawat_Guru_Satu' in r.headers['content-disposition']
    ws = load_workbook(io.BytesIO(r.content))['FORM-360-GR']
    assert (ws['E6'].value, ws['E8'].value) == (': Guru Satu', ': Wali Kelas 7A')
    u = d['pkg_unduhan'].docs[0]
    assert (u['jenis'], u['guru_id'], u['format']) == ('angket', 'g1', 'xlsx')
    assert d['pkg_audit_log'].docs[-1]['metadata']['jenis_angket'] == 'sejawat'
    pdf = k.post('/api/pkg/angket/cetak', json={'jenis_angket': 'ortu'})  # bawaan PDF, angket kosong
    assert pdf.status_code == 200 and pdf.content.startswith(b'%PDF')
    assert pdf.headers['x-pkg-mesin'] == 'reportlab' and pdf.headers['x-pkg-cadangan'] == 'libreoffice_tidak_tersedia'


def test_cetak_angket_otorisasi(klien):
    k, d, pengguna = klien
    pengguna['user'] = {'id': 'pen', 'roles': ['guru']}
    assert k.post('/api/pkg/angket/cetak', json={'jenis_angket': 'siswa', 'format': 'xlsx', 'guru_id': 'g1'}).status_code == 200
    assert k.post('/api/pkg/angket/cetak', json={'jenis_angket': 'siswa', 'guru_id': 'g2'}).status_code == 403  # bukan guru yang dinilainya
    pengguna['user'] = {'id': 'g2', 'roles': ['guru']}
    r = k.post('/api/pkg/angket/cetak', json={'jenis_angket': 'siswa'})
    assert r.status_code == 403 and r.json()['detail']['alasan'] == 'bukan_penilai'
    assert k.post('/api/pkg/angket/cetak', json={'jenis_angket': 'dudi'}).status_code == 422  # DUDI tidak dipakai
    assert any(a['aksi'] == 'cetak_gagal' for a in d['pkg_audit_log'].docs)
