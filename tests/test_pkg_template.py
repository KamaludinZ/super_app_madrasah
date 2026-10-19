"""Template PKG: berkas resmi tersedia di folder internal, path tetap, dan peringatan startup bila hilang."""

import logging
import tempfile
import os
import sys

import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'backend'))

import pkg_template  # noqa: E402


def test_berkas_template_resmi_tersedia():
    hasil = pkg_template.periksa_template()
    assert hasil['lengkap'] is True and hasil['versi'] == '2026'
    assert {b['nama_berkas'] for b in hasil['berkas']} == {'pkg_template_v2026.xlsx', 'rekap_hasil_pkg_v2026.xlsx'}
    assert all(b['ukuran'] > 0 for b in hasil['berkas'])
    assert pkg_template.path_template('rekap').startswith(pkg_template.FOLDER_TEMPLATE)
    with pytest.raises(ValueError):
        pkg_template.path_template('../rahasia')


def test_peringatan_bila_hilang(monkeypatch, tmp_path, caplog):
    monkeypatch.setattr(pkg_template, 'FOLDER_TEMPLATE', str(tmp_path))
    (tmp_path / 'pkg_template_v2026.xlsx').write_bytes(b'xlsx')
    with caplog.at_level(logging.WARNING, logger='matsandatama'):
        hasil = pkg_template.periksa_template()
    assert hasil['lengkap'] is False
    assert [b['ada'] for b in hasil['berkas']] == [True, False]
    assert 'rekap_hasil_pkg_v2026.xlsx tidak ditemukan' in caplog.text


def test_checksum_resmi_cocok_dan_terdeteksi_berubah(monkeypatch, tmp_path, caplog):
    hasil = pkg_template.periksa_template()
    assert hasil['utuh'] is True and all(b['status'] == 'cocok' for b in hasil['berkas'])
    instrumen = pkg_template.info_berkas('instrumen')
    assert instrumen['checksum'] == pkg_template.TEMPLATE_PKG['instrumen']['checksum_resmi'] and instrumen['versi'] == '2026'

    monkeypatch.setattr(pkg_template, 'FOLDER_TEMPLATE', str(tmp_path))
    (tmp_path / 'pkg_template_v2026.xlsx').write_bytes(b'berkas diubah')
    with caplog.at_level(logging.WARNING, logger='matsandatama'):
        hasil = pkg_template.periksa_template()
    status = {b['jenis']: b['status'] for b in hasil['berkas']}
    assert status == {'instrumen': 'berbeda', 'rekap': 'hilang'} and hasil['utuh'] is False
    assert 'berbeda dari checksum resmi' in caplog.text
    assert pkg_template.hitung_checksum(str(tmp_path / 'tidak-ada.xlsx')) is None


def test_pemeriksaan_startup_tersimpan():
    import asyncio
    from tests.pkg_palsu import Db
    db = Db()
    hasil = asyncio.run(pkg_template.periksa_template_startup(db))
    assert hasil['utuh'] is True and pkg_template.HASIL_TERAKHIR['diperiksa_pada'] == hasil['diperiksa_pada']
    dok = {d['jenis']: d for d in db['pkg_template'].docs}
    assert set(dok) == {'instrumen', 'rekap'} and dok['rekap']['status'] == 'cocok' and dok['rekap']['diperiksa_pada']
    asyncio.run(pkg_template.periksa_template_startup(db))  # idempoten: tetap satu dokumen per jenis
    assert len(db['pkg_template'].docs) == 2


def test_pemeriksaan_startup_tetap_jalan_bila_db_gagal():
    import asyncio

    class DbRusak(dict):
        def __getitem__(self, _):
            raise RuntimeError('db mati')
    assert asyncio.run(pkg_template.periksa_template_startup(DbRusak()))['lengkap'] is True


def test_baca_template_tanpa_menimpa(tmp_path, monkeypatch):
    path_asli = pkg_template.path_template('rekap')
    sebelum = pkg_template.hitung_checksum(path_asli)
    mtime = os.path.getmtime(path_asli)

    wb = pkg_template.buka_template('rekap')
    assert wb.sheetnames  # bisa dibaca
    wb.close()

    def pengisi(wb):
        wb.active['A1'] = 'DIUBAH UJI'
    hasil = pkg_template.isi_template_ke_bytes('rekap', pengisi)
    assert hasil[:2] == b'PK'  # berkas xlsx (zip) baru
    assert pkg_template.hitung_checksum(path_asli) == sebelum and os.path.getmtime(path_asli) == mtime

    with pkg_template.salinan_sementara(hasil, '../../keluar.xlsx') as p:
        folder = os.path.dirname(p)
        assert os.path.basename(p) == 'keluar.xlsx' and os.path.isfile(p) and folder.startswith(tempfile.gettempdir())
    assert not os.path.exists(folder)  # folder sementara dibersihkan


def test_template_tidak_tersedia_atau_berubah(tmp_path, monkeypatch):
    monkeypatch.setattr(pkg_template, 'FOLDER_TEMPLATE', str(tmp_path))
    with pytest.raises(pkg_template.TemplateTidakTersedia):
        pkg_template.baca_bytes_template('instrumen')
    (tmp_path / 'pkg_template_v2026.xlsx').write_bytes(b'bukan resmi')
    assert pkg_template.baca_bytes_template('instrumen') == b'bukan resmi'
    with pytest.raises(pkg_template.TemplateTidakTersedia):
        pkg_template.baca_bytes_template('instrumen', wajib_utuh=True)


def test_cetak_ke_salinan_sementara_dan_pembersihan(tmp_path, monkeypatch):
    monkeypatch.setenv('PKG_FOLDER_SEMENTARA', str(tmp_path))
    asli = pkg_template.hitung_checksum(pkg_template.path_template('rekap'))

    def pengisi(wb):
        wb.active['A1'] = 'REKAP UJI'
    with pkg_template.cetak_ke_salinan_sementara('rekap', pengisi, 'rekap periode.xlsx') as path:
        assert path.startswith(str(tmp_path)) and os.path.isfile(path)
        from openpyxl import load_workbook
        assert load_workbook(path).active['A1'].value == 'REKAP UJI'
    assert os.listdir(tmp_path) == []
    assert pkg_template.hitung_checksum(pkg_template.path_template('rekap')) == asli

    assert pkg_template.nama_berkas_cetak('Rekap PKG Ganjil 2026/2027', 'PDF') == 'Rekap_PKG_Ganjil_2026_2027.pdf'
    assert pkg_template.nama_berkas_cetak('../../', 'xlsx') == 'pkg.xlsx'

    lama = tmp_path / 'pkg_tertinggal'
    lama.mkdir()
    baru = tmp_path / 'pkg_berjalan'
    baru.mkdir()
    lain = tmp_path / 'bukan_pkg'
    lain.mkdir()
    os.utime(lama, (0, 0))
    os.utime(lain, (0, 0))
    assert pkg_template.bersihkan_sementara_lama() == 1
    assert sorted(os.listdir(tmp_path)) == ['bukan_pkg', 'pkg_berjalan']


def test_endpoint_info_template(monkeypatch):
    import asyncio
    import pkg_penjaga
    from routers import pkg as api
    from tests.pkg_palsu import Db, Koleksi, Request

    d = Db({'pkg_periode': Koleksi([{'id': 'p1', 'status': 'berjalan'}]),
            'pkg_sk': Koleksi([{'id': 's1', 'periode_id': 'p1', 'is_active': True, 'penilai_ids': [], 'dinilai_ids': ['g1']}])})
    monkeypatch.setattr(pkg_penjaga, 'db', d)
    monkeypatch.setattr(api, 'db', d)
    pkg_template.HASIL_TERAKHIR.clear()

    def panggil(user, ulang=False):
        async def jalan():
            akses = await pkg_penjaga.penjaga_pkg('anggota')(Request(), user)
            return await api.info_template_api(Request('/pkg/template/info'), ulang, akses)
        return asyncio.run(jalan())

    guru = {'id': 'g1', 'roles': ['guru']}
    info = panggil(guru)
    assert info['versi_template'] == '2026' and info['utuh'] is True and 'folder' not in info
    assert [b['status'] for b in info['berkas']] == ['cocok', 'cocok'] and 'ada' not in info['berkas'][0]
    pertama = info['diperiksa_pada']
    assert panggil(guru, ulang=True)['diperiksa_pada'] == pertama  # guru tidak dapat memicu periksa ulang
    assert panggil({'id': 'kep', 'roles': ['kepala_sekolah']}, ulang=True)['diperiksa_pada'] >= pertama
    with pytest.raises(Exception):
        panggil({'id': 'asing', 'roles': ['guru']})

    aksi = [(x['aksi'], x.get('jumlah')) for x in d['pkg_audit_log'].docs if x['aksi'] != 'tolak_akses']
    assert ('lihat_template', 2) in aksi and ('periksa_template', None) in aksi  # lihat dicatat harian (2x guru)


def test_catat_cetak_template():
    import asyncio
    import pkg_audit
    from tests.pkg_palsu import Db, Request
    d = Db()
    info = pkg_template.info_berkas('rekap')
    asyncio.run(pkg_audit.catat_cetak_template(
        d, {'id': 'kep', 'full_name': 'Kepala'}, 'rekap', 'pdf', 'Rekap_PKG.pdf', info['versi'], info['checksum'],
        periode_id='p1', metadata={'mesin': 'libreoffice'}, request=Request('/pkg/cetak')))
    audit = d['pkg_audit_log'].docs[0]
    assert audit['aksi'] == 'cetak_template' and audit['metadata']['mesin'] == 'libreoffice'
    assert audit['metadata']['checksum_template'] == info['checksum'] and audit['ip_address'] == '10.0.0.1'
    unduh = d['pkg_unduhan'].docs[0]
    assert (unduh['jenis'], unduh['format'], unduh['nama_berkas'], unduh['aktor_nama'], unduh['periode_id']) == \
        ('rekap', 'pdf', 'Rekap_PKG.pdf', 'Kepala', 'p1')


def test_konfigurasi_docker_menyertakan_template():
    akar = os.path.join(os.path.dirname(__file__), '..')
    for berkas, pola in (('backend/.dockerignore', '!templates/*.xlsx'), ('.dockerignore', '!backend/templates/*.xlsx')):
        with open(os.path.join(akar, berkas), encoding='utf-8') as f:
            assert pola in f.read().splitlines(), berkas
    for berkas in ('backend/Dockerfile', 'Dockerfile'):
        with open(os.path.join(akar, berkas), encoding='utf-8') as f:
            isi = f.read()
        assert 'pkg_template as t' in isi and "sys.exit(0 if h['utuh'] else 1)" in isi, berkas


def test_sheet_wajib_template(tmp_path, monkeypatch, caplog):
    assert pkg_template.SHEET_KOMPETENSI == ('KOMP-1', 'KOMP (2)', 'KOMP (3)', 'KOMP (4)')
    assert pkg_template.sheet_hilang('instrumen') == [] and pkg_template.sheet_hilang('rekap') == []
    assert pkg_template.info_berkas('instrumen')['sheet_hilang'] == []

    from openpyxl import Workbook
    monkeypatch.setattr(pkg_template, 'FOLDER_TEMPLATE', str(tmp_path))
    wb = Workbook()
    wb.active.title = 'KOMP-1'
    wb.save(tmp_path / 'pkg_template_v2026.xlsx')
    (tmp_path / 'rekap_hasil_pkg_v2026.xlsx').write_bytes(b'bukan xlsx')
    with caplog.at_level(logging.WARNING, logger='matsandatama'):
        hasil = {b['jenis']: b for b in pkg_template.periksa_template()['berkas']}
    assert hasil['instrumen']['sheet_hilang'] == ['KOMP (2)', 'KOMP (3)', 'KOMP (4)', '04-RKAP-PKG']
    assert hasil['rekap']['sheet_hilang'] is None
    assert 'tidak memuat sheet wajib' in caplog.text


def test_peta_cetak_cocok_dengan_template_resmi():
    assert pkg_template.periksa_struktur_cetak('rekap') == []
    assert pkg_template.periksa_struktur_cetak('instrumen') == []
    assert pkg_template.info_berkas('rekap')['struktur_cetak_salah'] == []


def test_struktur_cetak_berubah_terdeteksi(tmp_path, monkeypatch, caplog):
    from openpyxl import load_workbook
    wb = load_workbook(pkg_template.path_template('rekap'))
    wb['Rekap hasil PKG']['F5'] = 'Kolom lain'
    wb['Rekap hasil PKG'].insert_rows(3)  # geser seluruh tata letak
    import io as _io
    buf = _io.BytesIO()
    wb.save(buf)
    salah = pkg_template.periksa_struktur_cetak('rekap', buf.getvalue())
    assert salah and any('A5' in x for x in salah)
    monkeypatch.setattr(pkg_template, 'FOLDER_TEMPLATE', str(tmp_path))
    (tmp_path / 'rekap_hasil_pkg_v2026.xlsx').write_bytes(buf.getvalue())
    with caplog.at_level(logging.WARNING, logger='matsandatama'):
        b = pkg_template.info_berkas('rekap')
        pkg_template.periksa_template()
    assert b['status'] == 'berbeda' and b['struktur_cetak_salah']
    assert 'Tata letak berkas template rekap_hasil_pkg_v2026.xlsx berubah' in caplog.text


def test_riwayat_unduhan_lengkap():
    import asyncio
    import hashlib
    import pkg_audit
    from tests.pkg_palsu import Db, Request
    d = Db()
    isi = b'%PDF-1.7 isi'
    asyncio.run(pkg_audit.catat_cetak_template(
        d, {'id': 'pen', 'full_name': 'Penilai', 'active_role': 'guru'}, 'instrumen', 'pdf', 'Laporan_PKG_Budi.pdf', '2026',
        'abc', periode_id='p1', request=Request('/pkg/cetak'), jenis_dokumen='guru', guru_id='g1', mesin='reportlab', isi=isi))
    u = d['pkg_unduhan'].docs[0]
    assert (u['jenis'], u['jenis_template'], u['guru_id'], u['mesin'], u['ukuran']) == ('guru', 'instrumen', 'g1', 'reportlab', len(isi))
    assert u['checksum_berkas'] == hashlib.sha256(isi).hexdigest() and u['aktor_peran'] == 'guru' and u['ip_address'] == '10.0.0.1'
    meta = d['pkg_audit_log'].docs[0]['metadata']
    assert meta['mesin'] == 'reportlab' and meta['jenis_dokumen'] == 'guru' and meta['checksum_berkas'] == u['checksum_berkas']
