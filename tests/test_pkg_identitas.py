"""Lampiran identitas guru & tim penilai: sumber berlapis (snapshot periode, GTK, pengaturan) dan isian 00-DATA."""

import asyncio
import io
import os
import sys

from openpyxl import load_workbook

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'backend'))

import pkg_identitas as idn  # noqa: E402
import pkg_template  # noqa: E402
from tests.pkg_palsu import Db, Koleksi  # noqa: E402

PERIODE = {'id': 'p1', 'nama_periode': 'Ganjil', 'tanggal_buka': '2026-07-01', 'tanggal_tutup': '2026-12-31'}
SK = {'id': 's1', 'nomor_sk': 'B-12/2026', 'tanggal_sk': '2026-06-20', 'tentang': 'SK Tim PKG', 'penilai_ids': ['kep', 'pen'],
      'snapshot_anggota': {'g1': {'nama': 'Guru Lama', 'nip': 'NIP-lama', 'mapel': 'IPA', 'birth_place': 'Malang'},
                           'pen': {'nama': 'Penilai Satu', 'nip': 'NIP-pen'}}}
PENGATURAN = {'school_name': 'MTsN 2 Kota Malang', 'phone': '0341-1', 'address': 'Jl. Contoh 1', 'headmaster_name': 'Drs. Kepala',
              'leadership': [{'position': 'Kepala Madrasah', 'nip': 'NIP-kep'}]}


def _db():
    return Db({'users': Koleksi([
        {'id': 'g1', 'full_name': 'Guru Baru', 'birth_place': 'Batu', 'birth_date': '1980-04-29', 'gender': 'P', 'nuptk': '123',
         'nrg': '456', 'tmt_guru': '1999-04-01'},
        {'id': 'kep', 'full_name': 'Kepala Akun', 'roles': ['kepala_sekolah'], 'gender': 'L'},
        {'id': 'pen', 'full_name': 'Penilai Baru'}])})


def test_lampiran_identitas_berlapis():
    lamp = asyncio.run(idn.lampiran_identitas(_db(), PERIODE, SK, 'g1', 'pen', PENGATURAN, '2026-11-02T03:00:00+00:00'))
    g = lamp['guru']
    assert (g['nama'], g['nip'], g['mapel']) == ('Guru Lama', 'NIP-lama', 'IPA')  # snapshot periode diutamakan
    assert g['tempat_tanggal_lahir'] == 'Malang, 29 April 1980' and g['nuptk'] == '123 / 456'  # GTK melengkapi
    assert g['jenis_kelamin'] == 'Perempuan' and g['masa_kerja'] == '27 Tahun 8 Bulan'
    assert (lamp['kepala']['nama'], lamp['kepala']['nip'], lamp['kepala']['jenis_kelamin']) == ('Drs. Kepala', 'NIP-kep', 'Laki-laki')
    assert lamp['sekolah']['kota'] == 'Kota Malang' and lamp['sk']['tanggal'] == '20 Juni 2026'
    assert lamp['periode']['tanggal_penilaian'] == '2 November 2026'
    assert [(t['id'], t['peran'], t['penilai_lembar']) for t in lamp['tim_penilai']] == \
        [('pen', 'Penilai', True), ('kep', 'Kepala madrasah (penilai)', False)]


def test_isi_00_data():
    lamp = asyncio.run(idn.lampiran_identitas(_db(), PERIODE, SK, 'g1', 'pen', PENGATURAN))
    wb = load_workbook(io.BytesIO(pkg_template.isi_template_ke_bytes('instrumen', lambda w: idn.isi_data_identitas(w, lamp))))
    ws = wb['00-DATA']
    assert (ws['E5'].value, ws['E8'].value, ws['E9'].value, ws['F9'].value) == ('Guru Lama', 'Malang, 29 April 1980', '-', '-')
    assert (ws['E17'].value, ws['E19'].value, ws['E26'].value, ws['E28'].value) == ('Drs. Kepala', 'Kepala Madrasah', 'MTsN 2 Kota Malang', 'Jl. Contoh 1')
    assert (ws['E30'].value, ws['E36'].value, ws['C41'].value, ws['M2'].value) == ('Kota Malang', 'B-12/2026', '1 Juli 2026', 'MTs')
    assert ws['E11'].value == '27 Tahun 8 Bulan' and ws['H41'].value == '-'
    assert wb['02-IDNTY']['E6'].value == "='00-DATA'!E5&\"\""  # lembar identitas tetap membaca 00-DATA
