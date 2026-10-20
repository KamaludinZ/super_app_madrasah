"""Komponen kehadiran (INDEK-HDIR): rumus sesuai template dan API input rekap manual per guru."""

import os
import sys

import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'backend'))

import pkg_kehadiran as k  # noqa: E402
from routers import pkg_kehadiran_api as api  # noqa: E402
from routers import pkg_angket_api  # noqa: E402
from tests.pkg_palsu import Db, Koleksi  # noqa: E402

PERIODE = {'id': 'p1', 'status': 'berjalan', 'nama_periode': 'Ganjil', 'tanggal_buka': '2026-07-15', 'tanggal_tutup': '2026-12-20'}


def test_rumus_indek_hdir():
    assert k.bulan_periode(PERIODE) == ['2026-07', '2026-08', '2026-09', '2026-10', '2026-11', '2026-12']
    assert len(k.bulan_periode({'tanggal_buka': '2026-01-01', 'tanggal_tutup': '2027-06-30'})) == 12  # 12 baris sheet
    b = k.hitung_bulan({'terlambat_menit': 300, 'pulang_cepat_menit': 180, 'tanpa_keterangan': 2})
    assert (b['jumlah_menit'], b['konversi_hari'], b['jumlah_tidak_hadir']) == (480, 1.0, 3.0)  # F, G = F/480, I = G+H
    assert abs(b['persen_hadir'] - (1 - 3 / 46)) < 1e-12  # K = 1 - I/46
    kosong = k.hitung_kehadiran(k.baris_lengkap(k.bulan_periode(PERIODE), None))
    assert kosong['indeks'] == 1 and kosong['indeks_persen'] == 100.0  # template kosong -> 100%
    h = k.hitung_kehadiran(k.baris_lengkap(k.bulan_periode(PERIODE), [{'bulan': '2026-08', 'tanpa_keterangan': 1, 'izin': 2},
                                                    {'bulan': '2026-09', 'terlambat_menit': 240}]))
    assert h['jumlah']['jumlah_tidak_hadir'] == 1.5 and h['jumlah']['izin'] == 2  # izin tidak masuk rumus
    assert h['indeks_persen'] == round((1 - 1.5 / 46) * 100, 2)


@pytest.fixture
def klien(monkeypatch):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from core import get_current_user
    d = Db({
        'pkg_periode': Koleksi([dict(PERIODE)]),
        'pkg_sk': Koleksi([{'id': 's1', 'periode_id': 'p1', 'is_active': True, 'penilai_ids': ['pen'], 'dinilai_ids': ['g1', 'g2'],
                            'snapshot_anggota': {'g1': {'nama': 'Guru Satu'}, 'g2': {'nama': 'Guru Dua'}}}]),
        'pkg_penugasan': Koleksi([{'periode_id': 'p1', 'penilai_id': 'pen', 'guru_dinilai_ids': ['g1']}]),
        'pkg_penilaian': Koleksi([]),
        'users': Koleksi([]),
    })
    monkeypatch.setattr(api, 'db', d)
    monkeypatch.setattr(pkg_angket_api, 'db', d)
    app = FastAPI()
    app.include_router(api.router, prefix='/api')
    pengguna = {'user': {'id': 'kep', 'roles': ['kepala_sekolah'], 'full_name': 'Kepala'}}
    app.dependency_overrides[get_current_user] = lambda: pengguna['user']
    return TestClient(app), d, pengguna


def test_input_kehadiran(klien):
    c, d, pengguna = klien
    awal = c.get('/api/pkg/kehadiran/g1').json()
    assert awal['terisi'] is False and len(awal['bulan']) == 6 and awal['indeks_persen'] == 100.0
    assert awal['bulan'][0]['label'] == 'Juli 2026'
    r = c.put('/api/pkg/kehadiran/g1', json={'bulan': [{'bulan': '2026-08', 'tanpa_keterangan': 1, 'terlambat_menit': 240}]})
    assert r.status_code == 200 and r.json()['indeks_persen'] == round((1 - 1.5 / 46) * 100, 2)
    assert d['pkg_audit_log'].docs[-1]['aksi'] == 'kehadiran_simpan'
    assert c.put('/api/pkg/kehadiran/g1', json={'bulan': [{'bulan': '2027-01'}]}).status_code == 422  # di luar periode
    assert c.put('/api/pkg/kehadiran/g1', json={'bulan': [{'bulan': '2026-08'}, {'bulan': '2026-08'}]}).status_code == 422
    assert c.put('/api/pkg/kehadiran/g1', json={'bulan': [{'bulan': '2026-08', 'tanpa_keterangan': 40}]}).status_code == 422
    c.put('/api/pkg/kehadiran/g1', json={'bulan': [{'bulan': '2026-10', 'tanpa_keterangan': 2}]})  # mengganti data lama
    assert len(d['pkg_kehadiran'].docs) == 1 and c.get('/api/pkg/kehadiran/g1').json()['jumlah']['tanpa_keterangan'] == 2
    ringkas = {g['guru_id']: g for g in c.get('/api/pkg/kehadiran/ringkasan').json()['guru']}
    assert ringkas['g1']['terisi'] and not ringkas['g2']['terisi'] and ringkas['g2']['indeks_persen'] == 100.0
    pengguna['user'] = {'id': 'pen', 'roles': ['guru']}
    assert c.put('/api/pkg/kehadiran/g2', json={'bulan': []}).status_code == 403
    d['pkg_periode'].docs[0]['status'] = 'ditutup'
    assert c.put('/api/pkg/kehadiran/g1', json={'periode_id': 'p1', 'bulan': []}).status_code == 409


def test_nilai_akhir_05_rkap_pkgr():
    from pkg_nilai_akhir import hitung_nilai_akhir
    skor = {'per_jenis': {'sejawat': {'nilai_tepat': 90.0}, 'siswa': {'nilai_tepat': 85.0}, 'ortu': {'nilai_tepat': 100.0}}}
    hadir = {'indeks': 1 - 1.5 / 46, 'terisi': True, 'jumlah': {'jumlah_tidak_hadir': 1.5}}
    r = hitung_nilai_akhir(80.0, skor, hadir, jenjang='MTs')
    assert [k['nilai'] for k in r['komponen']] == [56.0, 9.0, 8.5, 10.0]  # F35:F38
    assert r['nilai_kompetensi'] == 83.5 and r['nilai_akhir'] == round(83.5 * (1 - 1.5 / 46), 2)  # F40, E43
    assert r['level_teks'] == 'pada Level 4' and r['predikat'].endswith('Ahli Madya') and r['lengkap']
    tanpa = hitung_nilai_akhir(80.0, {'per_jenis': {}}, None)
    assert tanpa['nilai_akhir'] == 56.0 and tanpa['indeks_kehadiran'] == 1.0 and not tanpa['kehadiran_terisi']


def test_isi_sheet_indek_hdir():
    import io
    from openpyxl import load_workbook
    import pkg_cetak
    import pkg_template
    baris = k.baris_lengkap(k.bulan_periode(PERIODE), [{'bulan': '2026-08', 'terlambat_menit': 30, 'pulang_cepat_menit': 15, 'tanpa_keterangan': 1}])
    wb = load_workbook(io.BytesIO(pkg_template.isi_template_ke_bytes('instrumen', lambda w: pkg_cetak.isi_kehadiran(w, baris))))
    ws = wb['INDEK-HDIR']
    assert ws['E9'].value == 'Guru' and ws['C16'].value == 'Juli 2026' and ws['C21'].value == 'Desember 2026'
    assert (ws['D17'].value, ws['E17'].value, ws['H17'].value) == (30, 15, 1)
    assert ws['C22'].value is None and ws['D27'].value == 0 and ws['H27'].value == 0  # baris bulan sisa dikosongkan
    assert ws['K29'].value == '=K28' and ws['G17'].value == '=F17/(8*60)'  # rumus template utuh


def test_bulan_kehadiran_mengikuti_semester():
    import asyncio
    sebulan = {'id': 'p2', 'tanggal_buka': '2026-10-01', 'tanggal_tutup': '2026-10-31', 'tahun_pelajaran': '2026/2027'}
    semester = ['2026-07', '2026-08', '2026-09', '2026-10', '2026-11', '2026-12']
    db = Db({'semesters': Koleksi([{'id': 'sm1', 'name': 'Ganjil', 'start_date': '2026-07-13', 'end_date': '2026-12-19'},
                                   {'id': 'sm2', 'name': 'Genap'}])})
    # 1) tanggal semester terkait; 2) nama semester + tahun pelajaran; 3) tanggal buka-tutup periode
    assert asyncio.run(k.bulan_kehadiran(db, {**sebulan, 'semester_id': 'sm1'})) == semester
    assert asyncio.run(k.bulan_kehadiran(db, {**sebulan, 'semester_id': 'sm2'})) == ['2027-01', '2027-02', '2027-03', '2027-04', '2027-05', '2027-06']
    assert asyncio.run(k.bulan_kehadiran(db, {**sebulan, 'semester': 'Ganjil'})) == semester
    assert asyncio.run(k.bulan_kehadiran(db, sebulan)) == ['2026-10']


def test_input_kehadiran_periode_sebulan(klien):
    c, d, _ = klien
    d['pkg_periode'].docs[0].update({'tanggal_buka': '2026-10-01', 'tanggal_tutup': '2026-10-31', 'semester': 'Ganjil',
                                     'tahun_pelajaran': '2026/2027'})
    data = c.get('/api/pkg/kehadiran/g1').json()
    assert [b['bulan'] for b in data['bulan']] == ['2026-07', '2026-08', '2026-09', '2026-10', '2026-11', '2026-12']
    assert c.put('/api/pkg/kehadiran/g1', json={'bulan': [{'bulan': '2026-07', 'tanpa_keterangan': 1}]}).status_code == 200
