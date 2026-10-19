"""Impor instrumen dari Excel: berkas resmi terbaca utuh; berkas rusak dilaporkan per lokasi sel."""

import io
import os
import sys

import pytest
from openpyxl import Workbook

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'backend'))

import pkg_instrumen_impor as impor  # noqa: E402
import pkg_template  # noqa: E402


def test_template_resmi_terbaca_utuh():
    p = impor.pratinjau_dari_template()
    assert p['cocok_resmi'] is True and p['sumber'] == 'template' and p['peringatan'] == []
    k = p['kompetensi']
    assert [x['nama'] for x in k] == ['Kompetensi Pedagogik', 'Kompetensi Kepribadian', 'Kompetensi Sosial', 'Kompetensi Profesional']
    sub = [s for x in k for s in x['sub']]
    assert [s['kode'] for s in sub] == [f'{a}.{b}' for a in range(1, 5) for b in range(1, 4)]
    assert all([i['level'] for i in s['indikator']] == [1, 2, 3, 4, 5] for s in sub)
    assert sum(len(s['indikator']) for s in sub) == 60
    assert sub[0]['nama'] == 'Lingkungan pembelajaran yang aman dan nyaman bagi peserta didik'
    assert sub[3]['nama'].startswith('Kematangan moral') and not sub[3]['nama'].endswith('.')
    assert not any(i['teks'].endswith(';') for s in sub for i in s['indikator'])


def _berkas(isi_sheet):
    """Workbook buatan: {nama_sheet: [(baris, kolom, nilai)]}."""
    wb = Workbook()
    wb.remove(wb.active)
    for nama, sel in isi_sheet.items():
        ws = wb.create_sheet(nama)
        for r, c, v in sel:
            ws.cell(r, c, v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _sheet(kode, n_indikator=5, kosong_level=None, kode_indikator_salah=False):
    sel = [(3, 2, f'{kode}. NAMA {kode}'), (5, 2, f'{kode}.1 Sub satu')]
    for lv in range(1, n_indikator + 1):
        kode_ind = f'{lv}.2' if kode_indikator_salah and lv == 1 else f'{lv}.1'
        sel += [(6 + lv, 2, 'v'), (6 + lv, 3, kode_ind), (6 + lv, 4, '' if lv == kosong_level else f'Teks level {lv};')]
    return sel


def test_berkas_buatan_sah():
    isi = _berkas({n: _sheet(str(i)) for i, n in enumerate(pkg_template.SHEET_KOMPETENSI, 1)})
    p = impor.pratinjau_dari_bytes(isi, 'unggah', 'instrumen.xlsx')
    assert p['cocok_resmi'] is False and p['peringatan'] and len(p['kompetensi']) == 4
    assert p['kompetensi'][0] == {'kode': '1', 'nama': 'Kompetensi Nama 1', 'sub': [{
        'kode': '1.1', 'nama': 'Sub satu', 'bobot': 1.0,
        'indikator': [{'level': lv, 'teks': f'Teks level {lv}'} for lv in range(1, 6)]}]}


def test_berkas_rusak_dilaporkan_per_lokasi():
    sheets = {n: _sheet(str(i)) for i, n in enumerate(pkg_template.SHEET_KOMPETENSI, 1)}
    sheets['KOMP (2)'] = _sheet('2', kosong_level=3)
    sheets['KOMP (3)'] = _sheet('3', n_indikator=4)
    sheets['KOMP (4)'] = _sheet('4', kode_indikator_salah=True)
    del sheets['KOMP-1']
    with pytest.raises(impor.GalatImporInstrumen) as e:
        impor.baca_instrumen_excel(_berkas(sheets))
    lokasi = {g['lokasi']: g['pesan'] for g in e.value.galat}
    assert 'KOMP-1' in lokasi
    assert 'KOMP (2)!D9' in lokasi and 'Level 3' in lokasi['KOMP (2)!D9']
    assert any(k.startswith('KOMP (3)!B') and '4 indikator' in v for k, v in lokasi.items())
    assert 'KOMP (4)!C7' in lokasi


@pytest.mark.parametrize('isi, kata', [(b'', 'kosong'), (b'bukan excel', 'bukan Excel'), (b'x' * (10 * 1024 * 1024 + 1), '10 MB')])
def test_berkas_tidak_terbaca(isi, kata):
    with pytest.raises(impor.GalatImporInstrumen) as e:
        impor.baca_instrumen_excel(isi)
    assert kata in e.value.pesan
