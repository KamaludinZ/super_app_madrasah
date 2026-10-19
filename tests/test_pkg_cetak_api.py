"""API cetak Excel: otorisasi & cakupan, rekap & laporan per guru dari template resmi, riwayat unduhan."""

import asyncio
import io
import os
import sys

import pytest
from openpyxl import load_workbook

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'backend'))

import pkg_instrumen_impor  # noqa: E402
import pkg_rekap  # noqa: E402
from routers import pkg_cetak_api as api  # noqa: E402
from tests.pkg_palsu import Db, Koleksi  # noqa: E402

INSTRUMEN = {'id': 'i1', 'versi_template': '2026', 'revisi': 1, 'is_active': True,
             'kompetensi': pkg_instrumen_impor.pratinjau_dari_template()['kompetensi']}
SUB = [s['kode'] for k in INSTRUMEN['kompetensi'] for s in k['sub']]


def _lembar(lid, guru, penilai, status='terkunci', skor=4):
    return {'id': lid, 'status': status, 'periode_id': 'p1', 'sk_id': 's1', 'guru_dinilai_id': guru, 'penilai_id': penilai,
            'instrumen_id': 'i1', 'snapshot_guru': {'nama': f'Guru {guru}', 'nip': f'NIP-{guru}'},
            'isian': {k: {'centang': list(range(1, skor + 1))} for k in SUB}, 'terkunci_pada': '2026-10-18T03:00:00+00:00'}


@pytest.fixture
def klien(monkeypatch):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from core import get_current_user
    d = Db({
        'pkg_periode': Koleksi([{'id': 'p1', 'status': 'berjalan', 'nama_periode': 'PKG Ganjil 2026/2027', 'semester': 'Ganjil',
                                 'tahun_pelajaran': '2026/2027', 'tanggal_buka': '2026-10-01', 'tanggal_tutup': '2026-11-30'}]),
        'pkg_sk': Koleksi([{'id': 's1', 'nomor_sk': 'B-1/2026', 'tanggal_sk': '2026-09-28', 'periode_id': 'p1', 'is_active': True,
                            'penilai_ids': ['pen'], 'dinilai_ids': ['g1', 'g2', 'g3'],
                            'snapshot_anggota': {'pen': {'nama': 'Penilai Satu', 'nip': 'NIP-pen'}}}]),
        'pkg_penugasan': Koleksi([{'periode_id': 'p1', 'penilai_id': 'pen', 'guru_dinilai_ids': ['g1', 'g2'], 'supervisi_selesai': True}]),
        'pkg_penilaian': Koleksi([_lembar('n1', 'g1', 'pen', skor=4), _lembar('n2', 'g2', 'pen', 'draf'), _lembar('n3', 'g3', 'kep', skor=5)]),
        'pkg_instrumen': Koleksi([INSTRUMEN]),
        'settings': Koleksi([{'id': 'global_config', 'headmaster_name': 'Drs. H. Mahmud, M.Pd.', 'school_name': 'MTsN 2 Kota Malang'}]),
    })
    monkeypatch.setattr(api, 'db', d)
    asyncio.run(pkg_rekap.bangun_ulang_rekap(d))
    app = FastAPI()
    app.include_router(api.router, prefix='/api')
    pengguna = {'user': {'id': 'kep', 'roles': ['kepala_sekolah'], 'full_name': 'Kepala'}}
    app.dependency_overrides[get_current_user] = lambda: pengguna['user']
    return TestClient(app), d, pengguna


def test_cetak_rekap_excel_kepala(klien):
    k, d, _ = klien
    r = k.post('/api/pkg/cetak', json={'jenis': 'rekap', 'format': 'xlsx'})
    assert r.status_code == 200 and r.headers['x-pkg-mesin'] == 'openpyxl'
    assert 'Rekap_Hasil_PKG_PKG_Ganjil_2026_2027.xlsx' in r.headers['content-disposition']
    ws = load_workbook(io.BytesIO(r.content))['Rekap hasil PKG']
    assert [ws[f'B{i}'].value for i in (6, 7, 8)] == ['Guru g1', 'Guru g3', None]  # hanya lembar terkunci
    assert ws['F6'].value == 80.0 and ws['F7'].value == 100.0 and ws['G24'].value == 'Drs. H. Mahmud, M.Pd.'
    u = d['pkg_unduhan'].docs[0]
    assert (u['jenis'], u['format'], u['mesin'], u['periode_id'], u['aktor_id']) == ('rekap', 'xlsx', 'openpyxl', 'p1', 'kep')
    assert u['ukuran'] == len(r.content) and d['pkg_audit_log'].docs[-1]['aksi'] == 'cetak_template'


def test_cetak_laporan_guru_excel(klien):
    k, d, pengguna = klien
    pengguna['user'] = {'id': 'pen', 'roles': ['guru'], 'full_name': 'Penilai Satu'}
    r = k.post('/api/pkg/cetak', json={'jenis': 'guru', 'format': 'xlsx', 'guru_id': 'g1'})
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    data = wb['00-DATA']
    assert (data['E5'].value, data['E6'].value, data['E33'].value, data['E36'].value) == ('Guru g1', 'NIP-g1', 'Penilai Satu', 'B-1/2026')
    assert data['E15'].value == '2026' and data['E26'].value == 'MTsN 2 Kota Malang' and data['C41'].value == '1 Oktober 2026'
    komp = wb['KOMP (3)']
    assert [komp[f'B{r}'].value for r in range(20, 25)] == ['v', 'v', 'v', 'v', None]  # skor 4: Level 1–4
    assert wb['04-RKAP-PKG']['H44'].value == '=SUM(H29:H43)'  # rumus resmi tetap
    assert d['pkg_unduhan'].docs[-1]['guru_id'] == 'g1' and d['pkg_unduhan'].docs[-1]['jenis'] == 'guru'


def test_cetak_ditolak_dan_divalidasi(klien):
    k, d, pengguna = klien
    pengguna['user'] = {'id': 'pen', 'roles': ['guru']}
    assert k.post('/api/pkg/cetak', json={'jenis': 'guru', 'format': 'xlsx', 'guru_id': 'g3'}).status_code == 403  # di luar cakupan
    assert k.post('/api/pkg/cetak', json={'jenis': 'guru', 'format': 'xlsx', 'guru_id': 'g2'}).status_code == 409  # draf
    pengguna['user'] = {'id': 'g1', 'roles': ['guru']}  # hanya dinilai
    r = k.post('/api/pkg/cetak', json={'jenis': 'rekap', 'format': 'xlsx'})
    assert r.status_code == 403 and r.json()['detail']['alasan'] == 'bukan_penilai'
    pengguna['user'] = {'id': 'asing', 'roles': ['guru']}
    assert k.post('/api/pkg/cetak', json={'jenis': 'rekap', 'format': 'xlsx'}).json()['detail']['alasan'] == 'bukan_anggota_sk'
    for isian in ({'jenis': 'guru', 'format': 'xlsx'}, {'jenis': 'rekap', 'format': 'doc'}, {'jenis': 'rekap', 'format': 'xlsx', 'periode_id': 'p$1'}):
        assert k.post('/api/pkg/cetak', json=isian).status_code == 422, isian
    assert d['pkg_unduhan'].docs == []  # gagal tidak tercatat sebagai unduhan


@pytest.mark.skipif(not __import__('pkg_pdf').libreoffice_tersedia(), reason='LibreOffice tidak tersedia')
def test_cetak_pdf_libreoffice(klien):
    k, d, _ = klien
    r = k.post('/api/pkg/cetak', json={'jenis': 'guru', 'format': 'pdf', 'guru_id': 'g3'})
    assert r.status_code == 200 and r.content.startswith(b'%PDF') and r.headers['content-type'] == 'application/pdf'
    assert r.headers['x-pkg-mesin'] == 'libreoffice' and '.pdf' in r.headers['content-disposition']
    assert d['pkg_unduhan'].docs[-1]['mesin'] == 'libreoffice' and d['pkg_unduhan'].docs[-1]['format'] == 'pdf'


def test_cetak_pdf_tanpa_libreoffice_503(klien, monkeypatch):
    import pkg_pdf
    k, d, _ = klien
    monkeypatch.setattr(pkg_pdf, 'cari_soffice', lambda: None)
    r = k.post('/api/pkg/cetak', json={'jenis': 'rekap', 'format': 'pdf'})
    assert r.status_code == 503 and 'PDF tidak dapat dibuat' in r.json()['detail']
    assert d['pkg_unduhan'].docs == []
